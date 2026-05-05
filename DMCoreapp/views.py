from django.shortcuts import render
from .models import *
from django.shortcuts import redirect
from django.contrib import messages
from django.contrib.auth.hashers import make_password
from django.http import JsonResponse
from django.views.decorators.csrf import csrf_exempt
import json
from django.shortcuts import render,get_object_or_404
from datetime import datetime
from DMCoreapp.models import LogRegister_Details, EmployeeRegister_Details
from django.db import transaction, models
from django.utils import timezone
from django.views.decorators.http import require_POST
from django.core.paginator import Paginator
from django.db import models
from django.http import HttpResponse
from datetime import date
from datetime import datetime
from DMCoreapp.models import ActionTaken,Feedback
from django.shortcuts import render, redirect
from django.db.models import Q
from datetime import datetime
from itertools import chain
from .models import EmployeeRegister_Details, EmployeeLeave, ActionTaken, Complaints, Feedback
import json
from django.shortcuts import render, get_object_or_404, redirect
from django.core.serializers.json import DjangoJSONEncoder
from django.http import JsonResponse
from django.urls import reverse
from django.db.models import Exists, OuterRef
from .models import WorkRegister, Work_Task, ClientTask_Register,LeadCollection, LeadAllocation, EmployeeRegister_Details,ClientRegister
import openpyxl
from django.http import HttpResponse
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib import messages
from .models import LeadCollection, LeadRow, LeadRowValue, LeadField, LogRegister_Details, EmployeeRegister_Details
from django.core.paginator import Paginator
from django.http import JsonResponse
import json
from django.utils.timezone import now
from datetime import timedelta
from .models import FollowUpStatus
from django.db.models import Subquery
from django.db.models import Count
from django.db.models.functions import TruncDate



# Create your views here.


def index(request):
    return render(request,'index.html')

def login_view(request):
    if request.method == 'POST':
        username = request.POST.get('username')
        password = request.POST.get('password')
        
        try:
            user = LogRegister_Details.objects.get(
                log_username=username, 
                log_password=password
            )

            if not user.active_status:
                messages.error(request, 'Your account is pending approval. Please wait for admin approval.')
                return render(request, 'login.html')
            
            # ---------------------- ADMIN LOGIN ----------------------
            if user.position == 'Admin':
                request.session['user_id'] = user.id
                request.session['username'] = user.log_username
                request.session['position'] = 'Admin'
                request.session['is_authenticated'] = True
                
                try:
                    business = BusinessRegister_Details.objects.get(login=user)
                    request.session['company_name'] = business.company_name
                    request.session['owner_fname'] = business.owner_fname
                    request.session['owner_lname'] = business.owner_lname
                except BusinessRegister_Details.DoesNotExist:
                    messages.error(request, 'Business details not found')
                    return render(request, 'login.html')
                
                return redirect('admin_dashboard')

            # ---------------------- EMPLOYEE LOGIN ----------------------
            else:
                try:
                    employee = EmployeeRegister_Details.objects.get(login=user)

                    if employee.active_status != 'Approved':
                        messages.error(request, 'Your account is pending approval. Please wait for admin approval.')
                        return render(request, 'login.html')

                    # IMPORTANT FIX: Read live role from designation table
                    actual_role = employee.designation.dashboard_id

                    # Set all sessions using actual_role, not user.position
                    request.session['user_id'] = user.id
                    request.session['username'] = user.log_username
                    request.session['position'] = actual_role
                    request.session['is_authenticated'] = True
                    request.session['employee_id'] = employee.id
                    request.session['employee_name'] = employee.name
                    request.session['company_id'] = employee.company.id

                    # Redirect dynamically using new designation
                    if actual_role == 'Digital_Marketing_Head':
                        return redirect('head_dashboard')
                    elif actual_role == 'Team_Lead':
                        return redirect('teamlead_dashboard')
                    elif actual_role == 'Executive':
                        return redirect('executive_dashboard')
                    elif actual_role == 'Data_Manager':
                        return redirect('data_manager_dashboard')
                    elif actual_role == 'Telecaller':
                        return redirect('telecaller_dashboard')
                    else:
                        messages.error(request, 'No dashboard available for your role')
                        return render(request, 'login.html')
                        
                except EmployeeRegister_Details.DoesNotExist:
                    messages.error(request, 'Employee details not found')
                
        except LogRegister_Details.DoesNotExist:
            messages.error(request, 'Invalid username or password')
    
    return render(request, 'login.html')

def business_register(request):
    if request.method == 'POST':
        first_name = request.POST.get('first_name')
        last_name = request.POST.get('last_name')
        company_name = request.POST.get('company_name')
        company_id = request.POST.get('company_id')
        contact_number = request.POST.get('contact_number')
        email = request.POST.get('email')
        location = request.POST.get('location')
        website = request.POST.get('website')
        username = request.POST.get('username')
        password = request.POST.get('password')
        
        if not all([first_name, last_name, company_name, company_id, contact_number, email, location, username, password]):
            messages.error(request, 'Please fill all required fields')
            return render(request, 'business_register.html')
        
        if LogRegister_Details.objects.filter(log_username=username).exists():
            messages.error(request, 'Username already exists')
            return render(request, 'business_register.html')
        
        if BusinessRegister_Details.objects.filter(email=email).exists():
            messages.error(request, 'Email already registered')
            return render(request, 'business_register.html')
        
        if BusinessRegister_Details.objects.filter(company_code=company_id).exists():
            messages.error(request, 'Company ID already exists')
            return render(request, 'business_register.html')
        
        password_requirements = {
            'length': len(password) >= 8,
            'uppercase': any(c.isupper() for c in password),
            'lowercase': any(c.islower() for c in password),
            'digit': any(c.isdigit() for c in password),
            'special': any(not c.isalnum() for c in password)
        }
        
        if not all(password_requirements.values()):
            messages.error(request, 'Password must contain at least 8 characters, one uppercase letter, one lowercase letter, one digit, and one special character')
            return render(request, 'business_register.html')
        
        try:
            login_record = LogRegister_Details(
                log_username=username,
                log_password=password,
                position='Admin',
                is_staff=True,
                active_status=True
            )
            login_record.save()
            
            business_record = BusinessRegister_Details(
                login=login_record,
                owner_fname=first_name,
                owner_lname=last_name,
                company_name=company_name,
                company_code=company_id,
                contact_number=contact_number,
                email=email,
                website=website or '',
                location=location,
                active_status=True
            )
            business_record.save()
            
            messages.success(request, 'Registration successful! You can now login.')
            return redirect('login')
            
        except Exception as e:
            messages.error(request, f'Registration failed: {str(e)}')
    
    return render(request, 'business_register.html')

def admin_dashboard(request):
    
    user_id = request.session.get('user_id')
    username = request.session.get('username')
    
    try:
        user = LogRegister_Details.objects.get(id=user_id)
        business = BusinessRegister_Details.objects.get(login=user)
        
        context = {
            'user': user,
            'business': business,
            'owner_full_name': f"{business.owner_fname} {business.owner_lname}",
            'company_name': business.company_name,
            'login_requests_urls': ['login_requests', 'login_requests_detail'],
            'departments_urls': ['departments', 'department_detail', 'add_department'],
            'designations_urls': ['designations', 'designation_detail', 'add_designation'],
            'employees_urls': ['employees', 'employee_detail', 'add_employee'],
            'data_bank_urls': ['data_bank', 'data_detail', 'add_data'],
        }
        
        return render(request, 'admin_dashboard.html', context)
        
    except (LogRegister_Details.DoesNotExist, BusinessRegister_Details.DoesNotExist):
        messages.error(request, 'User or business details not found')
        return redirect('login')
    
@csrf_exempt
def check_company_id(request):
    company_id = request.GET.get('company_id', '').strip()
    
    if not company_id:
        return JsonResponse({'exists': False})

    exists = BusinessRegister_Details.objects.filter(company_code=company_id).exists()
    
    return JsonResponse({'exists': exists})

@csrf_exempt
def check_email(request):
    email = request.GET.get('email', '').strip()
    
    if not email:
        return JsonResponse({'exists': False})

    exists = BusinessRegister_Details.objects.filter(email=email).exists()
    
    return JsonResponse({'exists': exists})


def head_dashboard(request):
    user_id = request.session.get('user_id')
    
    if not user_id or request.session.get('position') != 'Digital_Marketing_Head':
        return redirect('login')
    
    try:
        user = LogRegister_Details.objects.get(id=user_id)
        employee = EmployeeRegister_Details.objects.get(login=user)
        
        context = {
            'head_name': employee.name,  
            'employee': employee,
            'user': user
        }
        
        return render(request, 'head_dashboard.html', context)
        
    except (LogRegister_Details.DoesNotExist, EmployeeRegister_Details.DoesNotExist):
        messages.error(request, 'Employee details not found')
        return redirect('login')

def teamlead_dashboard(request):
    user_id = request.session.get('user_id')
    
    if not user_id or request.session.get('position') != 'Team_Lead':
        return redirect('login')
    
    try:
        user = LogRegister_Details.objects.get(id=user_id)
        employee = EmployeeRegister_Details.objects.get(login=user)
        
        context = {
            'teamlead_name': employee.name,
            'employee': employee,
            'user': user
        }
        
        return render(request, 'teamlead_dashboard.html', context)
        
    except (LogRegister_Details.DoesNotExist, EmployeeRegister_Details.DoesNotExist):
        messages.error(request, 'Employee details not found')
        return redirect('login')

def executive_dashboard(request):
    user_id = request.session.get('user_id')
    
    if not user_id or request.session.get('position') != 'Executive':
        return redirect('login')
    
    try:
        user = LogRegister_Details.objects.get(id=user_id)
        employee = EmployeeRegister_Details.objects.get(login=user)
        
        context = {
            'executive_name': employee.name,
            'employee': employee,
            'user': user
        }
        
        return render(request, 'executive_dashboard.html', context)
        
    except (LogRegister_Details.DoesNotExist, EmployeeRegister_Details.DoesNotExist):
        messages.error(request, 'Employee details not found')
        return redirect('login')

def data_manager_dashboard(request):
    user_id = request.session.get('user_id')
    
    if not user_id or request.session.get('position') != 'Data_Manager':
        return redirect('login')
    
    try:
        user = LogRegister_Details.objects.get(id=user_id)
        employee = EmployeeRegister_Details.objects.get(login=user)
        
        context = {
            'datamanager_name': employee.name,
            'employee': employee,
            'user': user
        }
        
        return render(request, 'data_manager_dashboard.html', context)
        
    except (LogRegister_Details.DoesNotExist, EmployeeRegister_Details.DoesNotExist):
        messages.error(request, 'Employee details not found')
        return redirect('login')

def telecaller_dashboard(request):
    user_id = request.session.get('user_id')
    
    if not user_id or request.session.get('position') != 'Telecaller':
        return redirect('login')
    
    try:
        user = LogRegister_Details.objects.get(id=user_id)
        employee = EmployeeRegister_Details.objects.get(login=user)
        leads = LeadRow.objects.filter(transferred_to = employee).exclude(status = "Opened")
        opened_count=LeadRow.objects.filter(transferred_to = employee , status = "Opened", is_waste = False).count()
        closed_count=LeadRow.objects.filter(transferred_to = employee, status = "closed").count()
        all_leads_count=LeadRow.objects.filter(transferred_to = employee).count()
        waste_count=LeadRow.objects.filter(transferred_to = employee, is_waste = True ).count()
        joined_count=LeadRow.objects.filter(transferred_to = employee, status = "Joined").count() 
        
        context = {
            'telecaller_name': employee.name,
            'employee': employee,
            'user': user,
            'new_leads':leads,
            'opened_count':opened_count,
            'closed_count':closed_count,
            'waste_count':waste_count,
            'joined_count':joined_count,
            'all_leads_count':all_leads_count,
        }
        
        return render(request, 'telecaller_dashboard.html', context)
        
    except (LogRegister_Details.DoesNotExist, EmployeeRegister_Details.DoesNotExist):
        messages.error(request, 'Employee details not found')
        return redirect('login')
     
def departments(request):
    user_id = request.session.get('user_id')
    username = request.session.get('username')
    
    try:
        user = LogRegister_Details.objects.get(id=user_id)
        business = BusinessRegister_Details.objects.get(login=user)
        
        departments = DepartmentRegister_Details.objects.filter(business=business).order_by('created_at')
        
        context = {
            'user': user,
            'business': business,
            'owner_full_name': f"{business.owner_fname} {business.owner_lname}",
            'company_name': business.company_name,
            'departments': departments,
            'login_requests_urls': ['login_requests', 'login_requests_detail'],
            'departments_urls': ['departments', 'department_detail', 'add_department'],
            'designations_urls': ['designations', 'designation_detail', 'add_designation'],
            'employees_urls': ['employees', 'employee_detail', 'add_employee'],
            'data_bank_urls': ['data_bank', 'data_detail', 'add_data'],
        }
        
        return render(request, 'departments.html', context)
        
    except (LogRegister_Details.DoesNotExist, BusinessRegister_Details.DoesNotExist):
        messages.error(request, 'User or business details not found')
        return redirect('login')
    
def add_department(request):
    if request.method == 'POST':
        user_id = request.session.get('user_id')
        
        try:
            user = LogRegister_Details.objects.get(id=user_id)
            business = BusinessRegister_Details.objects.get(login=user)
            
            department_name = request.POST.get('departmentName')
            department_description = request.POST.get('departmentDescription')
            
            if not department_name:
                messages.error(request, 'Department name is required')
                return redirect('departments')
            
            if DepartmentRegister_Details.objects.filter(business=business, name=department_name).exists():
                messages.error(request, 'Department with this name already exists')
                return redirect('departments')
            
            department = DepartmentRegister_Details(
                business=business,
                name=department_name,
                description=department_description or '',
                active_status=True
            )
            department.save()
            
            # messages.success(request, 'Department added successfully!')
            return redirect('departments')
            
        except (LogRegister_Details.DoesNotExist, BusinessRegister_Details.DoesNotExist):
            messages.error(request, 'User or business details not found')
            return redirect('login')
    
    return redirect('departments')

def edit_department(request, department_id):
    if request.method == 'POST':
        user_id = request.session.get('user_id')
        
        try:
            user = LogRegister_Details.objects.get(id=user_id)
            business = BusinessRegister_Details.objects.get(login=user)
            
            department = get_object_or_404(DepartmentRegister_Details, id=department_id, business=business)
            
            department_name = request.POST.get('departmentName')
            department_description = request.POST.get('departmentDescription')
            
            if not department_name:
                messages.error(request, 'Department name is required')
                return redirect('departments')
            
            if DepartmentRegister_Details.objects.filter(business=business, name=department_name).exclude(id=department_id).exists():
                messages.error(request, 'Department with this name already exists')
                return redirect('departments')
            
            department.name = department_name
            department.description = department_description or ''
            department.save()
            
            # messages.success(request, 'Department updated successfully!')
            return redirect('departments')
            
        except (LogRegister_Details.DoesNotExist, BusinessRegister_Details.DoesNotExist):
            messages.error(request, 'User or business details not found')
            return redirect('login')
    
    return redirect('departments')

def delete_department(request, department_id):
    if request.method == 'POST':
        user_id = request.session.get('user_id')
        
        try:
            user = LogRegister_Details.objects.get(id=user_id)
            business = BusinessRegister_Details.objects.get(login=user)
            
            department = get_object_or_404(DepartmentRegister_Details, id=department_id, business=business)
            department_name = department.name
            department.delete()
            
            # messages.success(request, f'Department "{department_name}" deleted successfully!')
            return redirect('departments')
            
        except (LogRegister_Details.DoesNotExist, BusinessRegister_Details.DoesNotExist):
            messages.error(request, 'User or business details not found')
            return redirect('login')
    
    return redirect('departments')

def get_department_details(request, department_id):
    user_id = request.session.get('user_id')
    
    try:
        user = LogRegister_Details.objects.get(id=user_id)
        business = BusinessRegister_Details.objects.get(login=user)
        
        department = get_object_or_404(DepartmentRegister_Details, id=department_id, business=business)
        
        department_data = {
            'id': department.id,
            'name': department.name,
            'description': department.description,
            'active_status': department.active_status,
            'created_at': department.created_at.strftime('%b. %d, %Y') if department.created_at else 'N/A'
        }
        
        return JsonResponse(department_data)
        
    except (LogRegister_Details.DoesNotExist, BusinessRegister_Details.DoesNotExist):
        return JsonResponse({'error': 'User or business details not found'}, status=404)
    
def check_department_exists(request):
    if request.method == 'POST':
        user_id = request.session.get('user_id')
        
        try:
            user = LogRegister_Details.objects.get(id=user_id)
            business = BusinessRegister_Details.objects.get(login=user)
            
            data = json.loads(request.body)
            department_name = data.get('departmentName', '').strip()
            exclude_id = data.get('excludeId') 
            
            if not department_name:
                return JsonResponse({'exists': False, 'message': ''})
            
            query = DepartmentRegister_Details.objects.filter(
                business=business, 
                name__iexact=department_name  
            )

            if exclude_id:
                query = query.exclude(id=exclude_id)
            exists = query.exists()
            
            if exists:
                return JsonResponse({
                    'exists': True, 
                    'message': f'Department "{department_name}" already exists in your company.'
                })
            else:
                return JsonResponse({
                    'exists': False, 
                })
                
        except (LogRegister_Details.DoesNotExist, BusinessRegister_Details.DoesNotExist):
            print("User or business not found")  
            return JsonResponse({'error': 'User or business details not found'}, status=404)
        except Exception as e:
            print(f"Error: {e}")  
            return JsonResponse({'error': str(e)}, status=500)
    
    return JsonResponse({'error': 'Invalid request method'}, status=400)
    
def designations(request):
    user_id = request.session.get('user_id')
    username = request.session.get('username')
    
    try:
        user = LogRegister_Details.objects.get(id=user_id)
        business = BusinessRegister_Details.objects.get(login=user)
        
        departments = DepartmentRegister_Details.objects.filter(business=business, active_status=True)
        
        designations = DesignationRegister_Details.objects.filter(
            business=business
        ).select_related('department').order_by('department__name', 'name')
        
        departments_with_designations = []
        for department in departments:
            dept_designations = designations.filter(department=department)
            if dept_designations.exists():
                departments_with_designations.append({
                    'department': department,
                    'designations': dept_designations
                })
        
        context = {
            'user': user,
            'business': business,
            'owner_full_name': f"{business.owner_fname} {business.owner_lname}",
            'company_name': business.company_name,
            'departments': departments,
            'departments_with_designations': departments_with_designations,
            'login_requests_urls': ['login_requests', 'login_requests_detail'],
            'departments_urls': ['departments', 'department_detail', 'add_department'],
            'designations_urls': ['designations', 'designation_detail', 'add_designation'],
            'employees_urls': ['employees', 'employee_detail', 'add_employee'],
            'data_bank_urls': ['data_bank', 'data_detail', 'add_data'],
        }
        
        return render(request, 'designations.html', context)
        
    except (LogRegister_Details.DoesNotExist, BusinessRegister_Details.DoesNotExist):
        messages.error(request, 'User or business details not found')
        return redirect('login')
    
def add_designation(request):
    if request.method == 'POST':
        user_id = request.session.get('user_id')
        
        try:
            user = LogRegister_Details.objects.get(id=user_id)
            business = BusinessRegister_Details.objects.get(login=user)
            
            department_id = request.POST.get('departmentSelect')
            dashboard_id = request.POST.get('dashboardSelect')
            designation_name = request.POST.get('designationSelect')
            designation_description = request.POST.get('designationDescription')
            
            if not all([department_id, dashboard_id, designation_name]):
                messages.error(request, 'All fields except description are required')
                return redirect('designations')
            
            department = get_object_or_404(DepartmentRegister_Details, id=department_id, business=business)
            
            if DesignationRegister_Details.objects.filter(
                business=business,
                department=department,
                name=designation_name
            ).exists():
                messages.error(request, f'Designation "{designation_name}" already exists in {department.name} department')
                return redirect('designations')
            
            designation = DesignationRegister_Details(
                business=business,
                department=department,
                dashboard_id=dashboard_id,
                name=designation_name,
                description=designation_description or '',
                active_status=True
            )
            designation.save()
            
            # messages.success(request, 'Designation added successfully!')
            return redirect('designations')
            
        except (LogRegister_Details.DoesNotExist, BusinessRegister_Details.DoesNotExist):
            messages.error(request, 'User or business details not found')
            return redirect('login')
        except DepartmentRegister_Details.DoesNotExist:
            messages.error(request, 'Department not found')
            return redirect('designations')
    
    return redirect('designations')

def edit_designation(request, designation_id):
    if request.method == 'POST':
        user_id = request.session.get('user_id')
        
        try:
            user = LogRegister_Details.objects.get(id=user_id)
            business = BusinessRegister_Details.objects.get(login=user)
            
            designation = get_object_or_404(DesignationRegister_Details, id=designation_id, business=business)
            
            department_id = request.POST.get('departmentSelect')
            dashboard_id = request.POST.get('dashboardSelect')
            designation_name = request.POST.get('designationSelect')
            designation_description = request.POST.get('designationDescription')
            
            if not all([department_id, dashboard_id, designation_name]):
                messages.error(request, 'All fields except description are required')
                return redirect('designations')
            
            department = get_object_or_404(DepartmentRegister_Details, id=department_id, business=business)
            
            if DesignationRegister_Details.objects.filter(
                business=business,
                department=department,
                name=designation_name
            ).exclude(id=designation_id).exists():
                messages.error(request, f'Designation "{designation_name}" already exists in {department.name} department')
                return redirect('designations')
            
            designation.department = department
            designation.dashboard_id = dashboard_id
            designation.name = designation_name
            designation.description = designation_description or ''
            designation.save()
            
            # messages.success(request, 'Designation updated successfully!')
            return redirect('designations')
            
        except (LogRegister_Details.DoesNotExist, BusinessRegister_Details.DoesNotExist):
            messages.error(request, 'User or business details not found')
            return redirect('login')
        except DepartmentRegister_Details.DoesNotExist:
            messages.error(request, 'Department not found')
            return redirect('designations')
    
    return redirect('designations')

def delete_designation(request, designation_id):
    if request.method == 'POST':
        user_id = request.session.get('user_id')
        
        try:
            user = LogRegister_Details.objects.get(id=user_id)
            business = BusinessRegister_Details.objects.get(login=user)
            
            designation = get_object_or_404(DesignationRegister_Details, id=designation_id, business=business)
            designation_name = designation.name
            department_name = designation.department.name
            designation.delete()
            
            # messages.success(request, f'Designation "{designation_name}" from "{department_name}" department deleted successfully!')
            return redirect('designations')
            
        except (LogRegister_Details.DoesNotExist, BusinessRegister_Details.DoesNotExist):
            messages.error(request, 'User or business details not found')
            return redirect('login')
    
    return redirect('designations')

def get_designation_details(request, designation_id):
    user_id = request.session.get('user_id')
    
    try:
        user = LogRegister_Details.objects.get(id=user_id)
        business = BusinessRegister_Details.objects.get(login=user)
        
        designation = get_object_or_404(DesignationRegister_Details, id=designation_id, business=business)
        
        designation_data = {
            'id': designation.id,
            'department_id': designation.department.id,
            'dashboard_id': designation.dashboard_id,
            'name': designation.name,
            'description': designation.description,
            'active_status': designation.active_status,
            'created_at': designation.created_at.strftime('%b. %d, %Y') if designation.created_at else 'N/A'
        }
        
        return JsonResponse(designation_data)
        
    except (LogRegister_Details.DoesNotExist, BusinessRegister_Details.DoesNotExist):
        return JsonResponse({'error': 'User or business details not found'}, status=404)

@csrf_exempt
def check_designation_exists(request):
    if request.method == 'POST':
        user_id = request.session.get('user_id')
        
        try:
            user = LogRegister_Details.objects.get(id=user_id)
            business = BusinessRegister_Details.objects.get(login=user)
            
            data = json.loads(request.body)
            department_id = data.get('departmentId')
            designation_name = data.get('designationName', '').strip()
            exclude_id = data.get('excludeId')  
            
            if not department_id or not designation_name:
                return JsonResponse({'exists': False, 'message': ''})
            
            department = get_object_or_404(DepartmentRegister_Details, id=department_id, business=business)
            
            query = DesignationRegister_Details.objects.filter(
                business=business,
                department=department,
                name__iexact=designation_name
            )
            
            if exclude_id:
                query = query.exclude(id=exclude_id)
            
            exists = query.exists()
            
            if exists:
                return JsonResponse({
                    'exists': True, 
                    'message': f'Designation "{designation_name}" already exists in {department.name} department.'
                })
            else:
                return JsonResponse({
                    'exists': False,
                })
                
        except (LogRegister_Details.DoesNotExist, BusinessRegister_Details.DoesNotExist):
            return JsonResponse({'error': 'User or business details not found'}, status=404)
        except DepartmentRegister_Details.DoesNotExist:
            return JsonResponse({'error': 'Department not found'}, status=404)
        except Exception as e:
            return JsonResponse({'error': str(e)}, status=500)
    
    return JsonResponse({'error': 'Invalid request method'}, status=400)

def login_requests(request):
    user_id = request.session.get('user_id')
    username = request.session.get('username')
    
    try:
        user = LogRegister_Details.objects.get(id=user_id)
        business = BusinessRegister_Details.objects.get(login=user)
        
        pending_employees = EmployeeRegister_Details.objects.filter(
            company=business,  
            active_status='Pending'
        ).select_related('company', 'department', 'designation', 'login')
        
        context = {
            'user': user,
            'business': business,
            'owner_full_name': f"{business.owner_fname} {business.owner_lname}",
            'company_name': business.company_name,
            'pending_employees': pending_employees,
            'login_requests_urls': ['login_requests', 'login_requests_detail'],
            'departments_urls': ['departments', 'department_detail', 'add_department'],
            'designations_urls': ['designations', 'designation_detail', 'add_designation'],
            'employees_urls': ['employees', 'employee_detail', 'add_employee'],
            'data_bank_urls': ['data_bank', 'data_detail', 'add_data'],
        }
        
        return render(request, 'login_requests.html', context)
        
    except (LogRegister_Details.DoesNotExist, BusinessRegister_Details.DoesNotExist):
        messages.error(request, 'User or business details not found')
        return redirect('login')
    
def check_company(request):
    company_id = request.GET.get('company_id', '').strip()
    exists = BusinessRegister_Details.objects.filter(company_code=company_id, active_status=True).exists()
    return JsonResponse({'exists': exists})

def get_departments(request):
    company_id = request.GET.get('company_id', '').strip()
    departments = []
    
    try:
        company = BusinessRegister_Details.objects.get(company_code=company_id, active_status=True)
        dept_objs = DepartmentRegister_Details.objects.filter(business=company, active_status=True)
        departments = [{'id': dept.id, 'name': dept.name} for dept in dept_objs]
    except BusinessRegister_Details.DoesNotExist:
        pass
    
    return JsonResponse({'departments': departments})

def get_designations(request):
    company_id = request.GET.get('company_id', '').strip()
    department_id = request.GET.get('department_id', '').strip()
    designations = []
    
    try:
        company = BusinessRegister_Details.objects.get(company_code=company_id, active_status=True)
        department = DepartmentRegister_Details.objects.get(id=department_id, business=company, active_status=True)
        designation_objs = DesignationRegister_Details.objects.filter(
            business=company, 
            department=department, 
            active_status=True
        )
        designations = [{'id': desig.id, 'name': desig.name} for desig in designation_objs]
    except (BusinessRegister_Details.DoesNotExist, DepartmentRegister_Details.DoesNotExist):
        pass
    
    return JsonResponse({'designations': designations})

def employee_register(request):
    if request.method == 'POST':
        company_id = request.POST.get('company_id')
        department_id = request.POST.get('department')
        designation_id = request.POST.get('designation')
        employee_name = request.POST.get('employee_name')
        contact_number = request.POST.get('contact_number')
        email = request.POST.get('email')
        username = request.POST.get('username')
        password = request.POST.get('password')
        
        try:
            company = BusinessRegister_Details.objects.get(company_code=company_id, active_status=True)
            
            department = DepartmentRegister_Details.objects.get(id=department_id, business=company, active_status=True)
            
            designation = DesignationRegister_Details.objects.get(id=designation_id, business=company, department=department, active_status=True)
            
            if EmployeeRegister_Details.objects.filter(email=email).exists():
                messages.error(request, 'Email already registered')
                return render(request, 'employee_register.html')
            
            if LogRegister_Details.objects.filter(log_username=username).exists():
                messages.error(request, 'Username already exists')
                return render(request, 'employee_register.html')
            
            password_requirements = {
                'length': len(password) >= 8,
                'uppercase': any(c.isupper() for c in password),
                'digit': any(c.isdigit() for c in password),
                'special': any(not c.isalnum() for c in password)
            }
            
            if not all(password_requirements.values()):
                messages.error(request, 'Password must contain at least 8 characters, one uppercase letter, one digit, and one special character')
                return render(request, 'employee_register.html')
            
            login_record = LogRegister_Details(
                log_username=username,
                log_password=password,
                position=designation.dashboard_id,  
                is_staff=False,
                active_status=False  
            )
            login_record.save()
            
            employee_count = EmployeeRegister_Details.objects.filter(company=company).count()
            employee_code = f"EMP{employee_count + 1:03d}"
            
            employee_record = EmployeeRegister_Details(
                login=login_record,
                company=company,
                department=department,
                designation=designation,
                name=employee_name,
                employee_code=employee_code,
                contact_number=contact_number,
                email=email,
                active_status='Pending',  
                verification_status=False
            )
            employee_record.save()
            
            messages.success(request, 'Registration request submitted successfully! Please wait for admin approval.', extra_tags='registration_success')
            return redirect('login')
            
        except BusinessRegister_Details.DoesNotExist:
            messages.error(request, 'Invalid company ID')
        except DepartmentRegister_Details.DoesNotExist:
            messages.error(request, 'Invalid department')
        except DesignationRegister_Details.DoesNotExist:
            messages.error(request, 'Invalid designation')
        except Exception as e:
            messages.error(request, f'Registration failed: {str(e)}')
    
    return render(request, 'employee_register.html')

def approve_employee_request(request, employee_id):
    if request.method == 'POST':
        try:
            user_id = request.session.get('user_id')
            admin_business = BusinessRegister_Details.objects.get(login_id=user_id)
            
            employee = EmployeeRegister_Details.objects.get(
                id=employee_id,
                company=admin_business  
            )

            employee.active_status = 'Approved'
            employee.verification_status = True
            employee.save()
            
            login_record = employee.login
            login_record.active_status = True
            login_record.save()
            
            
        except EmployeeRegister_Details.DoesNotExist:
            messages.error(request, 'Employee not found or you do not have permission to approve this employee')
        except BusinessRegister_Details.DoesNotExist:
            messages.error(request, 'Business details not found')
        except Exception as e:
            messages.error(request, f'Approval failed: {str(e)}')
    
    return redirect('login_requests')

def decline_employee_request(request, employee_id):
    if request.method == 'POST':
        try:
            user_id = request.session.get('user_id')
            admin_business = BusinessRegister_Details.objects.get(login_id=user_id)
            
            employee = EmployeeRegister_Details.objects.get(
                id=employee_id,
                company=admin_business  
            )
            
            employee_name = employee.name
            login_record = employee.login
            
            employee.delete()
            login_record.delete()
            
        except EmployeeRegister_Details.DoesNotExist:
            messages.error(request, 'Employee not found or you do not have permission to decline this employee')
        except BusinessRegister_Details.DoesNotExist:
            messages.error(request, 'Business details not found')
        except Exception as e:
            messages.error(request, f'Decline failed: {str(e)}')
    
    return redirect('login_requests')

@csrf_exempt
def check_email_employee(request):
    email = request.GET.get('email', '').strip()
    
    if not email:
        return JsonResponse({'exists': False})

    exists = EmployeeRegister_Details.objects.filter(email=email).exists()
    
    return JsonResponse({'exists': exists})
    
def dmh_works(request):
    user_id = request.session.get('user_id')
    
    if not user_id or request.session.get('position') != 'Digital_Marketing_Head':
        return redirect('login')
    
    try:
        user = LogRegister_Details.objects.get(id=user_id)
        employee = EmployeeRegister_Details.objects.get(login=user)
        
        context = {
            'head_name': employee.name,  
            'employee': employee,
            'user': user
        }
        
        return render(request, 'dmh_works.html', context)
        
    except (LogRegister_Details.DoesNotExist, EmployeeRegister_Details.DoesNotExist):
        messages.error(request, 'Employee details not found')
        return redirect('login')

def task_page(request):
    user_id = request.session.get('user_id')
    if not user_id or request.session.get('position') != 'Digital_Marketing_Head':
        return redirect('login')

    user = LogRegister_Details.objects.filter(id=user_id).first()
    employee = EmployeeRegister_Details.objects.filter(login=user).first()

    if not user or not employee:
        messages.error(request, 'Employee details not found')
        return redirect('login')

    # Fetch the correct company id from session
    comp_taskid = request.session.get('company_id')

    if request.method == "POST":
        task_id = request.POST.get("task_id")
        task_name = request.POST.get("task_name")
        task_description = request.POST.get("task_description")

        if task_id:  # If editing
            task = Work_Task.objects.get(id=task_id)
            task.task_name = task_name
            task.task_description = task_description
            task.save()
        else:  # If adding
            Work_Task.objects.create(
                task_name=task_name,
                task_description=task_description,
                comp_taskid_id=comp_taskid,  
            )
        messages.success(request, "Task added successfully!")
        return redirect('task_page')

    tasks = Work_Task.objects.filter(comp_taskid_id=comp_taskid).order_by('id')
    
    # If no tasks exist, create a default one
    if not tasks.exists():
        Work_Task.objects.create(
            task_name="Lead Collection",
            task_description="Effiecient Lead Collection is the corner stone of successful business growth",
            comp_taskid_id=comp_taskid
        )
        tasks = Work_Task.objects.filter(comp_taskid_id=comp_taskid).order_by('id')

    context = {
        'head_name': employee.name,
        'employee': employee,
        'user': user,
        'tasks': tasks,
    }

    return render(request, "task.html", context)

def delete_task(request, id):
    try:
        task = Work_Task.objects.get(id=id)
        task.delete()
        messages.success(request, "Task deleted successfully!")
    except Work_Task.DoesNotExist:
        messages.error(request, "Task not found!")
    return redirect('task_page')

def register_client(request):
    user_id = request.session.get('user_id')
    if not user_id or request.session.get('position') != 'Digital_Marketing_Head':
        return redirect('login')

    user = LogRegister_Details.objects.filter(id=user_id).first()
    employee = EmployeeRegister_Details.objects.filter(login=user).first()

    if not user or not employee:
        messages.error(request, 'Employee details not found')
        return redirect('login')

    comp_id = employee.company.id if employee.company else None

    clients = ClientRegister.objects.filter(compId_id=comp_id).order_by('-id')
    tasks = Work_Task.objects.filter(comp_taskid_id=comp_id).order_by('id')

    context = {
        'head_name': employee.name,
        'employee': employee,
        'user': user,
        'clients': clients,
        'tasks': tasks,
    }

    return render(request, "register_client.html", context)

def add_client(request):
    if request.method == "POST":
        try:
            # Get the logged-in employee
            user_id = request.session.get('user_id')
            employee = EmployeeRegister_Details.objects.filter(login_id=user_id).first()

            if not employee or not employee.company:
                return JsonResponse({'status': 'error', 'message': 'Employee or company not found'}, status=400)

            company = employee.company

            # Create the client
            client = ClientRegister(
                compId=company,  # directly assign the company object
                client_name=request.POST.get('client_name'),
                client_email_primary=request.POST.get('client_email_primary'),
                client_email_alter=request.POST.get('client_email_alter'),
                client_phone=request.POST.get('client_phone'),
                client_phone_alter=request.POST.get('client_phone_alter'),
                client_address1=request.POST.get('client_address1'),
                client_address2=request.POST.get('client_address2'),
                client_address3=request.POST.get('client_address3'),
                client_place=request.POST.get('client_place'),
                client_district=request.POST.get('client_district'),
                client_state=request.POST.get('client_state'),
                client_profile=request.FILES.get('client_profile'),

                client_bussiness_name=request.POST.get('client_bussiness_name'),
                client_bussiness_email_primary=request.POST.get('client_bussiness_email_primary'),
                client_bussiness_email_alter=request.POST.get('client_bussiness_email_alter'),
                client_bussiness_phone=request.POST.get('client_bussiness_phone'),
                client_bussiness_phone_alter=request.POST.get('client_bussiness_phone_alter'),
                client_bussiness_website=request.POST.get('client_bussiness_website'),
                client_bussiness_address1=request.POST.get('client_bussiness_address1'),
                client_bussiness_address2=request.POST.get('client_bussiness_address2'),
                client_bussiness_address3=request.POST.get('client_bussiness_address3'),
                client_bussiness_place=request.POST.get('client_bussiness_place'),
                client_bussiness_district=request.POST.get('client_bussiness_district'),
                client_bussiness_state=request.POST.get('client_bussiness_state'),
                client_bussiness_files=request.FILES.get('client_bussiness_files'),
                bussiness_logo=request.FILES.get('bussiness_logo'),
                more_description=request.POST.get('more_description')
            )
            client.save()

            return JsonResponse({'status': 'success', 'message': 'Client registered successfully'})

        except Exception as e:
            print("Error in add_client:", e)
            return JsonResponse({'status': 'error', 'message': str(e)}, status=500)

    return JsonResponse({'status': 'error', 'message': 'Invalid request'}, status=400)

def get_client_data(request, id):
    client = ClientRegister.objects.get(id=id)
    data = {
        "client_name": client.client_name,
        "client_email": client.client_email_primary,
        "client_alt_email": client.client_email_alter,
        "client_phone": client.client_phone,
        "client_alt_phone": client.client_phone_alter,
        "client_address1": client.client_address1,
        "client_address2": client.client_address2,
        "client_address3": client.client_address3,
        "client_place": client.client_place,
        "client_district": client.client_district,
        "client_state": client.client_state,
        "client_profile": client.client_profile.url if client.client_profile else "",

        "business_name": client.client_bussiness_name,
        "business_email": client.client_bussiness_email_primary,
        "business_alt_email": client.client_bussiness_email_alter,
        "business_phone": client.client_bussiness_phone,
        "business_alt_phone": client.client_bussiness_phone_alter,
        "business_website": client.client_bussiness_website,
        "business_address1": client.client_bussiness_address1,
        "business_address2": client.client_bussiness_address2,
        "business_address3": client.client_bussiness_address3,
        "business_place": client.client_bussiness_place,
        "business_district": client.client_bussiness_district,
        "business_state": client.client_bussiness_state,
        "more_about": client.more_description,
        "business_logo" : client.bussiness_logo.url if client.bussiness_logo else "",
        "business_file" : client.client_bussiness_files.url if client.client_bussiness_files else ""
    }
    return JsonResponse(data)

def update_client(request, id):
    client = ClientRegister.objects.get(id=id)

    if request.method == "POST":
        client.client_name = request.POST.get("client_name")
        client.client_email_primary = request.POST.get("client_email")
        client.client_email_alter = request.POST.get("client_alt_email")
        client.client_phone = request.POST.get("client_phone")
        client.client_phone_alter = request.POST.get("client_alt_phone")
        client.client_address1 = request.POST.get("client_address1")
        client.client_address2 = request.POST.get("client_address2")
        client.client_address3 = request.POST.get("client_address3")
        client.client_place = request.POST.get("client_place")
        client.client_district = request.POST.get("client_district")
        client.client_state = request.POST.get("client_state")

        client.client_bussiness_name = request.POST.get("business_name")
        client.client_bussiness_email_primary = request.POST.get("business_email")
        client.client_bussiness_email_alter = request.POST.get("business_alt_email")
        client.client_bussiness_phone = request.POST.get("business_phone")
        client.client_bussiness_phone_alter = request.POST.get("business_alt_phone")
        client.client_bussiness_website = request.POST.get("business_website")
        client.client_bussiness_address1 = request.POST.get("business_address1")
        client.client_bussiness_address2 = request.POST.get("business_address2")
        client.client_bussiness_address3 = request.POST.get("business_address3")
        client.client_bussiness_place = request.POST.get("business_place")
        client.client_bussiness_district = request.POST.get("business_district")
        client.client_bussiness_state = request.POST.get("business_state")

        client.more_description = request.POST.get("more_about")

        # File Handling
        if request.FILES.get("business_logo"):
            client.bussiness_logo = request.FILES["business_logo"]  # Correct field
        if request.FILES.get("client_profile"):
            client.client_profile = request.FILES["client_profile"]

        client.save()
        return JsonResponse({"status": "success"})

    return JsonResponse({"status": "invalid request"}, status=400)

def check_client_field(request):
    value = request.GET.get("value", "").strip()
    field_type = request.GET.get("type", "").strip()  # "email", "phone", or "website"
    client_id = request.GET.get("client_id")  # optional, only for edit modal

    qs = ClientRegister.objects.all()
    if client_id:
        qs = qs.exclude(id=client_id)  # ignore current client when editing

    exists = False

    if field_type == "email":
        exists = (
            qs.filter(client_email_primary=value).exists() or
            qs.filter(client_email_alter=value).exists() or
            qs.filter(client_bussiness_email_primary=value).exists() or
            qs.filter(client_bussiness_email_alter=value).exists()
        )
    elif field_type == "phone":
        exists = (
            qs.filter(client_phone=value).exists() or
            qs.filter(client_phone_alter=value).exists() or
            qs.filter(client_bussiness_phone=value).exists() or
            qs.filter(client_bussiness_phone_alter=value).exists()
        )
    elif field_type == "website":
        exists = qs.filter(client_bussiness_website=value).exists()

    return JsonResponse({"exists": exists})

def delete_client(request, client_id):
    if request.method == "POST":
        try:
            client = ClientRegister.objects.get(id=client_id)
            client.delete()
            return JsonResponse({"status": "success"})
        except ClientRegister.DoesNotExist:
            return JsonResponse({"status": "error", "message": "Client not found"})
    return JsonResponse({"status": "error", "message": "Invalid request"})

def create_work(request, client_id):
    if request.method == "POST":
        client = ClientRegister.objects.get(id=client_id)
        comp = BusinessRegister_Details.objects.get(id=request.POST.get("comp_id"))

        # Create Work
        work = WorkRegister.objects.create(
            clientId=client,
            wcompId=comp,
            work_create_date=request.POST.get("work_create_date"),
            work_end_date=request.POST.get("work_end_date"),
            work_description=request.POST.get("work_description"),  # From textarea
            work_file=request.FILES.get("work_file"),
            work_status=1
        )

        # Get all selected tasks (checkbox + dynamic entries)
        task_names = request.POST.getlist("task_name")

        for task_name in task_names:
            if task_name.strip():

                # Try to fetch task from Work_Task model
                try:
                    task_obj = Work_Task.objects.get(task_name=task_name)
                    task_description = task_obj.task_description
                except Work_Task.DoesNotExist:
                    task_description = ""  # If manually added task and not in Work_Task

                # Create ClientTask record
                ClientTask_Register.objects.create(
                    cTcompId=comp,
                    client_Id=client,
                    work_Id=work,
                    task_name=task_name,
                    task_description=task_description,  # Saved here now
                    task_status=0,
                    task_create_date=request.POST.get("work_create_date")
                )

        client.work_reg_status = 1
        client.save()

        return JsonResponse({"status": "success"})

def registered_work(request):
    user_id = request.session.get('user_id')
    if not user_id or request.session.get('position') != 'Digital_Marketing_Head':
        return redirect('login')

    user = LogRegister_Details.objects.filter(id=user_id).first()
    employee = EmployeeRegister_Details.objects.filter(login=user).first()

    if not user or not employee:
        messages.error(request, 'Employee details not found')
        return redirect('login')

    comp_id = employee.company.id if employee.company else None

    if request.method == "POST":


        if "collection_head" in request.POST:
            work_id = request.POST.get("work_id")
            work = WorkRegister.objects.get(id=work_id)

            LeadCollection.objects.create(
                work_Id=work,
                collection_head=request.POST.get("collection_head"),
                collection_description=request.POST.get("collection_desc"),
                target=request.POST.get("target"),
                file=request.FILES.get("file")
            )

            return redirect("registered_work")


        else:
            work_id = request.POST.get("work_id")
            task_type = request.POST.get("task_type")
            description = request.POST.get("description")
            file = request.FILES.get("task_file")

            work = WorkRegister.objects.get(id=work_id)
            client = work.clientId
            company = work.wcompId

            task_name = ""

            if task_type == "client":
                task_name = request.POST.get("client_task")

            elif task_type == "company":
                task_id = request.POST.get("task_name")

                if task_id:
                    task_obj = Work_Task.objects.filter(id=task_id).first()
                    if task_obj:
                        task_name = task_obj.task_name

            ClientTask_Register.objects.create(
                cTcompId=company,
                client_Id=client,
                work_Id=work,
                task_name=task_name,
                task_description=description,
                task_file=file,
                task_status=0
            )

            return redirect("registered_work")


    clients = ClientRegister.objects.filter(compId_id=comp_id).order_by('-id')
    tasks = Work_Task.objects.filter(comp_taskid_id=comp_id).order_by('id')

    works = WorkRegister.objects.filter(wcompId_id=comp_id).order_by('-id')
    for work in works:
        work.tasks = ClientTask_Register.objects.filter(work_Id=work)
        work.leads = LeadCollection.objects.filter(work_Id=work)

    context = {
        'head_name': employee.name,
        'employee': employee,
        'user': user,
        'clients': clients,
        'tasks': tasks,
        'works': works,
    }

    return render(request, "registered_work.html", context)

def edit_work(request, id):
    work = get_object_or_404(WorkRegister, id=id)

    if request.method == "POST":
        # Update work description
        work.work_description = request.POST.get("description")

        # Convert date strings to date objects
        start_str = request.POST.get("start_date")
        end_str = request.POST.get("end_date")

        if start_str:
            work.work_create_date = datetime.strptime(start_str, "%Y-%m-%d").date()
        if end_str:
            work.work_end_date = datetime.strptime(end_str, "%Y-%m-%d").date()
        
        # Update file if uploaded
        if request.FILES.get("work_file"):
            work.work_file = request.FILES["work_file"]

        work.save()

        # ðŸ”¹ Get tasks associated with this work
        tasks = list(ClientTask_Register.objects.filter(work_Id=work)
                     .values_list('task_name', flat=True))

        return JsonResponse({
            "status": "success",
            "start_date": work.work_create_date.strftime("%b. %#d, %Y") if work.work_create_date else "",
            "end_date": work.work_end_date.strftime("%b. %#d, %Y") if work.work_end_date else "",
            "tasks": tasks
        })


    return JsonResponse({"status": "failed"}, status=400)

def delete_work(request, work_id):
    if request.method == "POST":  # or "GET" depending on your setup
        work = get_object_or_404(WorkRegister, id=work_id)
        client = work.clientId  # Get the related client

        # Delete the work (this also cascades to ClientTask_Register if set)
        try:
            work.delete()  # This will also delete related ClientTask_Register
            # Update client's work_reg_status
            client.work_reg_status = 0
            client.save()
            return JsonResponse({"status": "success", "message": "Work deleted successfully."})
        except Exception as e:
            return JsonResponse({"status": "error", "message": str(e)}, status=500)

        

        return JsonResponse({"status": "success"})

    return JsonResponse({"status": "failed"}, status=400)
    
def dm_employees_home(request):

    user_id = request.session.get('user_id')

    if not user_id:
        return redirect('login')  # not logged in

    try:
        user = LogRegister_Details.objects.get(id=user_id)
    except LogRegister_Details.DoesNotExist:
        return redirect('login')  # user deleted?

    # Now you can fetch employee details if needed
    employee = None
    if user.position != 'Admin':
        try:
            employee = EmployeeRegister_Details.objects.get(login=user)
        except EmployeeRegister_Details.DoesNotExist:
            employee = None

    return render(request, "dm_employees_home.html", {"user": user, "employee": employee,'head_name': employee.name})

def head_view_employees(request):

    user_id = request.session.get('user_id')
    if not user_id:
        return redirect('login')

    user = LogRegister_Details.objects.get(id=user_id)
    employee = EmployeeRegister_Details.objects.get(login=user)

    business_id = request.session.get('company_id')
    business = BusinessRegister_Details.objects.get(id=business_id)

    #  SORT PARAMS
    sort = request.GET.get("sort", "")
    order = request.GET.get("order", "asc")
    q = request.GET.get('q', '').strip()
    per_page_param = request.GET.get('per_page', '10')   # keep string here
    page_number = request.GET.get('page', 1)

    # Base queryset
    qs = EmployeeRegister_Details.objects.filter(
        company=business
    ).select_related('department', 'designation')

    # Search filter
    if q:
        qs = qs.filter(
            models.Q(name__icontains=q) |
            models.Q(employee_code__icontains=q) |
            models.Q(email__icontains=q) |
            models.Q(department__name__icontains=q) |
            models.Q(designation__name__icontains=q)
        )

    #  APPLY SORTING
    if sort:
        if sort == "department":
            sort_field = "department__name"
        elif sort == "designation":
            sort_field = "designation__name"
        else:
            sort_field = sort

        if order == "desc":
            sort_field = "-" + sort_field

        qs = qs.order_by(sort_field)
    else:
        qs = qs.order_by("name")  # default
    # Handle per_page = ALL
    if per_page_param == "all":
        per_page = qs.count()    # show all rows
    else:
        per_page = int(per_page_param)

    paginator = Paginator(qs, per_page)
    page_obj = paginator.get_page(page_number)

    designations = DesignationRegister_Details.objects.filter(business=business, active_status=True)

    context = {
        'user': user,
        'business': business,
        'employees_page': page_obj,
        'q': q,
        'per_page': per_page_param,     
        'designations': designations,
        'head_name': employee.name,
        'sort': sort,
        'order': order,
    }
    return render(request, 'head_view_employees.html', context)

def head_allocate_employees(request):

    user_id = request.session.get('user_id')
    if not user_id:
        return redirect('login')

    user = LogRegister_Details.objects.get(id=user_id)
    employee = EmployeeRegister_Details.objects.get(login=user)

    # FIX: Get company from session
    business_id = request.session.get("company_id")
    business = BusinessRegister_Details.objects.get(id=business_id)

    q = request.GET.get('q', '').strip()
    per_page = int(request.GET.get('per_page', 10))
    page_number = request.GET.get('page', 1)

    # Already allocated employee IDs
    allocated_ids = Allocation_Details.objects.values_list('allocatEmp_id', flat=True)

    # Unallocated Executives
    unallocated_qs = EmployeeRegister_Details.objects.filter(
        company=business,
        designation__dashboard_id='Executive'
    ).exclude(id__in=allocated_ids)\
     .select_related('department', 'designation')\
     .order_by('name')

    if q:
        unallocated_qs = unallocated_qs.filter(
            models.Q(name__icontains=q) |
            models.Q(employee_code__icontains=q) |
            models.Q(department__name__icontains=q)
        )

    paginator = Paginator(unallocated_qs, per_page)
    page_obj = paginator.get_page(page_number)

    # Team Leads with count
    team_leads = EmployeeRegister_Details.objects.filter(
        company=business,
        designation__dashboard_id='Team_Lead'
    ).order_by('name')

    tl_ids = [tl.id for tl in team_leads]
    counts = Allocation_Details.objects.filter(
        allocat_to_id__in=tl_ids
    ).values('allocat_to_id').annotate(count=models.Count('id'))

    counts_map = {c['allocat_to_id']: c['count'] for c in counts}

    team_leads_with_count = [
        {'tl': tl, 'count': counts_map.get(tl.id)}
        for tl in team_leads
        if counts_map.get(tl.id, 0) > 0
    ]

    context = {
        'user': user,
        'business': business,
        'unallocated_page': page_obj,
        'team_leads': team_leads, 
        'team_leads_with_count': team_leads_with_count,
        'q': q,
        'per_page': per_page,
        'head_name': employee.name,
    }
    return render(request, 'head_allocate_employees.html', context)

@require_POST
def head_api_allocate(request):

    user_id = request.session.get('user_id')
    if not user_id:
        return JsonResponse({'error': 'Unauthorized'}, status=401)

    user = LogRegister_Details.objects.get(id=user_id)

    # FIX: Use session company_id
    business_id = request.session.get("company_id")
    business = BusinessRegister_Details.objects.get(id=business_id)

    # Get employee ids list
    employee_ids = request.POST.getlist('employee_ids[]') or request.POST.get('employee_ids') or ''
    if isinstance(employee_ids, str):
        employee_ids = [s.strip() for s in employee_ids.split(',') if s.strip()]

    target_tl_id = request.POST.get('target_tl_id')
    if not employee_ids or not target_tl_id:
        return JsonResponse({'error': 'Missing parameters'}, status=400)

    try:
        target_tl = EmployeeRegister_Details.objects.get(
            id=target_tl_id,
            company=business,
            designation__dashboard_id='Team_Lead'
        )
    except EmployeeRegister_Details.DoesNotExist:
        return JsonResponse({'error': 'Invalid team lead'}, status=400)

    created = []
    errors = []

    with transaction.atomic():
        for eid in employee_ids:
            try:
                emp = EmployeeRegister_Details.objects.get(
                    id=eid,
                    company=business,
                    designation__dashboard_id='Executive'
                )
            except EmployeeRegister_Details.DoesNotExist:
                errors.append({'id': eid, 'error': 'Invalid executive'})
                continue

            # Skip already allocated
            if Allocation_Details.objects.filter(allocatEmp_id=emp).exists():
                errors.append({'id': eid, 'error': 'Already allocated'})
                continue

            Allocation_Details.objects.create(
                allocatEmp_id=emp,
                allocat_to=target_tl,
                allocate_status=1,
                allocation_date=timezone.now().date()
            )
            created.append(emp.id)

    return JsonResponse({'created': created, 'errors': errors, 'message': 'Allocation processed'})

def head_allocated_list(request):
    """Show allocated employees (company-wide)"""
    user_id = request.session.get('user_id')
    if not user_id:
        return redirect('login')

    user = LogRegister_Details.objects.get(id=user_id)
    employee = EmployeeRegister_Details.objects.get(login=user)

    # FIX: Fetch company using session
    business_id = request.session.get("company_id")
    business = BusinessRegister_Details.objects.get(id=business_id)

    q = request.GET.get('q', '').strip()
    per_page = int(request.GET.get('per_page', 10))
    page_number = request.GET.get('page', 1)

    allocations = Allocation_Details.objects.select_related(
        'allocatEmp_id__department',
        'allocatEmp_id__designation',
        'allocat_to__designation'
    ).filter(
        allocatEmp_id__company=business
    ).order_by('-allocation_date')

    if q:
        allocations = allocations.filter(
            models.Q(allocatEmp_id__name__icontains=q) |
            models.Q(allocat_to__name__icontains=q)
        )

    paginator = Paginator(allocations, per_page)
    page_obj = paginator.get_page(page_number)

    team_leads = EmployeeRegister_Details.objects.filter(
        company=business,
        designation__dashboard_id='Team_Lead'
    ).order_by('name')

    context = {
        'user': user,
        'business': business,
        'allocations_page': page_obj,
        'team_leads': team_leads,
        'q': q,
        'per_page': per_page,
        'head_name': employee.name,
    }
    return render(request, 'head_allocated_list.html', context)

@require_POST
def head_api_reallocate(request):

    user_id = request.session.get('user_id')
    if not user_id:
        return JsonResponse({'error': 'Unauthorized'}, status=401)

    user = LogRegister_Details.objects.get(id=user_id)

    # FIX: Fetch company via session
    business_id = request.session.get("company_id")
    business = BusinessRegister_Details.objects.get(id=business_id)

    employee_ids = request.POST.getlist('employee_ids[]') or request.POST.get('employee_ids') or ''
    if isinstance(employee_ids, str):
        employee_ids = [e.strip() for e in employee_ids.split(',') if e.strip()]

    target_tl_id = request.POST.get('target_tl_id')

    if not employee_ids or not target_tl_id:
        return JsonResponse({'error': 'Missing parameters'}, status=400)

    try:
        target_tl = EmployeeRegister_Details.objects.get(
            id=target_tl_id,
            company=business,
            designation__dashboard_id='Team_Lead'
        )
    except EmployeeRegister_Details.DoesNotExist:
        return JsonResponse({'error': 'Invalid team lead'}, status=400)

    updated = []
    errors = []

    with transaction.atomic():
        for eid in employee_ids:
            try:
                emp = EmployeeRegister_Details.objects.get(id=eid, company=business)
            except EmployeeRegister_Details.DoesNotExist:
                errors.append({'id': eid, 'error': 'Invalid employee'})
                continue

            allocation = Allocation_Details.objects.filter(allocatEmp_id=emp).first()
            if not allocation:
                errors.append({'id': eid, 'error': 'Not allocated yet'})
                continue

            allocation.allocat_to = target_tl
            allocation.allocate_status = 1
            allocation.allocation_date = timezone.now().date()
            allocation.save()

            updated.append(emp.id)

    return JsonResponse({
        'updated': updated,
        'errors': errors,
        'message': 'Re-allocation processed'
    })

@csrf_exempt
def update_employee_designation(request):
    if request.method == 'POST':
        data = json.loads(request.body)
        emp_id = data.get('emp_id')
        designation_name = data.get('designation_name')

        try:
            emp = EmployeeRegister_Details.objects.get(id=emp_id)

            # ✅ Get existing or create new designation automatically
            desig, created = DesignationRegister_Details.objects.get_or_create(
                dashboard_id=designation_name,
                
                defaults={'name': designation_name,}  # optional: set display name same as dashboard_id
            )

            emp.designation = desig
            emp.save()

            return JsonResponse({'success': True})
        except EmployeeRegister_Details.DoesNotExist:
            return JsonResponse({'success': False, 'error': 'Employee not found'})

    return JsonResponse({'success': False, 'error': 'Invalid request'})
    
def head_myschedule(request):
    user_id = request.session.get('user_id')
    user = LogRegister_Details.objects.get(id=user_id)
    employee = EmployeeRegister_Details.objects.get(login=user)
    empschedule = EmployeeSchedule.objects.filter(emp_id=employee)
    return render(request, 'head_myschedule.html', {'emps': empschedule})

@csrf_exempt
def save_schedule(request):
    if request.method == "POST":
        try:
            user_id = request.session.get('user_id')
            user = LogRegister_Details.objects.get(id=user_id)
            employee = EmployeeRegister_Details.objects.get(login=user)
            EmployeeSchedule.objects.create(
                emp_id=employee,
                start_time=request.POST.get("start_time"),
                end_time=request.POST.get("end_time"),
                schedule_head=request.POST.get("schedule_head"),
                todo_content=request.POST.get("todo_content"),
                schedule_date=request.POST.get("schedule_date")
            )
            messages.success(request, "Schedule added successfully!")
            return redirect("head_myschedule")
        except Exception as e:
            messages.error(request, f"Error: {str(e)}")
            return redirect("head_myschedule")

def filter_today_schedule(request):
    if request.method == "GET":
        try:
            user_id = request.session.get('user_id')
            user = LogRegister_Details.objects.get(id=user_id)
            employee = EmployeeRegister_Details.objects.get(login=user)
            today = date.today()
            schedules = EmployeeSchedule.objects.filter(
                emp_id=employee,
                schedule_date=today
            )
            data = []
            for s in schedules:
                data.append({
                    "id": s.id,
                    "schedule_date": s.schedule_date.strftime("%Y-%m-%d"),
                    "start_time": s.start_time.strftime("%H:%M"),
                    "end_time": s.end_time.strftime("%H:%M"),
                    "schedule_head": s.schedule_head,
                    "todo_content": s.todo_content,
                    "schedule_status": s.schedule_status,
                })
            return JsonResponse({"status": "success", "data": data})
        except Exception as e:
            return JsonResponse({"status": "error", "message": str(e)})

def filter_schedule_range(request):
    if request.method == "GET":
        try:
            start = request.GET.get("start_date")
            end = request.GET.get("end_date")
            if not start or not end:
                return JsonResponse({"status": "error", "message": "Missing dates"})
            start_date = datetime.strptime(start, "%Y-%m-%d").date()
            end_date = datetime.strptime(end, "%Y-%m-%d").date()
            user_id = request.session.get('user_id')
            user = LogRegister_Details.objects.get(id=user_id)
            employee = EmployeeRegister_Details.objects.get(login=user)
            schedules = EmployeeSchedule.objects.filter(
                emp_id=employee,
                schedule_date__range=[start_date, end_date]
            ).order_by("schedule_date", "start_time")
            data = []
            for s in schedules:
                data.append({
                    "id": s.id,
                    "schedule_date": s.schedule_date.strftime("%Y-%m-%d"),
                    "start_time": s.start_time.strftime("%H:%M"),
                    "end_time": s.end_time.strftime("%H:%M"),
                    "schedule_head": s.schedule_head,
                    "todo_content": s.todo_content,
                    "schedule_status": s.schedule_status,
                })
            return JsonResponse({"status": "success", "data": data})
        except Exception as e:
            return JsonResponse({"status": "error", "message": str(e)})

@csrf_exempt
def update_schedule_status(request):
    if request.method == "POST":
        try:
            schedule_id = request.POST.get("id")
            checked = request.POST.get("checked")
            status_value = 1 if checked == "true" else 0
            obj = EmployeeSchedule.objects.get(id=schedule_id)
            obj.schedule_status = status_value
            obj.save()
            return JsonResponse({"status": "success", "new_status": status_value})
        except Exception as e:
            return JsonResponse({"status": "error", "message": str(e)})

def get_schedule(request, id):
    schedule = EmployeeSchedule.objects.get(id=id)
    data = {
        "id": schedule.id,
        "schedule_date": schedule.schedule_date,
        "start_time": str(schedule.start_time),
        "end_time": str(schedule.end_time),
        "schedule_head": schedule.schedule_head,
        "todo_content": schedule.todo_content,
        "schedule_status": schedule.schedule_status,
    }
    return JsonResponse({"status": "success", "data": data})

def get_schedule(request, id):
    schedule = EmployeeSchedule.objects.get(id=id)
    data = {
        "id": schedule.id,
        "schedule_date": schedule.schedule_date,
        "start_time": str(schedule.start_time),
        "end_time": str(schedule.end_time),
        "schedule_head": schedule.schedule_head,
        "todo_content": schedule.todo_content,
        "schedule_status": schedule.schedule_status,
    }
    return JsonResponse({"status": "success", "data": data})

@csrf_exempt
def update_schedule(request, id):
    schedule = EmployeeSchedule.objects.get(id=id)

    schedule.schedule_date = request.POST.get("schedule_date")
    schedule.start_time = request.POST.get("start_time")
    schedule.end_time = request.POST.get("end_time")
    schedule.schedule_head = request.POST.get("schedule_head")
    schedule.todo_content = request.POST.get("todo_content")

    schedule.save()
    return JsonResponse({"status": "success"})

def delete_schedule(request, pk):
    if request.method == "POST":
        try:
            obj = EmployeeSchedule.objects.get(id=pk)
            obj.delete()
            return JsonResponse({"status": "success"})
        except EmployeeSchedule.DoesNotExist:
            return JsonResponse({"status": "error", "message": "Not found"})
    return JsonResponse({"status": "invalid"})

def head_schemployee(request):
    user_id = request.session.get('user_id')
    user = LogRegister_Details.objects.get(id=user_id)
    employee = EmployeeRegister_Details.objects.get(login=user)
    company = employee.company
    department = employee.department
    empschedule = EmployeeSchedule.objects.filter(
        emp_id__company=company,
        emp_id__department=department,
        emp_id__designation__dashboard_id='Team_Lead'
    ).exclude(emp_id=employee).order_by('start_time')
    empl = EmployeeRegister_Details.objects.filter(
        company=company,
        department=department,
        verification_status=1,
        designation__dashboard_id='Team_Lead'
    ).exclude(login=user)
    return render(request, 'head_schemployee.html', {
        'emps': empschedule,
        'empl': empl
    })

def save_employee_schedule(request):
    if request.method == "POST":
        emp_id = request.POST.get("schedule_emp")
        start_time = request.POST.get("start_time")
        end_time = request.POST.get("end_time")
        schedule_head = request.POST.get("schedule_head")
        todo_content = request.POST.get("todo_content")
        schedule_date = request.POST.get("schedule_date")
        if not emp_id:
            messages.error(request, "Please select an employee.")
            return redirect('head_schemployee')
        try:
            employee = EmployeeRegister_Details.objects.get(id=emp_id)
            EmployeeSchedule.objects.create(
                emp_id=employee,
                start_time=start_time,
                end_time=end_time,
                schedule_head=schedule_head,
                todo_content=todo_content,
                schedule_date=schedule_date,
                schedule_status=0
            )
            messages.success(request, f"Schedule added successfully for {employee.name}!")
        except EmployeeRegister_Details.DoesNotExist:
            messages.error(request, "Employee not found!")
        return redirect("head_schemployee")
    return redirect("head_schemployee")

def filter_emptoday_schedule(request):
    if request.method == "GET":
        try:
            user_id = request.session.get('user_id')
            user = LogRegister_Details.objects.get(id=user_id)
            employee = EmployeeRegister_Details.objects.get(login=user)
            company = employee.company
            department = employee.department
            today = date.today()
            schedules = EmployeeSchedule.objects.filter(
                schedule_date=today,
                emp_id__company=company,
                emp_id__department=department,
                emp_id__designation__dashboard_id='Team_Lead'
            ).exclude(emp_id=employee).order_by('start_time')
            data = []
            for s in schedules:
                data.append({
                    "id": s.id,
                    "employee_name": s.emp_id.name,
                    "schedule_date": s.schedule_date.strftime("%Y-%m-%d"),
                    "start_time": s.start_time.strftime("%H:%M"),
                    "end_time": s.end_time.strftime("%H:%M"),
                    "schedule_head": s.schedule_head,
                    "todo_content": s.todo_content,
                    "schedule_status": s.schedule_status,
                })
            return JsonResponse({"status": "success", "data": data})
        except Exception as e:
            return JsonResponse({"status": "error", "message": str(e)})

def filter_emp_schedule(request):
    try:
        emp_id = request.GET.get("emp_id")
        if not emp_id:
            return JsonResponse({"status": "error", "message": "Employee ID missing"})
        schedules = EmployeeSchedule.objects.filter(emp_id=emp_id).order_by("schedule_date", "start_time")
        data = []
        for s in schedules:
            data.append({
                "id": s.id,
                "schedule_date": s.schedule_date.strftime("%Y-%m-%d"),
                "start_time": s.start_time.strftime("%H:%M"),
                "end_time": s.end_time.strftime("%H:%M"),
                "schedule_head": s.schedule_head,
                "todo_content": s.todo_content,
                "schedule_status": s.schedule_status,
                "employee_name": s.emp_id.name,
            })
        return JsonResponse({"status": "success", "data": data})
    except Exception as e:
        return JsonResponse({"status": "error", "message": str(e)})

def get_empschedule(request, id):
    sch = EmployeeSchedule.objects.get(id=id)
    return JsonResponse({
        "status": "success",
        "data": {
            "id": sch.id,
            "emp_id": sch.emp_id.id,
            "schedule_date": sch.schedule_date.strftime("%Y-%m-%d"),
            "start_time": sch.start_time.strftime("%H:%M"),
            "end_time": sch.end_time.strftime("%H:%M"),
            "schedule_head": sch.schedule_head,
            "todo_content": sch.todo_content,
        }
    })

def update_empschedule(request, id):
    sch = EmployeeSchedule.objects.get(id=id)
    if request.method == "POST":
        sch.emp_id_id = request.POST.get("emp_id")
        sch.schedule_date = request.POST.get("schedule_date")
        sch.start_time = request.POST.get("start_time")
        sch.end_time = request.POST.get("end_time")
        sch.schedule_head = request.POST.get("schedule_head")
        sch.todo_content = request.POST.get("todo_content")
        sch.save()
        return JsonResponse({"status": "success"})

def head_schefilter(request):
    user_id = request.session.get('user_id')
    user = LogRegister_Details.objects.get(id=user_id)
    employee = EmployeeRegister_Details.objects.get(login=user)
    company = employee.company
    department = employee.department
    today = date.today()
    empschedule = EmployeeSchedule.objects.filter(
        schedule_date=today,
        emp_id__company=company,
        emp_id__department=department,
        emp_id__designation__dashboard_id='Team_Lead'
        ).exclude(emp_id=employee).order_by("schedule_date", "start_time")
    empl = EmployeeRegister_Details.objects.filter(
        company=company,
        department=department,
        verification_status=1,
        designation__dashboard_id='Team_Lead'
        ).exclude(login=user)
    return render(request, 'head_schefilter.html', {'emps': empschedule, 'empl': empl})

def filter_schedules(request):
    try:
        emp_id = request.GET.get("emp_id")
        from_date = request.GET.get("from_date")
        to_date = request.GET.get("to_date")
        user_id = request.session.get('user_id')
        user = LogRegister_Details.objects.get(id=user_id)
        logged_employee = EmployeeRegister_Details.objects.get(login=user)
        schedules = EmployeeSchedule.objects.exclude(emp_id=logged_employee)
        # Filter by employee
        if emp_id:
            schedules = schedules.filter(emp_id_id=emp_id)
        # Filter by date range
        if from_date and to_date:
            schedules = schedules.filter(schedule_date__range=[from_date, to_date])
        schedules = schedules.order_by("schedule_date", "start_time")
        result = []
        for s in schedules:
            result.append({
                "schedule_date": s.schedule_date.strftime("%Y-%m-%d"),
                "start_time": s.start_time.strftime("%H:%M"),
                "end_time": s.end_time.strftime("%H:%M"),
                "schedule_head": s.schedule_head,
                "todo_content": s.todo_content,
                "employee_name": s.emp_id.name,
                "schedule_status": s.schedule_status
            })
        return JsonResponse({"status": "success", "data": result})
    except Exception as e:
        return JsonResponse({"status": "error", "message": str(e)})

def head_myleave(request):
    user_id = request.session.get('user_id')
    user = LogRegister_Details.objects.get(id=user_id)
    employee = EmployeeRegister_Details.objects.get(login=user)
    empleave = EmployeeLeave.objects.filter(emp_id=employee).order_by("start_date")
    return render(request, 'head_myleave.html', {'empl': empleave})

def save_leave(request):
    if request.method == "POST":
        user_id = request.session.get('user_id')
        user = LogRegister_Details.objects.get(id=user_id)
        employee = EmployeeRegister_Details.objects.get(login=user)
        start_date = request.POST.get("leave_from_date")
        end_date = request.POST.get("leave_to_date")
        leave_type = request.POST.get("leave_type")
        leave_reason = request.POST.get("leave_for")
        leave_file = request.FILES.get("leave_file")
        # Convert to date objects
        start = datetime.strptime(start_date, "%Y-%m-%d").date()
        end = datetime.strptime(end_date, "%Y-%m-%d").date()
        # Calculate no of days (inclusive)
        no_of_days = (end - start).days + 1
        leave_entry = EmployeeLeave(
            emp_id = employee,
            start_date = start,
            end_date = end,
            leave_type = leave_type,
            leave_reason = leave_reason,
            no_of_days = no_of_days,
            leave_status = 0,   # pending initially
            leave_apply_date = datetime.today().date(),
            leave_statuChange_date = None,
        )
        # Save file if exists
        if leave_file:
            leave_entry.leave_request_file = leave_file
        leave_entry.save()
        messages.success(request, "Leave Applied Successfully!")
        return redirect("head_myleave")
    return redirect("head_myleave")

def head_leaverequest(request):
    user_id = request.session.get('user_id')
    user = LogRegister_Details.objects.get(id=user_id)
    employee = EmployeeRegister_Details.objects.get(login=user)
    company = employee.company
    empleave = EmployeeLeave.objects.filter(emp_id__company=company).order_by("start_date")
    return render(request, 'head_leaverequest.html', {'empl': empleave})

def update_leave_status(request):
    if request.method == "POST":
        leave_id = request.POST.get("leave_id")
        action = request.POST.get("action")  # approve or reject
        try:
            leave = EmployeeLeave.objects.get(id=leave_id)
            if action == "approve":
                leave.leave_status = 1
            elif action == "reject":
                leave.leave_status = 2
            leave.leave_statuChange_date = timezone.now().date()
            leave.save()
            return JsonResponse({"status": "success"})
        except EmployeeLeave.DoesNotExist:
            return JsonResponse({"status": "error", "message": "Leave not found"})
    return JsonResponse({"status": "error", "message": "Invalid request"})

def action_taken(request):
    if not request.session.get('is_authenticated'):
        return redirect('login')
    employee_id = request.session.get('employee_id')
    company_id = request.session.get('company_id')
    logged_employee = EmployeeRegister_Details.objects.select_related(
        'company', 'department'
    ).get(id=employee_id)

    # queryset: same company + same department, exclude self
    employees = EmployeeRegister_Details.objects.filter(
        company=logged_employee.company,
        department=logged_employee.department
    ).exclude(
        id=logged_employee.id
    ).select_related('department', 'designation')
    actions = ActionTaken.objects.filter(act_emp_id__company_id = company_id)
    today = date.today()

    return render(request, "action_taken.html", {
        "employees": employees,
        "actions": actions,
        "head_name":logged_employee.name,
        "today":today
    })

def save_actionTaken(request):
    if request.method == "POST":
        # ------------------- Logged In User (From Session) -------------------
        employeeId = request.session['employee_id']
        employeeName = request.session['employee_name']

        # ------------------- Employee Selected in Dropdown -------------------
        to_employee_id = request.POST.get("actionEmployee")

        try:
            to_employee = EmployeeRegister_Details.objects.get(id=to_employee_id)
        except EmployeeRegister_Details.DoesNotExist:
            messages.error(request, "Selected employee does not exist.")
            return redirect("action_taken")

        # ------------------- Form Fields -------------------
        action_date = request.POST.get("action_date")
        reason = request.POST.get("action_reason", "")
        reason_for_action = request.POST.get("reason_forAction", "")
        what_action_taken = request.POST.get("what_actionTaken", "")

        # ------------------- Save Record -------------------
        action = ActionTaken.objects.create(
            act_emp_id=to_employee,         
            act_from_id=employeeId,        
            act_from_name=employeeName,     
            act_reason=reason,               
            act_head=reason_for_action,      
            act_content=what_action_taken,   
            action_date=action_date,
            status=0                         
        )

        messages.success(request, "Action submitted successfully.")
        return redirect("action_taken")

    # If GET request:
    return redirect("action_taken")

def edit_action_taken(request,id):
    action = get_object_or_404(ActionTaken, id=id)
    employee_id = request.session.get('employee_id')
    logged_employee = EmployeeRegister_Details.objects.get(id=employee_id)
    employees = EmployeeRegister_Details.objects.filter(
        company=logged_employee.company,department=logged_employee.department
    ).exclude(id=employee_id)
    return render(request,"edit_action_taken.html",{'employees':employees,'action':action})

def edit_actionTaken(request, id):
    action = get_object_or_404(ActionTaken, id=id)
    if request.method == 'POST':
        emp_id = request.POST.get('editActionEmployee')
        employee = get_object_or_404(EmployeeRegister_Details, id=emp_id)
        action.act_emp_id = employee
        action.action_date = request.POST.get('edit_action_date')
        action.act_reason = request.POST.get('edit_action_reason')
        action.act_head = request.POST.get('edit_reason_forAction')
        action.act_content = request.POST.get('edit_what_actionTaken')
        action.save()
        messages.success(request, "Action updated successfully.")
        return redirect('action_taken')

    return render(request, 'edit_actionTaken.html')

def feedback(request):
    company_id = request.session.get("company_id")
    feedback = Feedback.objects.filter(feedback_emp_id__company_id=company_id)
    employee_id = request.session.get('employee_id')
    logged_employee = EmployeeRegister_Details.objects.get(id=employee_id)
    employees = EmployeeRegister_Details.objects.filter(
        company=logged_employee.company,department=logged_employee.department
    ).exclude(id=employee_id)
    return render(request,'feedback.html',{'feedbacks':feedback,'employees':employees,"head_name":logged_employee.name,})

def save_feedback(request):
    if request.method == 'POST':
        toEmployeeId = request.POST.get("feedbackToEmployee")
        toEmployee = EmployeeRegister_Details.objects.get(id=toEmployeeId)
        feedback = request.POST.get("feedback_text")
        fromEmployee_id = request.session.get('employee_id')
        fromEmployee_name = request.session.get('employee_name')
        fb = Feedback.objects.create(
            feedback_emp_id = toEmployee,
            from_id = fromEmployee_id,
            from_name = fromEmployee_name,
            feedback_content = feedback,
            feedback_date =  date.today()
            )
        messages.success(request, "Feedback Added successfully.")
        return redirect('feedback')
    return render(request, 'feedback.html')

def head_view_leaves(request):
    employeeId = request.session.get('employee_id')
    employeeName = request.session.get('employee_name')
    company_id = request.session.get('company_id')
    leaves = EmployeeLeave.objects.filter( emp_id__company_id=company_id).select_related('emp_id')
    employees = EmployeeRegister_Details.objects.filter(company_id=company_id).exclude(id=employeeId)
     # ---------------- FILTER PARAMETERS ----------------
    emp_filter = request.GET.get('employee')
    from_date = request.GET.get('from_date')
    to_date = request.GET.get('to_date')

    # ---------------- APPLY FILTERS ----------------
    if emp_filter:
        leaves = leaves.filter(emp_id_id=emp_filter)

    if from_date and to_date:
        leaves = leaves.filter(
            start_date__gte=from_date,
            end_date__lte=to_date
        )
    elif from_date:
        leaves = leaves.filter(start_date__gte=from_date)
    elif to_date:
        leaves = leaves.filter(end_date__lte=to_date)

    context = {
        'leaves': leaves,
        'employees': employees,
        'head_name': employeeName,
        'selected_employee': emp_filter,
        'from_date': from_date,
        'to_date': to_date,
    }
    return render(request,'head_view_leaves.html',context)

def head_view_actionTaken(request):
    company_id = request.session.get('company_id')
    actions = ActionTaken.objects.filter(act_emp_id__company_id = company_id)
    employeeName = request.session.get('employee_name')
    employeeId = request.session.get('employee_id')
    employees = EmployeeRegister_Details.objects.filter(company_id=company_id).exclude(id=employeeId)
     # ---------------- FILTER PARAMETERS ----------------
    emp_filter = request.GET.get('employee')
    from_date = request.GET.get('from_date')
    to_date = request.GET.get('to_date')

    # ---------------- APPLY FILTERS ----------------
    if emp_filter:
        actions = actions.filter(act_emp_id=emp_filter)

    if from_date and to_date:
        actions = actions.filter(
            action_date__gte=from_date,
            action_date__lte=to_date
        )
    elif from_date:
        actions = actions.filter(action_date__gte=from_date)
    elif to_date:
        actions = actions.filter(action_date__lte=to_date)

    return render(request,'head_view_actionTaken.html',{'actions':actions,'head_name':employeeName,'employees':employees})

def admin_employees(request):
    user_id = request.session.get('user_id')

    try:
        user = LogRegister_Details.objects.get(id=user_id)
        business = BusinessRegister_Details.objects.get(login=user)

        context = {
            'user': user,
            'business': business,
            'owner_full_name': f"{business.owner_fname} {business.owner_lname}",
            'company_name': business.company_name,
        }

        return render(request, "admin_employees.html", context)

    except (LogRegister_Details.DoesNotExist, BusinessRegister_Details.DoesNotExist):
        messages.error(request, "User or business details not found")
        return redirect("login")

def employee_views(request):
    user_id = request.session.get('user_id')

    try:
        user = LogRegister_Details.objects.get(id=user_id)
        business = BusinessRegister_Details.objects.get(login=user)

        employees = EmployeeRegister_Details.objects.filter(
    company=business,
    active_status='Approved'
).order_by('-created_at')

        context = {
            'user': user,
            'business': business,
            'owner_full_name': f"{business.owner_fname} {business.owner_lname}",
            'company_name': business.company_name,
            'employees': employees,
        }
        
        return render(request, "employee_views.html", context)

    except (LogRegister_Details.DoesNotExist, BusinessRegister_Details.DoesNotExist):
        messages.error(request, "User or business details not found")
        return redirect("login")

def edit_employee(request):
    if request.method == "POST":
        emp_id = request.POST.get("employee_id")

        user_id = request.session.get("user_id")
        if not user_id:
            messages.error(request, "Please login first")
            return redirect("login")
        user = LogRegister_Details.objects.get(id=user_id)
        company = BusinessRegister_Details.objects.get(login=user)

        employee = get_object_or_404(EmployeeRegister_Details, pk=emp_id, company=company)

        # Update address details
        employee.address_line1 = request.POST.get("address_line1", employee.address_line1)
        employee.address_line2 = request.POST.get("address_line2", employee.address_line2)
        employee.address_line3 = request.POST.get("address_line3", employee.address_line3)
        employee.pin = request.POST.get("pin", employee.pin)
        employee.location = request.POST.get("location", employee.location)
        employee.district = request.POST.get("district", employee.district)
        employee.state = request.POST.get("state", employee.state)

        employee.verification_status = False
        employee.save()
        messages.success(request, f"{employee.name}'s details updated successfully")
        return redirect("employee_views")
    
def resigned(request):
    user_id = request.session.get('user_id')

    if not user_id:
        messages.error(request, "Please login first")
        return redirect("login")

    try:
        # Logged-in user
        user = LogRegister_Details.objects.get(id=user_id)

        # Business linked to this admin
        business = BusinessRegister_Details.objects.get(login=user)

        # ONLY resigned employees under THIS company
        resigned_employees = EmployeeRegister_Details.objects.filter(
            company_id=business.id,     # 🔒 strict company filter
            active_status="Resigned"
        ).order_by('-created_at')

        context = {
            'user': user,
            'business': business,
            'company_name': business.company_name,
            'owner_full_name': f"{business.owner_fname} {business.owner_lname}",
            'resigned_employees': resigned_employees,
        }

        return render(request, "resigned.html", context)

    except LogRegister_Details.DoesNotExist:
        messages.error(request, "User not found")
        return redirect("login")

    except BusinessRegister_Details.DoesNotExist:
        messages.error(request, "Business not found")
        return redirect("login")

def update_employee(request):
    if request.method == "POST":
        emp_id = request.POST.get("employee_id")
        active_status = request.POST.get("active_status")

        user_id = request.session.get("user_id")
        if not user_id:
            messages.error(request, "Please login first")
            return redirect("login")
        user = LogRegister_Details.objects.get(id=user_id)
        company = BusinessRegister_Details.objects.get(login=user)

        employee = get_object_or_404(EmployeeRegister_Details, pk=emp_id, company=company)
        employee.active_status = active_status
        employee.save()
        messages.success(request, f"{employee.name} status changed to {active_status}")
        return redirect("employee_views")

def re_approve(request):
    if request.method == "POST":
        employee_id = request.POST.get("employee_id")

        user_id = request.session.get("user_id")
        user = LogRegister_Details.objects.get(id=user_id)
        company = BusinessRegister_Details.objects.get(login=user)

        employee = get_object_or_404(
            EmployeeRegister_Details,
            id=employee_id,
            company=company
        )

        employee.active_status = "Approved"
        employee.verification_status = True
        employee.save()

        messages.success(request, "Employee re-approved successfully")
        return redirect("resigned")

def allocated_lists(request):
    user_id = request.session.get('user_id')

    if not user_id:
        messages.error(request, "Please login first")
        return redirect("login")

    try:
        # Logged-in user
        user = LogRegister_Details.objects.get(id=user_id)

        # Business linked to user
        business = BusinessRegister_Details.objects.get(login=user)

        # Selected Team Lead from dropdown
        team_lead_id = request.GET.get('team_lead_id')

        # Base queryset (company-wide allocations)
        allocations = Allocation_Details.objects.select_related(
            'allocatEmp_id__designation',
            'allocat_to__designation'
        ).filter(
            allocatEmp_id__company=business
        )

        # Filter allocations when team lead selected
        if team_lead_id:
            allocations = allocations.filter(
                allocatEmp_id_id=team_lead_id
            )

        allocations = allocations.order_by('-allocation_date')

        # Team Leads for dropdown
        team_leads = EmployeeRegister_Details.objects.filter(
            company=business,
            designation__dashboard_id='Team_Lead'
        ).order_by('name')

        context = {
            'user': user,
            'business': business,
            'allocations': allocations,
            'team_leads': team_leads,
            'selected_team_lead': team_lead_id,
            'owner_full_name': f"{business.owner_fname} {business.owner_lname}",
            'company_name': business.company_name,
        }

        return render(request, "allocated_lists.html", context)

    except LogRegister_Details.DoesNotExist:
        messages.error(request, "User not found")
        return redirect("login")

    except BusinessRegister_Details.DoesNotExist:
        messages.error(request, "Business details not found")
        return redirect("login")

def employee_leaves(request):
    user_id = request.session.get('user_id')

    if not user_id:
        messages.error(request, "Please login first")
        return redirect("login")

    user = LogRegister_Details.objects.get(id=user_id)
    business = BusinessRegister_Details.objects.get(login=user)

    # Employees for dropdown
    employees = EmployeeRegister_Details.objects.filter(
        company=business
    ).order_by('name')

    # GET values
    employee_id = request.GET.get('employee_id')
    from_date = request.GET.get('from_date')
    to_date = request.GET.get('to_date')

    leaves = EmployeeLeave.objects.none()

    if employee_id:
        leaves = EmployeeLeave.objects.filter(
            emp_id_id=employee_id
        )

        # ✅ FROM DATE filter (independent)
        if from_date:
            leaves = leaves.filter(start_date__gte=from_date)

        # ✅ TO DATE filter (independent)
        if to_date:
            leaves = leaves.filter(end_date__lte=to_date)

        leaves = leaves.order_by('-leave_apply_date')

    return render(request, 'employee_leaves.html', {
        'user': user,
        'business': business,
        'employees': employees,
        'leaves': leaves,
        'selected_employee': employee_id,
        'owner_full_name': f"{business.owner_fname} {business.owner_lname}",
        'company_name': business.company_name,
    })

def employee_schedule(request):
    user_id = request.session.get('user_id')

    if not user_id:
        messages.error(request, "Please login first")
        return redirect("login")

    user = LogRegister_Details.objects.get(id=user_id)
    business = BusinessRegister_Details.objects.get(login=user)

    # Employees for dropdown
    employees = EmployeeRegister_Details.objects.filter(
        company=business
    ).order_by('name')

    # GET values
    employee_id = request.GET.get('employee_id')
    from_date = request.GET.get('from_date')
    to_date = request.GET.get('to_date')

    schedules = EmployeeSchedule.objects.none()

    if employee_id:
        schedules = EmployeeSchedule.objects.filter(
            emp_id_id=employee_id
        )

        # ✅ FROM DATE filter (independent)
        if from_date:
            schedules = schedules.filter(schedule_date__gte=from_date)

        # ✅ TO DATE filter (independent)
        if to_date:
            schedules = schedules.filter(schedule_date__lte=to_date)

        schedules = schedules.order_by('-schedule_date')

    return render(request, 'employee_schedule.html', {
        'user': user,
        'business': business,
        'employees': employees,
        'schedules': schedules,
        'selected_employee': employee_id,
        'owner_full_name': f"{business.owner_fname} {business.owner_lname}",
        'company_name': business.company_name,
    })
    
def tl_schedule(request):

    user_id = request.session.get('user_id')

    if not user_id:
        messages.error(request, "Please login first")
        return redirect("login")

    try:
        user = LogRegister_Details.objects.get(id=user_id)
        employee = EmployeeRegister_Details.objects.get(login=user)

        # =========================
        # CREATE / UPDATE SCHEDULE
        # =========================
        if request.method == "POST":

            start_time = request.POST.get("start_time")
            end_time = request.POST.get("end_time")
            schedule_id = request.POST.get("schedule_id")

            # ---- TIME VALIDATION ----
            if start_time >= end_time:
                messages.error(request, "End time must be greater than start time")
                return redirect("tl_schedule")

            # ---------- UPDATE ----------
            if schedule_id:
                try:
                    schedule = EmployeeSchedule.objects.get(
                        id=schedule_id,
                        emp_id=employee
                    )

                    schedule.start_time = start_time
                    schedule.end_time = end_time
                    schedule.schedule_head = request.POST.get("schedule_head")
                    schedule.todo_content = request.POST.get("todo_content")
                    schedule.schedule_date = request.POST.get("schedule_date") or date.today()
                    schedule.save()

                    messages.success(request, "Schedule updated successfully")

                except EmployeeSchedule.DoesNotExist:
                    messages.error(request, "Schedule not found")

            # ---------- CREATE ----------
            else:
                EmployeeSchedule.objects.create(
                    emp_id=employee,
                    start_time=start_time,
                    end_time=end_time,
                    schedule_head=request.POST.get("schedule_head"),
                    todo_content=request.POST.get("todo_content"),
                    schedule_date=request.POST.get("schedule_date") or date.today()
                )

                messages.success(request, "Schedule added successfully")

            return redirect("tl_schedule")

        # =========================
        # FILTER SECTION
        # =========================
        from_date = request.GET.get("from_date")
        to_date = request.GET.get("to_date")

        if from_date:
            from_date = datetime.strptime(from_date, "%Y-%m-%d").date()

        if to_date:
            to_date = datetime.strptime(to_date, "%Y-%m-%d").date()

        schedules = EmployeeSchedule.objects.filter(emp_id=employee)

        # ---- DATE FILTER LOGIC ----
        if from_date and to_date:

            if from_date > to_date:
                messages.error(request, "From date cannot be greater than To date")
                return redirect("tl_schedule")

            schedules = schedules.filter(schedule_date__range=[from_date, to_date])

        elif from_date:
            schedules = schedules.filter(schedule_date=from_date)

        elif to_date:
            schedules = schedules.filter(schedule_date=to_date)

        else:
            # DEFAULT → TODAY ONLY
            schedules = schedules.filter(schedule_date=timezone.localdate())

        schedules = schedules.order_by("schedule_date", "start_time")

        # =========================
        # RENDER PAGE
        # =========================
        return render(request, "tl_schedule.html", {
            "teamlead_name": employee.name,
            "employee": employee,
            "user": user,
            "schedules": schedules
        })

    except (LogRegister_Details.DoesNotExist, EmployeeRegister_Details.DoesNotExist):
        messages.error(request, "Employee details not found")
        return redirect("login") 
    
def toggle_schedule(request,id):
    task = EmployeeSchedule.objects.get(id=id)
    task.schedule_status = 0 if task.schedule_status else 1
    task.save()
    return JsonResponse({"status":"ok"})

# tl_delete schedules
def delete_tlschedule(request,id):
    EmployeeSchedule.objects.get(id=id).delete()
    messages.success(request, "Schedule deleted successfully")
    return redirect("tl_schedule")

# tl_schedules employees
def tl_empschedule(request):

    user_id = request.session.get("user_id")

    if not user_id:
        messages.error(request, "Please login first")
        return redirect("login")

    try:
        user = LogRegister_Details.objects.get(id=user_id)
        tl_employee = EmployeeRegister_Details.objects.get(login=user)

        # =========================
        # GET EMPLOYEES UNDER TL
        # =========================
        allocated_employees = EmployeeRegister_Details.objects.filter(
            id__in=Allocation_Details.objects.filter(
                allocat_to=tl_employee,
                allocate_status=1
            ).values_list("allocatEmp_id", flat=True)
        )

        # =========================
        # CREATE / UPDATE SCHEDULE
        # =========================
        if request.method == "POST":

            start_time = request.POST.get("start_time")
            end_time = request.POST.get("end_time")
            schedule_id = request.POST.get("schedule_id")
            emp_id = request.POST.get("employee")

            # validate time
            if start_time >= end_time:
                messages.error(request, "End time must be greater than start time")
                return redirect("tl_emp_schedule")

            # validate employee belongs to TL
            if not allocated_employees.filter(id=emp_id).exists():
                messages.error(request, "Invalid employee selected")
                return redirect("tl_emp_schedule")

            # UPDATE
            if schedule_id:
                try:
                    schedule = EmployeeSchedule.objects.get(id=schedule_id,
                        emp_id__in=allocated_employees   )

                    schedule.emp_id_id = emp_id
                    schedule.start_time = start_time
                    schedule.end_time = end_time
                    schedule.schedule_head = request.POST.get("schedule_head")
                    schedule.todo_content = request.POST.get("todo_content")
                    schedule.schedule_date = request.POST.get("schedule_date") or date.today()
                    schedule.save()

                    messages.success(request, "Schedule updated successfully")

                except EmployeeSchedule.DoesNotExist:
                    messages.error(request, "Schedule not found")

            # CREATE
            else:
                EmployeeSchedule.objects.create(
                    emp_id_id=emp_id,
                    start_time=start_time,
                    end_time=end_time,
                    schedule_head=request.POST.get("schedule_head"),
                    todo_content=request.POST.get("todo_content"),
                    schedule_date=request.POST.get("schedule_date") or date.today()
                )

                messages.success(request, "Schedule added successfully")

            return redirect("tl_empschedule")

        # =========================
        # FILTER SECTION
        # =========================
        from_date = request.GET.get("from_date")
        to_date = request.GET.get("to_date")

        if from_date:
            from_date = datetime.strptime(from_date, "%Y-%m-%d").date()

        if to_date:
            to_date = datetime.strptime(to_date, "%Y-%m-%d").date()

        # only show schedules of employees under TL
        # 
        emp_filter = request.GET.get("emp")

        schedules = EmployeeSchedule.objects.filter(
            emp_id__in=allocated_employees
        )

        if emp_filter:
            schedules = schedules.filter(emp_id=emp_filter)

        if from_date and to_date:
            if from_date > to_date:
                messages.error(request, "From date cannot be greater than To date")
                return redirect("tl_empschedule")

            schedules = schedules.filter(schedule_date__range=[from_date, to_date])

        elif from_date:
            schedules = schedules.filter(schedule_date=from_date)

        elif to_date:
            schedules = schedules.filter(schedule_date=to_date)

        else:
            schedules = schedules.filter(schedule_date=timezone.localdate())

        schedules = schedules.order_by("schedule_date", "start_time")

        return render(request, "tl_empschedule.html", {
            "teamlead_name": tl_employee.name,
            "employee": tl_employee,
            "user": user,
            "schedules": schedules,
            "employees": allocated_employees
        })

    except (LogRegister_Details.DoesNotExist, EmployeeRegister_Details.DoesNotExist):
        messages.error(request, "Employee details not found")
        return redirect("login")

# tl_deletes employees schedules
def delete_empschedule(request, id):
    schedule = get_object_or_404(EmployeeSchedule, id=id)
    schedule.delete()
    messages.success(request, "Schedule deleted successfully")
    return redirect('tl_empschedule')


def tl_leaves(request):

    user_id = request.session.get("user_id")

    # ---------- LOGIN CHECK ----------
    if not user_id:
        messages.error(request, "Please login first")
        return redirect("login")

    try:
        user = LogRegister_Details.objects.get(id=user_id)
        employee = EmployeeRegister_Details.objects.get(login=user)

    except LogRegister_Details.DoesNotExist:
        messages.error(request, "User not found")
        return redirect("login")

    except EmployeeRegister_Details.DoesNotExist:
        messages.error(request, "Employee profile not found")
        return redirect("login")

    # =====================================================
    # SAVE LEAVE REQUEST
    # =====================================================
    if request.method == "POST":

        from_date = request.POST.get("from_date")
        to_date = request.POST.get("to_date")
        leave_type = request.POST.get("leave_type")
        reason = request.POST.get("reason")
        attachment = request.FILES.get("attachment")

        try:
            start = datetime.strptime(from_date, "%Y-%m-%d").date()
            end = datetime.strptime(to_date, "%Y-%m-%d").date()
        except:
            messages.error(request, "Invalid date format")
            return redirect("tl_leaves")

        # ---------- VALIDATIONS ----------

        # 1. Past date block
        if start < date.today() or end < date.today():
            messages.error(request, "Past dates are not allowed")
            return redirect("tl_leaves")

        # 2. End < Start
        if end < start:
            messages.error(request, "To date must be after From date")
            return redirect("tl_leaves")

        # 3. Half day only for same day
        if (end - start).days > 0 and leave_type == "Half Day":
            messages.error(request, "Half day leave allowed only for single day")
            return redirect("tl_leaves")

        # 4. Overlapping leave check
        exists = EmployeeLeave.objects.filter(
            emp_id=employee,
            start_date__lte=end,
            end_date__gte=start
        ).exists()

        if exists:
            messages.error(request, "You already applied leave for these dates")
            return redirect("tl_leaves")

        # ---------- CALCULATE DAYS ----------
        total_days = (end - start).days + 1

        # ---------- SAVE ----------
        EmployeeLeave.objects.create(
            emp_id=employee,
            start_date=start,
            end_date=end,
            leave_type=leave_type,
            leave_reason=reason,
            no_of_days=total_days,
            leave_status=0,
            leave_apply_date=date.today(),
            leave_request_file=attachment
        )

        messages.success(request, "Leave request submitted successfully")
        return redirect("tl_leaves")

    # =====================================================
    # GET LEAVE LIST + FILTER
    # =====================================================
    my_leaves = EmployeeLeave.objects.filter(emp_id=employee)

    start_filter = request.GET.get("start")
    end_filter = request.GET.get("end")

    if start_filter and end_filter:
        my_leaves = my_leaves.filter(
            start_date__gte=start_filter,
            end_date__lte=end_filter
        )

    my_leaves = my_leaves.order_by("-id")
    # =====================================================
    # EMPLOYEES LEAVES SECTION
    # =====================================================

    assigned_ids = Allocation_Details.objects.filter(
        allocat_to=employee
    ).values_list("allocatEmp_id", flat=True)

    employees = EmployeeRegister_Details.objects.filter(
        id__in=assigned_ids
    )

    employees_leaves = EmployeeLeave.objects.filter(
        emp_id__in=employees
    )

    # filters
    emp_filter = request.GET.get("employee")
    emp_start = request.GET.get("emp_start")
    emp_end = request.GET.get("emp_end")

    if emp_filter:
        employees_leaves = employees_leaves.filter(emp_id_id=emp_filter)

    if emp_start and emp_end:
        employees_leaves = employees_leaves.filter(
            start_date__gte=emp_start,
            end_date__lte=emp_end
        )

    employees_leaves = employees_leaves.order_by("-id")

    # ---------- CONTEXT ----------
    context = {
        "user": user,
        "employee": employee,
        "teamlead_name": employee.name,
        "my_leaves": my_leaves,
        "employees": employees,
        "employees_leaves": employees_leaves
    }

    return render(request, "tl_leaves.html", context)


def tl_action(request):

    user_id = request.session.get("user_id")

    if not user_id:
        messages.error(request,"Login required")
        return redirect("login")

    user = LogRegister_Details.objects.get(id=user_id)
    employee = EmployeeRegister_Details.objects.get(login=user)


    # =========================
    # EMPLOYEES UNDER TEAMLEAD
    # =========================
    assigned_ids = Allocation_Details.objects.filter(
        allocat_to=employee
    ).values_list("allocatEmp_id", flat=True)

    employees = EmployeeRegister_Details.objects.filter(id__in=assigned_ids)

    edit_id = request.GET.get("edit")
    edit_action = None

    if edit_id:
        edit_action = ActionTaken.objects.filter(id=edit_id).first()
    # =========================
    # SAVE ACTION
    # =========================
    if request.method == "POST":

        emp_id = request.POST.get("employee")
        action_date = request.POST.get("action_date")
        reason_title = request.POST.get("reason_title")
        reason_action = request.POST.get("reason_action")
        action_taken = request.POST.get("action_taken")

        # ---------- Validation ----------
        if not emp_id:
            messages.error(request,"Please select employee")
            return redirect("tl_action")

        try:
            emp = EmployeeRegister_Details.objects.get(id=emp_id)
        except:
            messages.error(request,"Invalid employee")
            return redirect("tl_action")

        # date validation
        try:
            action_date_obj = datetime.strptime(action_date,"%Y-%m-%d").date()
        except:
            messages.error(request,"Invalid date")
            return redirect("tl_action")

        if action_date_obj > date.today():
            messages.error(request,"Future date not allowed")
            return redirect("tl_action")


        # ---------- SAVE ----------
        edit_id = request.POST.get("edit_id")

        if edit_id:
            action = ActionTaken.objects.get(id=edit_id)
        else:
            action = ActionTaken()

        action.act_emp_id = emp
        action.act_from_id = employee.id
        action.act_from_name = employee.name
        action.act_head = reason_title
        action.act_reason = reason_action
        action.act_content = action_taken
        action.action_date = action_date_obj
        action.status = 0
        action.save()

        messages.success(request,"Action recorded successfully")
        return redirect("tl_action")

    # =========================
    # ACTIONS TAKEN BY TEAMLEAD
    # =========================
    actions_taken = ActionTaken.objects.filter(
        act_from_id=employee.id
    ).select_related("act_emp_id").order_by("-id")

    actions_received = ActionTaken.objects.filter(
        act_emp_id=employee
    ).select_related("act_emp_id").order_by("-id")
    # =========================
    # CONTEXT
    # =========================
    context = {
        "employees": employees,
        "teamlead_name": employee.name,
        "actions_taken": actions_taken,
        "edit_action": edit_action,
        "actions_received": actions_received
    }

    return render(request,"tl_action.html",context)
    
##################### Executive ###################

def executive_schedule(request): 
    user_id = request.session.get('user_id')

    if not user_id:
        messages.error(request, "Please login first")
        return redirect("login")

    try:
        user = LogRegister_Details.objects.get(id=user_id)
        employee = EmployeeRegister_Details.objects.get(login=user)

        current_day=date.today()
        day=current_day.strftime("%A")
        c_day=current_day.strftime('%d %B, %Y')
        

        schedules_today=EmployeeSchedule.objects.filter(schedule_date=current_day,emp_id=employee)
        return render(request,'executive_schedule.html',{'day':day,'date':c_day,'schedules':schedules_today,'executive_name':employee.name})
    
    except (LogRegister_Details.DoesNotExist, EmployeeRegister_Details.DoesNotExist):
        messages.error(request, "Employee details not found")
        return redirect("login") 



def register_executive_schedule(request):

    user_id = request.session.get('user_id')

    if not user_id:
        messages.error(request, "Please login first")
        return redirect("login")

    try:
        user = LogRegister_Details.objects.get(id=user_id)
        employee = EmployeeRegister_Details.objects.get(login=user)
        if request.method == "POST":
            schedule_date=request.POST.get('date')
            start_time=request.POST.get('start_time')
            end_time=request.POST.get('end_time')
            head=request.POST.get('task_head')
            description=request.POST.get('task')

            ######### update  

            schedule_id=request.POST.get('schedule_id')
            if schedule_id:
                try:
                    schedule=EmployeeSchedule.objects.get(id=schedule_id)
                    schedule.schedule_date=schedule_date
                    schedule.start_time=start_time
                    schedule.end_time=end_time
                    schedule.schedule_head=head
                    schedule.todo_content=description
                    schedule.save()
                    messages.success(request,'Schedule updated successfully')
                    return redirect('executive_schedule')
                except EmployeeSchedule.DoesNotExist:
                    messages.error(request,'Schedule not found')
                    return redirect('executive_schedule')

            else:

            
            
                EmployeeSchedule.objects.create(emp_id=employee,schedule_date=schedule_date,start_time=start_time,end_time=end_time,schedule_head=head,todo_content=description)
                messages.success(request,'Schedule added successfully')
                return redirect('executive_schedule')
    except (LogRegister_Details.DoesNotExist, EmployeeRegister_Details.DoesNotExist):
        messages.error(request, "Employee details not found")
        return redirect("login") 
        

def exclusive_schedule_finished(request):
    if request.method == 'POST':
        schedule_id = request.POST.get('id')
        is_checked =int( request.POST.get('checked_or_not'))


        if not schedule_id:
            messages.error(request,'Schedule not found')
            return redirect('executive_schedule')


        updated = EmployeeSchedule.objects.filter(id=schedule_id).update(
            schedule_status=1 if is_checked==1 else 0
        )

        if not updated:
            messages.error(request,'Schedule not found')
            return redirect('executive_schedule')
        messages.success(request,'Schedule Completed') if is_checked==1 else messages.error(request,'Schedule not completed')

        return redirect('executive_schedule')



def exclusive_schedule_edit(request):
    schedule_id = request.GET.get('id')
    print(schedule_id)
    schedule = get_object_or_404(EmployeeSchedule, id=schedule_id)
    format_date = schedule.schedule_date.strftime('%Y-%m-%d')
    data={'id':schedule_id,'start_time':schedule.start_time,'end_time':schedule.end_time,'schedule_head':schedule.schedule_head,'todo_content':schedule.todo_content,'date':format_date}
    return JsonResponse(data)

def delete_executive_schedule(request,id):
    try:
        data=EmployeeSchedule.objects.get(id=id)
        data.delete()
        messages.success(request,'Schedule deleted successfully')
        return redirect('executive_schedule')
    except EmployeeSchedule.DoesNotExist:
        messages.error(request,'Schedule not found')
        return redirect('executive_schedule')


def executive_view_schedule(request):
    user_id = request.session.get('user_id')

    if not user_id:
        messages.error(request, "Please login first")
        return redirect("login")

    try:
        user = LogRegister_Details.objects.get(id=user_id)
        employee = EmployeeRegister_Details.objects.get(login=user)

        alldata=EmployeeSchedule.objects.filter(emp_id=employee).order_by('-schedule_date')
        from_date = request.POST.get("from_date", "")
        to_date = request.POST.get("to_date", "")
        searchtype = request.POST.get('searchtype', 'all')

        if request.method == "POST":
            try:
                from_date_obj = datetime.strptime(from_date, "%Y-%m-%d").date() if from_date else None
                to_date_obj = datetime.strptime(to_date, "%Y-%m-%d").date() if to_date else None
            except ValueError:
                    messages.error(request, 'Invalid search dates')
                    return redirect('executive_view_schedule')

            if from_date and to_date and from_date_obj > to_date_obj:
                    messages.error(request, 'To date must be greater than or equal to From date')
                    return redirect('executive_view_schedule')
            if from_date and to_date:
                
                alldata = alldata.filter(schedule_date__range=[from_date, to_date])
            elif from_date:
                alldata = alldata.filter(schedule_date__gte=from_date)
            elif to_date:
                alldata = alldata.filter(schedule_date__lte=to_date)

            if searchtype == 'completed':
                alldata = alldata.filter(schedule_status=1)
            elif searchtype == 'upcoming':
                alldata = alldata.filter(schedule_status=0, schedule_date__gte=date.today())

        return render(
            request,'executive_view_schedule.html',
            {'tasks': alldata,'from_date': from_date,'to_date': to_date,'searchtype': searchtype,'user':employee,'executive_name':employee.name,})
    
    except (LogRegister_Details.DoesNotExist, EmployeeRegister_Details.DoesNotExist):
        messages.error(request, "Employee details not found")
        return redirect("login")
    

def executive_leave(request):
    user_id = request.session.get('user_id')

    if not user_id:
        messages.error(request, "Please login first")
        return redirect("login")

    try:
        user = LogRegister_Details.objects.get(id=user_id)
        employee = EmployeeRegister_Details.objects.get(login=user)

        show_view_leave = request.session.pop('show_view_leave', False)

        all_leave=EmployeeLeave.objects.filter(emp_id=employee)


        if request.method == "POST":
            from_date = request.POST.get("from_date")
            to_date = request.POST.get("to_date")
            search= request.POST.get("search")
            leave_type=request.POST.get("leave_type")
            reason=request.POST.get("reason")
            file=request.FILES.get("file")
            

            if search=='no':
                try:
                    start_date = datetime.strptime(from_date, "%Y-%m-%d").date()
                    end_date = datetime.strptime(to_date, "%Y-%m-%d").date()
                except (TypeError, ValueError):
                    messages.error(request, 'Invalid leave dates')
                    return redirect('executive_leave')

                if end_date < start_date:
                    messages.error(request, 'To date must be greater than or equal to From date')
                    return redirect('executive_leave')

                no_of_days = (end_date - start_date).days + 1
                print(no_of_days)

                EmployeeLeave.objects.create(
                    emp_id=employee,
                    start_date=start_date,
                    end_date=end_date,
                    leave_type=leave_type,
                    leave_reason=reason,
                    no_of_days=no_of_days,
                    leave_request_file=file
                )
                messages.success(request,'Leave applied successfully')
                request.session['show_view_leave'] = True
                return redirect('executive_leave')
            elif search == 'yes':
                show_view_leave = True
                try:
                    from_date_obj = datetime.strptime(from_date, "%Y-%m-%d").date() if from_date else None
                    to_date_obj = datetime.strptime(to_date, "%Y-%m-%d").date() if to_date else None
                except ValueError:
                    messages.error(request, 'Invalid search dates')
                    request.session['show_view_leave'] = True
                    return redirect('executive_leave')

                if from_date_obj and to_date_obj and to_date_obj < from_date_obj:
                    messages.error(request, 'To date must be greater than or equal to From date')
                    request.session['show_view_leave'] = True
                    return redirect('executive_leave')

                if from_date_obj and to_date_obj:
                    all_leave = all_leave.filter(start_date__gte=from_date_obj, end_date__lte=to_date_obj)
                    
                elif from_date_obj:
                    all_leave = all_leave.filter(start_date__gte=from_date_obj)
                    
                elif to_date_obj:
                    all_leave = all_leave.filter(end_date__lte=to_date_obj)
                    
                else:
                    messages.error(request, 'Invalid search dates')
                
        return render(
            request,
            "executive_leave.html",
            {'leave': all_leave, 'show_view_leave': show_view_leave,'executive_name':employee.name}
        )
    
    except (LogRegister_Details.DoesNotExist, EmployeeRegister_Details.DoesNotExist):
        messages.error(request, "Employee details not found")
        return redirect("login")
    

def executive_feedback(request):
    user_id = request.session.get('user_id')

    if not user_id:
        messages.error(request, "Please login first")
        return redirect("login")

    try:
        user = LogRegister_Details.objects.get(id=user_id)
        employee = EmployeeRegister_Details.objects.get(login=user)

        department=DepartmentRegister_Details.objects.get(id=employee.department.id)
        company=BusinessRegister_Details.objects.get(id=employee.company.id)
        users=EmployeeRegister_Details.objects.filter(department_id=department,company_id=company).exclude(id=employee.id)

        given_feedback=Feedback.objects.filter(from_id=employee.id)
        recived_feedback=Feedback.objects.filter(feedback_emp_id=employee)
        show_view_feedback = request.session.pop('show_view_feedback', False)



        if request.method == "POST":
            feedback=request.POST.get('feedback')
            emp=request.POST.get('emp')
            
            feedback_emp=EmployeeRegister_Details.objects.get(id=emp)
            
            Feedback.objects.create(feedback_emp_id=feedback_emp,feedback_content=feedback,from_id=employee.id,from_name=employee.name,feedback_date=date.today())
            messages.success(request,'Feedback submitted successfully')
            request.session['show_view_feedback'] = True
            return redirect('executive_feedback')

        return render(request,'executive_feedback.html',{'users':users,'given_feedback':given_feedback,'recived_feedback':recived_feedback,'executive_name':employee.name,'view_feedback':show_view_feedback})

    except (LogRegister_Details.DoesNotExist, EmployeeRegister_Details.DoesNotExist):
        messages.error(request, "Employee details not found")
        return redirect("login")


def executive_complaints(request):
    user_id = request.session.get('user_id')

    if not user_id:
        messages.error(request, "Please login first")
        return redirect("login")

    try:
        user = LogRegister_Details.objects.get(id=user_id)
        employee = EmployeeRegister_Details.objects.get(login=user)


        show_view_complaint = request.session.pop('show_view_complaint', False)


        if request.method == "POST":
            complaint=request.POST.get('complaint')
            complaint_date=date.today()
            complaint_head=request.POST.get('complaint_head')
            complaint_status=0
            Complaints.objects.create(complaint_emp_id=employee,
                                    compaint_content=complaint,
                                    complaint_date=complaint_date,
                                    compaint_head=complaint_head,
                                    status=complaint_status)
            show_view_complaint
            
            messages.success(request,'Complaint submitted successfully')
            request.session['show_view_complaint'] = True
            return redirect('executive_complaints')
        
        complaint_details=Complaints.objects.filter(complaint_emp_id=employee)

        
        return render(request,'executive_complaints.html',{'executive_name':employee.name,'complaint_details':complaint_details,'show_view_complaint':show_view_complaint})
    
    except (LogRegister_Details.DoesNotExist, EmployeeRegister_Details.DoesNotExist):
        messages.error(request, "Employee details not found")
        return redirect("login")
#End

def telecaller_leave(request):
    user_id = request.session.get('user_id')  
    if not user_id:
        return redirect('login') 

    emp = EmployeeRegister_Details.objects.filter(login_id=user_id).first()
    if not emp:
        return redirect('login')

    leaves = EmployeeLeave.objects.filter(emp_id=emp)

    from_date = request.GET.get("from_date")
    to_date = request.GET.get("to_date")
    if from_date and to_date:
        leaves = leaves.filter(start_date__gte=from_date, end_date__lte=to_date)

    search = request.GET.get("search")
    if search:
        leaves = leaves.filter(
            Q(leave_reason__icontains=search) |
            Q(leave_type__icontains=search)
        )

    if request.method == "POST":
        start_date = request.POST.get("start_date")
        end_date = request.POST.get("end_date")
        leave_type = request.POST.get("leave_type")
        reason = request.POST.get("leave_reason")
        file = request.FILES.get("leave_request_file")

        sd = datetime.strptime(start_date, "%Y-%m-%d")
        ed = datetime.strptime(end_date, "%Y-%m-%d")
        no_days = (ed - sd).days + 1

        EmployeeLeave.objects.create(
            emp_id=emp,
            start_date=start_date,
            end_date=end_date,
            leave_type=leave_type,
            leave_reason=reason,
            no_of_days=no_days,
            leave_status=0,
            leave_apply_date=start_date,
            leave_request_file=file
        )

        return redirect('telecaller_leave')

    return render(request, "telecaller_leave.html", {"leaves": leaves})






def telecaller_action_taken(request):

    employee_id = request.session.get("employee_id")

    actions = ActionTaken.objects.filter(
        act_emp_id_id = employee_id
    ).order_by('-action_date')

    action_count = actions.count()

    return render(request, "telecaller_action_taken.html", {
        "actions": actions,
        "action_count": action_count
    })





def telecaller_feedback(request):

    employee_id = request.session.get('employee_id')
    if not employee_id:
        return redirect('login')

    logged_employee = EmployeeRegister_Details.objects.get(id=employee_id)


    if request.method == "POST":

        to_emp = request.POST.get("feedbackToEmployee")
        text = request.POST.get("feedback_text")

        if to_emp and text:
            Feedback.objects.create(
                feedback_emp_id_id = to_emp,
                from_id = employee_id,
                from_name = logged_employee.name,
                feedback_content = text,
                feedback_date = date.today()
            )

        return redirect("telecaller_feedback")


    received_feedback = Feedback.objects.filter(feedback_emp_id_id = employee_id).order_by('-feedback_date')

    sent_feedback = Feedback.objects.filter(from_id = employee_id).order_by('-feedback_date')

    feedbacks = sorted(
        chain(received_feedback, sent_feedback),
        key=lambda x: x.feedback_date or date.today(),
        reverse=True
    )

    employees = EmployeeRegister_Details.objects.filter(
        company = logged_employee.company,
        department = logged_employee.department
    ).exclude(id=employee_id).order_by('name')


    context = {
        "received_feedback": received_feedback,
        "sent_feedback": sent_feedback,
        "feedbacks": feedbacks,
        "employees": employees,
        "head_name": logged_employee.name,
    }

    return render(request, 'telecaller_feedback.html', context)


def telecaller_complaints(request):

    employee_id = request.session.get('employee_id')

    if not employee_id:
        return redirect('login')

    employee = EmployeeRegister_Details.objects.filter(id=employee_id).first()

    if not employee:
        return redirect('login')


    complaints = Complaints.objects.filter(
        complaint_emp_id__company=employee.company
    ).order_by('-id')


    telecallers = EmployeeRegister_Details.objects.filter(
        company=employee.company,
        department=employee.department
    ).exclude(id=employee_id).order_by('name')


    if request.method == "POST":

        emp_id = request.POST.get("employee")
        complaint = request.POST.get("complaint")

        Complaints.objects.create(
            complaint_emp_id_id=emp_id,
            compaint_content=complaint,
            status=0
        )

        return redirect('telecaller_complaints')


    return render(request,'telecaller_complaints.html',{
        'complaints':complaints,
        'telecallers':telecallers
    })

def employee_views_tele(request):
    log_id = request.session.get('user_id')  

    if not log_id:
        return redirect('login')

    logged_user = LogRegister_Details.objects.filter(id=log_id).first()
    if not logged_user:
        return HttpResponse("Login user not found")

    logged_employee = EmployeeRegister_Details.objects.filter(
        login=logged_user
    ).select_related('company').first()

    if not logged_employee:
        return HttpResponse("Employee record not found")

    employees = EmployeeRegister_Details.objects.select_related(
        'department', 'designation', 'company'
    ).filter(
        company=logged_employee.company,
        active_status='Approved',
        designation__dashboard_id='Telecaller' 
    ).order_by('-created_at')

    return render(request, "data_manager_telle.html", {
        'employees': employees,'datamanager_name': logged_employee.name
    })

def employee_views_executive(request):
    log_id = request.session.get('user_id')   

    if not log_id:
        return redirect('login')

    logged_user = LogRegister_Details.objects.filter(id=log_id).first()
    if not logged_user:
        return HttpResponse("Login user not found")

    logged_employee = EmployeeRegister_Details.objects.filter(
        login=logged_user
    ).select_related('company').first()

    if not logged_employee:
        return HttpResponse("Employee record not found")

    employees = EmployeeRegister_Details.objects.select_related(
        'department', 'designation', 'company'
    ).filter(
        company=logged_employee.company,
        active_status='Approved',
        designation__dashboard_id__in=['Team_Lead', 'Executive'] 
    ).order_by('-created_at')

    return render(request, "data_manager_executive.html", {
        'employees': employees,'datamanager_name': logged_employee.name
    })


def employee_detail(request, id):
    log_id = request.session.get('user_id')   
    logged_user = LogRegister_Details.objects.filter(id=log_id).first()
    if not logged_user:
        return HttpResponse("Login user not found")

    logged_employee = EmployeeRegister_Details.objects.filter(
        login=logged_user
    ).select_related('company').first()
    employee = EmployeeRegister_Details.objects.select_related(
        'department', 'designation', 'company', 'login'
    ).get(id=id)

    return render(request, "employee_details.html", {
        'employee': employee,'datamanager_name': logged_employee.name
    })


def view_employee_schedules(request):

    user_id = request.session.get('user_id')
    company_id = request.session.get('company_id')
    department_id = request.session.get('department_id')
    if not user_id or not company_id:
        return redirect('login')
    user = LogRegister_Details.objects.filter(id=user_id).first()
    if not user:
        return redirect('login')

    logged_employee = EmployeeRegister_Details.objects.filter(
        login=user
    ).select_related('company').first()

    if not logged_employee:
        return HttpResponse("Employee record not found")

    if not company_id:
        return redirect('login')
    employees = EmployeeRegister_Details.objects.filter(
        company_id=company_id,
        
        active_status='Approved'
    )
    if department_id:
        employees = employees.filter(department_id=department_id)

    schedules = EmployeeSchedule.objects.none()

    emp_id = request.GET.get('employee')
    from_date = request.GET.get('from_date')
    to_date = request.GET.get('to_date')

    if emp_id:
        schedules = EmployeeSchedule.objects.select_related('emp_id').filter(
            emp_id_id=emp_id,
            emp_id__company_id=company_id,
            )
        if from_date:
            schedules = schedules.filter(schedule_date__gte=from_date)

        if to_date:
            schedules = schedules.filter(schedule_date__lte=to_date)

        schedules = schedules.order_by('-schedule_date', '-start_time')

    return render(request, 'head_view_shedule.html', {
        'employees': employees,
        'emps': schedules,
        'head_name': logged_employee.name 
    })

def head_view_resigned(request):
    user_id = request.session.get('user_id')
    company_id = request.session.get('company_id')
    department_id = request.session.get('department_id')
    if not user_id or not company_id:
        return redirect('login')
    user = LogRegister_Details.objects.filter(id=user_id).first()
    if not user:
        return redirect('login')
    logged_employee = EmployeeRegister_Details.objects.filter(
        login=user
    ).select_related('company').first()

    if not logged_employee:
        return HttpResponse("Employee record not found")
    resigned_employees = EmployeeRegister_Details.objects.select_related(
        'department', 'designation'
    ).filter(
        company_id=company_id,
        active_status='Resigned'
    )
    if department_id:
        resigned_employees = resigned_employees.filter(
            department_id=department_id
        )
    resigned_employees = resigned_employees.order_by('-created_at')

    return render(request, 'head_view_resigned.html', {
        'resigned_employees': resigned_employees,
        'head_name': logged_employee.name 
    })

from django.db.models import Q

def view_employee_feedback(request):
    log_id = request.session.get('user_id')

    if not log_id:
        return redirect('login')

    logged_user = LogRegister_Details.objects.filter(id=log_id).first()
    if not logged_user:
        return HttpResponse("Login user not found")

    logged_employee = EmployeeRegister_Details.objects.filter(
        login=logged_user
    ).select_related('company').first()

    if not logged_employee:
        return HttpResponse("Employee record not found")

    # Employees in same company
    employees = EmployeeRegister_Details.objects.filter(
        company=logged_employee.company
    )

    company_emp_ids = employees.values_list('id', flat=True)
    feedback = Feedback.objects.none()

    emp_id = request.GET.get('employee')
    f_type = request.GET.get('type')
    if emp_id or f_type:

        feedback = Feedback.objects.filter(
            Q(feedback_emp_id__in=company_emp_ids) |
            Q(from_id__in=company_emp_ids)
        )
        if emp_id:
            feedback = feedback.filter(
                Q(feedback_emp_id_id=emp_id) |
                Q(from_id=emp_id)
            )
        if f_type == "given":
            feedback = feedback.filter(from_id=emp_id)

        elif f_type == "recived":
            feedback = feedback.filter(feedback_emp_id=emp_id)

        elif f_type == "all":
            feedback = feedback.filter(
                Q(from_id=emp_id) |
                Q(feedback_emp_id=emp_id)
            )
            

    return render(request, "head_view_feedback.html", {
        'employees': employees,
        'feedback': feedback.order_by('-feedback_date'),
        'head_name': logged_employee.name
    })
    
def dm_work_allocate(request):

    user_id = request.session.get('user_id')
    user = LogRegister_Details.objects.filter(id=user_id).first()
    employee = EmployeeRegister_Details.objects.filter(login=user).first()

    comp_id = employee.company.id if employee and employee.company else None   

    works = WorkRegister.objects.filter(
        wcompId_id=comp_id
    ).select_related('clientId').prefetch_related('allocated_emp')

    for work in works:
        work.is_allocated = work.allocated_emp.exists()  

    return render(request, "dm_work_allocate.html", {
        "works": works,
        "head_name": employee.name if employee else ""
    })

def remove_allocation(request, work_id, emp_id):

    user_id = request.session.get('user_id')
    user = LogRegister_Details.objects.filter(id=user_id).first()
    employee = EmployeeRegister_Details.objects.filter(login=user).first()

    comp_id = employee.company.id if employee and employee.company else None

    work = get_object_or_404(
        WorkRegister,
        id=work_id,
        wcompId_id=comp_id
    )

    emp = get_object_or_404(EmployeeRegister_Details, id=emp_id)

    work.allocated_emp.remove(emp)

    LeadAllocation.objects.filter(
        work=work,
        team_lead=emp
    ).delete()

    if work.allocated_emp.count() == 0:
        work.work_allocate_status = 0
        work.work_status = 0
        work.work_progress = 0
        work.save()

    return redirect('dm_work_allocate')




def work_allocate_page(request, id):

    user_id = request.session.get('user_id')
    user = LogRegister_Details.objects.filter(id=user_id).first()
    employee = EmployeeRegister_Details.objects.filter(login=user).first()
    comp_id = employee.company.id if employee and employee.company else None

    work = get_object_or_404(
        WorkRegister,
        id=id,
        wcompId_id=comp_id
    )

    tasks = ClientTask_Register.objects.filter(work_Id=work)

    task_list = [
        {
            "id": t.id,
            "task_name": t.task_name,
            "is_lead_collection": "lead collection" in t.task_name.lower()
        }
        for t in tasks
    ]

    lead_collections = LeadCollection.objects.filter(work_Id=work)


    lead_collection_list = [
        {"id": lead.id, "collection_head": lead.collection_head}
        for lead in lead_collections
    ]
    lead_collection_json = json.dumps(lead_collection_list, cls=DjangoJSONEncoder)

    team_leads = EmployeeRegister_Details.objects.filter(
        company_id=comp_id,
        designation__dashboard_id='Team_Lead'
    )

    if request.method == "POST":
        emp_id = request.POST.get("team_lead")
        task_id = request.POST.get("task")
        work_type = request.POST.get("type", "single")
        lead_collection_id = request.POST.get("lead_collection")

        if emp_id and task_id:
            emp = EmployeeRegister_Details.objects.get(id=emp_id)
            task = ClientTask_Register.objects.get(id=task_id)
            work.allocated_emp.add(emp)

            lead_collection = LeadCollection.objects.get(id=lead_collection_id) if lead_collection_id else None

            LeadAllocation.objects.create(
                team_lead=emp,
                work=work,
                task=task,
                lead_collection=lead_collection,
                display_name=lead_collection.collection_head if lead_collection else "",
                description=request.POST.get("description"),
                instagram=request.POST.get("instagram_link"),
                facebook=request.POST.get("facebook_link"),
                file=request.FILES.get("file"),
                target=request.POST.get("target") or 0,
                work_type=work_type,
                start_date=request.POST.get("start_date") or None,
                end_date=request.POST.get("end_date") or None,
            )

            work.work_status = 1
            work.work_allocate_status = 1
            work.save()

            return redirect('dm_work_allocate')

    return render(request, "dm_allocate.html", {
        "work": work,
        "team_leads": team_leads,
        "task_list": task_list,
        "lead_collection_list": lead_collection_json,
        "head_name": employee.name if employee else ""
    })

def edit_task(request):
    if request.method == "POST":

        user_id = request.session.get('user_id')
        user = LogRegister_Details.objects.filter(id=user_id).first()
        employee = EmployeeRegister_Details.objects.filter(login=user).first()
        comp_id = employee.company.id if employee and employee.company else None

        task_id = request.POST.get("task_id")

        try:
            task = ClientTask_Register.objects.get(
                id=task_id,
                cTcompId_id=comp_id
            )

            task.task_name = request.POST.get("new_task")
            task.task_description = request.POST.get("description")

            if request.FILES.get("task_file"):
                task.task_file = request.FILES.get("task_file")

            task.save()

            return JsonResponse({"status": "success"})

        except ClientTask_Register.DoesNotExist:
            return JsonResponse({"status": "error", "message": "Task not found"})

def get_task_details(request):

    user_id = request.session.get('user_id')
    user = LogRegister_Details.objects.filter(id=user_id).first()
    employee = EmployeeRegister_Details.objects.filter(login=user).first()
    comp_id = employee.company.id if employee and employee.company else None

    task_id = request.GET.get("task_id")

    try:
        task = ClientTask_Register.objects.get(
            id=task_id,
            cTcompId_id=comp_id
        )

        return JsonResponse({
            "description": task.task_description,
            "file": task.task_file.url if task.task_file else ""
        })

    except ClientTask_Register.DoesNotExist:
        return JsonResponse({"error": "Task not found"}, status=404)


def delete_client_task(request):

    if request.method == "POST":

        user_id = request.session.get('user_id')
        user = LogRegister_Details.objects.filter(id=user_id).first()
        employee = EmployeeRegister_Details.objects.filter(login=user).first()
        comp_id = employee.company.id if employee and employee.company else None

        task_id = request.POST.get("task_id")

        try:
            task = ClientTask_Register.objects.get(
                id=task_id,
                cTcompId_id=comp_id
            )

            task.delete()

            return JsonResponse({
                "status": "success",
                "message": "Task deleted"
            })

        except ClientTask_Register.DoesNotExist:
            return JsonResponse({
                "status": "error",
                "message": "Task not found"
            })
        





def delete_task_from_allocation(request, allocation_id):

    if request.method == "POST":

        user_id = request.session.get('user_id')
        user = LogRegister_Details.objects.filter(id=user_id).first()
        employee = EmployeeRegister_Details.objects.filter(login=user).first()
        comp_id = employee.company.id if employee and employee.company else None

        allocation = get_object_or_404(
            LeadAllocation,
            id=allocation_id,
            work__wcompId_id=comp_id
        )

        if allocation.task:
            allocation.task = None
            allocation.save()

            return JsonResponse({
                "status": "success",
                "message": "Task removed from allocation"
            })
        else:
            return JsonResponse({
                "status": "error",
                "message": "No task assigned"
            })

    return JsonResponse({
        "status": "error",
        "message": "Invalid request"
    })

def edit_lead(request):

    if request.method == "POST":

        user_id = request.session.get('user_id')
        user = LogRegister_Details.objects.filter(id=user_id).first()
        employee = EmployeeRegister_Details.objects.filter(login=user).first()
        comp_id = employee.company.id if employee and employee.company else None

        lead_id = request.POST.get("lead_id")

        try:
            lead = LeadCollection.objects.get(
                id=lead_id,
                work_Id__wcompId_id=comp_id
            )

            lead.collection_head = request.POST.get("collection_head")
            lead.collection_description = request.POST.get("collection_description")
            lead.target = request.POST.get("target")

            if request.FILES.get("file"):
                lead.file = request.FILES.get("file")

            lead.save()

            return JsonResponse({"status": "success"})

        except LeadCollection.DoesNotExist:
            return JsonResponse({
                "status": "error",
                "message": "Lead not found"
            })
    


def delete_lead(request):

    if request.method == "POST":

        user_id = request.session.get('user_id')
        user = LogRegister_Details.objects.filter(id=user_id).first()
        employee = EmployeeRegister_Details.objects.filter(login=user).first()
        comp_id = employee.company.id if employee and employee.company else None

        lead_id = request.POST.get("lead_id")

        try:
            lead = LeadCollection.objects.get(
                id=lead_id,
                work_Id__wcompId_id=comp_id
            )

            lead.delete()

            return JsonResponse({
                "status": "success",
                "message": "Lead deleted successfully"
            })

        except LeadCollection.DoesNotExist:
            return JsonResponse({
                "status": "error",
                "message": "Lead not found"
            })

    return JsonResponse({
        "status": "error",
        "message": "Invalid request"
    })




def delete_category_allocation(request, id):

    user_id = request.session.get('user_id')
    user = LogRegister_Details.objects.filter(id=user_id).first()
    employee = EmployeeRegister_Details.objects.filter(login=user).first()
    comp_id = employee.company.id if employee and employee.company else None

    try:
        alloc = LeadAllocation.objects.get(
            id=id,
            work__wcompId_id=comp_id
        )

        lead_id = alloc.team_lead.id
        work_id = alloc.work.id

        alloc.lead_collection = None
        alloc.display_name = None
        alloc.target = 0
        alloc.save()

    except LeadAllocation.DoesNotExist:
        return redirect('dm_work_allocate')  

    return redirect('team_lead_tasks', lead_id=lead_id, work_id=work_id)

def team_lead_tasks(request, lead_id, work_id):

    user_id = request.session.get('user_id')
    user = LogRegister_Details.objects.filter(id=user_id).first()
    employee = EmployeeRegister_Details.objects.filter(login=user).first()
    comp_id = employee.company.id if employee and employee.company else None

    lead = get_object_or_404(
        EmployeeRegister_Details,
        id=lead_id,
        company_id=comp_id
    )

    allocations = LeadAllocation.objects.filter(
        team_lead=lead,
        work_id=work_id,
        work__wcompId_id=comp_id
    ).select_related(
        "work",
        "lead_collection",
        "task"
    ).order_by("-allocated_date")

    return render(request, "dm_tl_tasks.html", {
        "lead": lead,
        "allocations": allocations,
        "head_name": employee.name if employee else ""
    })


def delete_work_allocation(request, allocation_id):

    user_id = request.session.get('user_id')
    user = LogRegister_Details.objects.filter(id=user_id).first()
    employee = EmployeeRegister_Details.objects.filter(login=user).first()
    comp_id = employee.company.id if employee and employee.company else None

    allocation = get_object_or_404(
        LeadAllocation,
        id=allocation_id,
        work__wcompId_id=comp_id
    )

    work = allocation.work
    team_lead = allocation.team_lead

    lead_id = team_lead.id
    work_id = work.id

    allocation.delete()

    still_exists = LeadAllocation.objects.filter(
        work=work,
        team_lead=team_lead
    ).exists()

    if not still_exists:
        work.allocated_emp.remove(team_lead)

    if work.allocated_emp.count() == 0:
        work.work_allocate_status = 0
        work.work_status = 0
        work.work_progress = 0
        work.save()

    return redirect(reverse('team_lead_tasks', args=[lead_id, work_id]))

def delete_task_allocation(request, allocation_id):

    user_id = request.session.get('user_id')
    user = LogRegister_Details.objects.filter(id=user_id).first()
    employee = EmployeeRegister_Details.objects.filter(login=user).first()
    comp_id = employee.company.id if employee and employee.company else None

    allocation = get_object_or_404(
        LeadAllocation,
        id=allocation_id,
        work__wcompId_id=comp_id
    )

    lead_id = allocation.team_lead.id
    work_id = allocation.work.id

    allocation.task = None
    allocation.save()

    return redirect(reverse('team_lead_tasks', args=[lead_id, work_id]))



def edit_lead_collection(request, allocation_id):

    user_id = request.session.get('user_id')
    user = LogRegister_Details.objects.filter(id=user_id).first()
    employee = EmployeeRegister_Details.objects.filter(login=user).first()
    comp_id = employee.company.id if employee and employee.company else None

    allocation = get_object_or_404(
        LeadAllocation,
        id=allocation_id,
        work__wcompId_id=comp_id
    )

    if request.method == "POST":
        allocation.display_name = request.POST.get("collection_head")
        allocation.target = request.POST.get("target")
        allocation.description = request.POST.get("description")
        allocation.instagram = request.POST.get("instagram")
        allocation.facebook = request.POST.get("facebook")

        if request.FILES.get("file"):
            allocation.file = request.FILES.get("file")

        allocation.start_date = request.POST.get("start_date")
        allocation.end_date = request.POST.get("end_date")

        allocation.save()

    return redirect(reverse(
        'team_lead_tasks',
        args=[allocation.team_lead.id, allocation.work.id]
    ))

def delete_lead_collection(request, allocation_id):

    user_id = request.session.get('user_id')
    user = LogRegister_Details.objects.filter(id=user_id).first()
    employee = EmployeeRegister_Details.objects.filter(login=user).first()
    comp_id = employee.company.id if employee and employee.company else None

    allocation = get_object_or_404(
        LeadAllocation,
        id=allocation_id,
        work__wcompId_id=comp_id
    )

    allocation.lead_collection = None
    allocation.display_name = None
    allocation.save()

    return redirect(reverse(
        'team_lead_tasks',
        args=[allocation.team_lead.id, allocation.work.id]
    ))


def pending_works(request):

    user_id = request.session.get('user_id')
    user = LogRegister_Details.objects.filter(id=user_id).first()
    employee = EmployeeRegister_Details.objects.filter(login=user).first()
    comp_id = employee.company.id if employee and employee.company else None

    allocations = LeadAllocation.objects.filter(
        work__wcompId_id=comp_id
    ).select_related(
        'work', 'team_lead', 'task', 'lead_collection', 'work__clientId'
    ).prefetch_related(
        'work__allocated_emp'
    ).order_by('-allocated_date')

    works = []
    for alloc in allocations:
        work = alloc.work

        executives = (
            work.allocated_emp.exclude(id=alloc.team_lead.id)
            if alloc.team_lead else work.allocated_emp.all()
        )

        exec_names = ", ".join([
            emp.name
            for emp in executives
            if emp.designation.dashboard_id == "Executive"
        ]) or "N/A"

        works.append({
            "work": work,
            "team_lead": alloc.team_lead,
            "task": alloc.task,
            "allocated_date": alloc.allocated_date,
            "work_type": alloc.work_type,
            "executive_names": exec_names,
            "end_date": work.work_end_date,
        })

    return render(request, 'dm_pending_works.html', {'works': works, "head_name": employee.name if employee else ""})


def work_progress(request):

    user_id = request.session.get('user_id')
    user = LogRegister_Details.objects.filter(id=user_id).first()
    employee = EmployeeRegister_Details.objects.filter(login=user).first()
    comp_id = employee.company.id if employee and employee.company else None

    allocated_work = WorkRegister.objects.filter(
        clientId=OuterRef('pk'),
        work_allocate_status=1
    )

    clients = ClientRegister.objects.filter(
        compId_id=comp_id
    ).annotate(
        has_work=Exists(allocated_work)
    )

    return render(request, "dm_work_progress.html", {
        "clients": clients,
         "head_name": employee.name if employee else ""
    })
def client_work_details(request, client_id):

    user_id = request.session.get('user_id')
    user = LogRegister_Details.objects.filter(id=user_id).first()
    employee = EmployeeRegister_Details.objects.filter(login=user).first()
    comp_id = employee.company.id if employee and employee.company else None

    client = get_object_or_404(
        ClientRegister,
        id=client_id,
        compId_id=comp_id
    )

    allocations = LeadAllocation.objects.filter(
        work__clientId=client,
        work__wcompId_id=comp_id
    )

    is_allocated = allocations.exists()

    if is_allocated:
        work = allocations.first().work
        tasks = ClientTask_Register.objects.filter(
            id__in=LeadAllocation.objects.filter(
                work__clientId=client,
                work__wcompId_id=comp_id
            ).values_list('task_id', flat=True)
        )
    else:
        work = None
        tasks = []

    progress = 0
    if is_allocated:
        works = WorkRegister.objects.filter(
            clientId=client,
            wcompId_id=comp_id
        )
        progress = (
            sum([w.work_progress for w in works]) // works.count()
            if works.exists() else 0
        )

    return render(request, "dm_client_details.html", {
        "client": client,
        "work": work,
        "tasks": tasks,
        "progress": progress,
        "is_allocated": is_allocated,
         "head_name": employee.name if employee else ""
    })
    
def lead_fields_page(request, lead_id):
    user_id = request.session.get('user_id')
    if not user_id:
        return redirect('login')

    user = LogRegister_Details.objects.filter(id=user_id).first()
    employee = EmployeeRegister_Details.objects.filter(login=user).first()

    if not user or not employee:
        messages.error(request, 'Employee details not found')
        return redirect('login')

    selected_lead = LeadCollection.objects.get(id=lead_id)


    leads = LeadCollection.objects.filter(
        work_Id__wcompId=employee.company
    ).order_by('-id')

    data = []
    for lead in leads:
        achieved = LeadAllocation.objects.filter(
            lead_collection=lead
        ).count()

        target = lead.target or 0
        progress = int((achieved / target) * 100) if target > 0 else 0

        data.append({
            "lead": lead,
            "achieved": achieved,
            "progress": progress
        })

    return render(request, "dm_lead_fields_page.html", {
        "data": data,
        "head_name": employee.name,
        "department": selected_lead.collection_head 
    })

def add_required_field(request):
    if request.method == "POST":
        lead_id = request.POST.get("lead_id")
        name = request.POST.get("field_name")
        desc = request.POST.get("field_description")

        lead = LeadCollection.objects.get(id=lead_id)  

        LeadField.objects.create(
            lead=lead,  
            field_name=name,
            field_description=desc
        )

    return redirect('lead_fields_page', lead_id=lead_id)



def lead_data_page(request):
    user_id = request.session.get('user_id')
    if not user_id:
        return redirect('login')

    user = LogRegister_Details.objects.filter(id=user_id).first()
    employee = EmployeeRegister_Details.objects.filter(login=user).first()
    if not user or not employee:
        messages.error(request, 'Employee details not found')
        return redirect('login')

    leads = LeadCollection.objects.filter(work_Id__wcompId=employee.company).order_by('-id')

    data = {}

    for lead in leads:
        client_name = lead.work_Id.clientId.client_name

        if client_name not in data:
            data[client_name] = {
                "start_date": lead.work_Id.work_create_date,
                "end_date": lead.work_Id.work_end_date,
                "category_fields": {}  
            }

        category = lead
        if category not in data[client_name]["category_fields"]:
            data[client_name]["category_fields"][category] = []

        for f in lead.fields.all():
            data[client_name]["category_fields"][category].append({
                "field_name": f.field_name
            })

    return render(request, "lead_data_page.html", {
        "data": data,
        "head_name": employee.name
    })




def download_lead_format(request, lead_id):

    lead = LeadCollection.objects.get(id=lead_id)


    client_name = lead.work_Id.clientId.client_name
    category_name = lead.collection_head


    file_name = f"{client_name}_{category_name}.xlsx"
    file_name = file_name.replace(" ", "_").replace("/", "_")

    dynamic_fields = lead.fields.all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Lead Format"

 
    headers = [
        "Full Name",
        "Email Id",
        "Contact Number",
        "Lead Source"
    ]


    for f in dynamic_fields:
        headers.append(f.field_name)


    ws.append(headers)


    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    response["Content-Disposition"] = f'attachment; filename="{file_name}"'
    response["Cache-Control"] = "no-cache"
    response["Pragma"] = "no-cache"
    response["Expires"] = "0"

    wb.save(response)

    return response




def upload_leads_excel(request, lead_id):

    if request.method == "POST":

        file = request.FILES.get("file")

        wb = openpyxl.load_workbook(file)
        sheet = wb.active

        lead_collection = LeadCollection.objects.get(id=lead_id)

        dynamic_fields = list(
            LeadField.objects.filter(lead=lead_collection)
        )

        user_id = request.session.get('user_id')
        user = LogRegister_Details.objects.get(id=user_id)
        employee = EmployeeRegister_Details.objects.get(login=user)

        for row in sheet.iter_rows(min_row=2, values_only=True):

            lead_row = LeadRow.objects.create(
                lead_collection=lead_collection,
                full_name=row[0],
                email=row[1],
                contact=row[2],
                source=row[3],
                collected_by=employee,
                status="Unverified"
            )

            for i, field in enumerate(dynamic_fields):
                value = row[4 + i] if len(row) > (4 + i) else ""
                LeadRowValue.objects.create(
                lead_row=lead_row,
                field=field,
                value=str(value) if value else "")
    return redirect("lead_data_page")

def client_lead_data(request, lead_id):

    lead_collection = get_object_or_404(LeadCollection, id=lead_id)
    fields = list(lead_collection.fields.all())


    user_id = request.session.get('user_id')
    if not user_id:
        messages.error(request, "Please login first.")
        return redirect("login")

    user = LogRegister_Details.objects.filter(id=user_id).first()
    employee = EmployeeRegister_Details.objects.filter(login=user).first()


    platforms = Platform.objects.filter(company=employee.company)

    if not employee:
        messages.error(request, "Employee details not found.")
        return redirect("login")


    if request.method == "POST":
        form_type = request.POST.get("form_type")
        if form_type == "single":
            source_name = request.POST.get("lead_source", "").strip()
            
            if source_name:
                Platform.objects.get_or_create(
                    company=employee.company, 
                    name=source_name
                )
            row = LeadRow.objects.create(
                lead_collection=lead_collection,
                full_name=request.POST.get("full_name"),
                email=request.POST.get("email"),
                contact=request.POST.get("contact"),
                source=request.POST.get("lead_source"),
                collected_by=employee,
                status="Unverified"
            )
            for field in fields:
                value = request.POST.get(f"field_{field.id}")
                if value:
                    LeadRowValue.objects.create(
                        lead_row=row,
                        field=field,
                        value=value
                    )
            return redirect("client_lead_data", lead_id=lead_id)

        elif form_type == "excel":
            file = request.FILES.get("excel_file")
            if file:
                workbook = openpyxl.load_workbook(file)
                sheet = workbook.active
                for excel_row in sheet.iter_rows(min_row=2, values_only=True):
                    source_val = str(excel_row[3]).strip() if len(excel_row) > 3 and excel_row[3] else "Manual"
                    Platform.objects.get_or_create(
                        company=employee.company, 
                        name=source_val
                    )
                    row = LeadRow.objects.create(
                        lead_collection=lead_collection,
                        full_name=excel_row[0] if len(excel_row) > 0 else "",
                        email=excel_row[1] if len(excel_row) > 1 else "",
                        contact=excel_row[2] if len(excel_row) > 2 else "",
                        source=excel_row[3] if len(excel_row) > 3 else "",
                        collected_by=employee
                    )
                    for i, field in enumerate(fields, start=4):
                        if len(excel_row) > i:
                            LeadRowValue.objects.create(
                                lead_row=row,
                                field=field,
                                value=str(excel_row[i]) if excel_row[i] else ""
                            )
            return redirect("client_lead_data", lead_id=lead_id)


    rows = LeadRow.objects.filter(lead_collection=lead_collection).order_by("-id")


    selected_status = request.GET.get("status")

    if selected_status == "Transferred":
        rows = rows.filter(is_transferred=True)

    elif selected_status:
        rows = rows.filter(status=selected_status)

    else:
        rows = rows.filter(status="Unverified")

    selected_employee = request.GET.get("employee", "")
    if selected_employee:
        rows = rows.filter(collected_by_id=selected_employee)

    unverified_count = LeadRow.objects.filter(
    lead_collection=lead_collection,
    status="Unverified"
).count()

 
    page_size = request.GET.get('page_size', 10)
    try:
        page_size = int(page_size)
    except ValueError:
        page_size = 10

    paginator = Paginator(rows, page_size)
    page_number = request.GET.get('page')
    page_obj = paginator.get_page(page_number)
    for lead in page_obj:
        if lead.created_at:
            lead.created_at_ist = lead.created_at + timedelta(hours=5, minutes=30)

    table_rows = []
    for row in page_obj:
        row_values = [
            row.full_name,
            row.email,
            row.contact,
            row.source
        ]
        for field in fields:
            value = row.values.filter(field=field).first()
            row_values.append(value.value if value else "")
        created_info = f"{row.collected_by.name} ({row.created_at.strftime('%d %b %Y')})" if row.collected_by else str(row.created_at.date())
        table_rows.append({
            "id": row.id,
            "created_info": created_info,
            "values": row_values,
            "status": row.status
        })

   
    employees = EmployeeRegister_Details.objects.filter(company=employee.company,department=employee.department)
    print(employees) 
    

    return render(request, "client_lead_data.html", {
        "lead_collection": lead_collection,
        "fields": fields,
        "table_rows": table_rows,
        "page_obj": page_obj,
        "head_name": employee.name,
        "platforms": platforms,
        "unverified_count": unverified_count,
        "selected_status": selected_status,
        "selected_employee": selected_employee,
        "page_size": page_size,
        "employees": employees
    })




@csrf_exempt
def apply_lead_action(request):
    if request.method == "POST":
        data = json.loads(request.body)
        lead_ids = data.get("lead_ids", [])
        action = data.get("action")

        if not lead_ids or not action:
            return JsonResponse({"success": False, "error": "Missing data"})

        leads = LeadRow.objects.filter(id__in=lead_ids)

        for lead in leads:

            if action == "Delete":
                lead.delete()
                continue

            elif action == "Verified":
                lead.status = "Verified"

            elif action == "Incomplete":
                lead.status = "Incomplete"

            elif action == "Waste":
                lead.status = "Waste"

            elif action == "Unverified":
                lead.status = "Unverified"

            elif action == "Transferred":
                lead.is_transferred = True  

            lead.save()

        return JsonResponse({"success": True})

    return JsonResponse({"success": False, "error": "Invalid request"})





def verified_leads_page(request):
    user_id = request.session.get('user_id')

    if not user_id:
        return redirect('login')

    user = LogRegister_Details.objects.filter(
        id=user_id
    ).first()

    employee = EmployeeRegister_Details.objects.filter(
        login=user
    ).first()

    if not employee:
        return redirect('login')

    
    clients = ClientRegister.objects.filter(
    compId=employee.company,
    workregister__leadcollection__isnull=False
).distinct()

 
    same_department_employees = EmployeeRegister_Details.objects.filter(
        company=employee.company,
        department=employee.department,
        active_status="Approved" )

    verified_leads = LeadRow.objects.filter(
    status="Verified",
    is_transferred=False,   
    lead_collection__work_Id__clientId__compId=employee.company
).select_related(
    'lead_collection',
    'collected_by'
).order_by('-created_at')

   
    page_size = request.GET.get("page_size", 10)
    page_number = request.GET.get("page")

    paginator = Paginator(verified_leads, int(page_size))
    page_obj = paginator.get_page(page_number)
    for lead in page_obj:
        if lead.created_at:
            lead.created_at_ist = lead.created_at + timedelta(hours=5, minutes=30)

    context = {
        'page_obj': page_obj,
        'employee': employee,
        'clients': clients,
        'same_department_employees': same_department_employees,
        "head_name": employee.name,
    }

    return render(
        request,
        'dm_verified_leads_page.html',
        context
    )


def get_client_categories(request):
    client_id = request.GET.get("client_id")

    categories = LeadCollection.objects.filter(
        work_Id__clientId_id=client_id
    ).values("id", "collection_head")

    return JsonResponse(list(categories), safe=False)



def transfer_verified_leads(request):

    if request.method=="POST":

        data=json.loads(request.body)

        lead_ids=data.get("lead_ids",[])

        employee_id=data.get("employee_id")

        if not lead_ids:
            return JsonResponse({"success":False})

        LeadRow.objects.filter(
            id__in=lead_ids
        ).update(
            is_transferred=True,
            transfer_date=timezone.now(),
            transferred_to_id=employee_id
        )

        return JsonResponse({
            "success":True,
            "message":"Leads transferred successfully",
        })

    return JsonResponse({"success":False})



def transferred_leads_page(request):
    user_id = request.session.get('user_id')
    if not user_id:
        return redirect('login')

    
    user = LogRegister_Details.objects.filter(id=user_id).first()
    employee = EmployeeRegister_Details.objects.filter(login=user).first()
    if not employee:
        return redirect('login')


    today_filter = request.GET.get("today", "today")      
    client_id = request.GET.get("client")                
    employee_id = request.GET.get("employee")           
    category_id = request.GET.get("category")            
    start_date = request.GET.get("startDate")             
    end_date = request.GET.get("endDate")                


    transferred_leads = LeadRow.objects.filter(
    is_transferred=True,
    lead_collection__work_Id__clientId__compId=employee.company
)


    if today_filter == "today":
        today_date = now().date()
        transferred_leads = transferred_leads.filter(transfer_date__date=today_date)

    if client_id:
        transferred_leads = transferred_leads.filter(
            lead_collection__work_Id__clientId__id=client_id
        )

    if employee_id:
        transferred_leads = transferred_leads.filter(collected_by__id=employee_id)

    if category_id:
        transferred_leads = transferred_leads.filter(lead_collection__id=category_id)

    if start_date:
        transferred_leads = transferred_leads.filter(transfer_date__date__gte=start_date)

    if end_date:
        transferred_leads = transferred_leads.filter(transfer_date__date__lte=end_date)

 
    transferred_leads = transferred_leads.select_related(
        "collected_by",
        "transferred_to",
        "lead_collection",
        "lead_collection__work_Id",
        "lead_collection__work_Id__clientId"
    ).order_by("-transfer_date")

    
    clients = ClientRegister.objects.filter(
    compId=employee.company,
    workregister__leadcollection__isnull=False
).distinct()

    
    if client_id:
        categories = LeadCollection.objects.filter(work_Id__clientId__id=client_id).distinct()
    else:
        categories = LeadCollection.objects.all().distinct()


    employees = EmployeeRegister_Details.objects.filter(
        company=employee.company,
        department=employee.department,
        active_status="Approved"
    ).distinct()

   
    return render(request, "dm_transferred_leads_page.html", {
        "transferred_leads": transferred_leads,
        "clients": clients,
        "employees": employees,
        "categories": categories,
        "today_filter": today_filter,
        "selected_client": client_id or "",
        "selected_employee": employee_id or "",
        "selected_category": category_id or "",
        "start_date": start_date or "",
        "end_date": end_date or "",
         "head_name": employee.name,
    })
    
def tl_works_home(request):
    user_id = request.session.get('user_id')
    
    if not user_id or request.session.get('position', '').lower() != 'team_lead':
        return redirect('login')

    try:
        user = LogRegister_Details.objects.get(id=user_id)
        employee = EmployeeRegister_Details.objects.get(login=user)
        
        context = {
            'teamlead_name': employee.name,  
            'employee': employee,
            'user': user
        }
        
        return render(request, 'tl_work_home.html', context)
        
    except (LogRegister_Details.DoesNotExist, EmployeeRegister_Details.DoesNotExist):
        messages.error(request, 'Employee details not found')
        return redirect('login')
    

from .models import teamleadallocation    
from django.utils.timezone import now
def teamlead_allocated_works(request):
    user_id = request.session.get('user_id')

    if not user_id or request.session.get('position', '').lower() != 'team_lead':
        return redirect('login')

    try:
        user = LogRegister_Details.objects.get(id=user_id)
        employee = EmployeeRegister_Details.objects.get(login=user)
        if request.method == "POST":
            work_id = request.POST.get('work_id') 
            task_id = request.POST.get('task')
            employee_ids = request.POST.getlist('employees')

            start_date = request.POST.get('start_date')
            end_date = request.POST.get('end_date')
            target = request.POST.get('target')
            description = request.POST.get('description')

            

            file = request.FILES.get('file')
            lead_collection_id = request.POST.get('lead_collection')

            for emp_id in employee_ids:
                teamleadallocation.objects.create(
                    work_id=work_id, 
                    team_lead=employee,           
                    assigned_to_id=emp_id,    
                    task_id=task_id,
                    start_date=start_date,
                    end_date=end_date,
                    target=target or 0,
                    description=description,
                    
                    file=file,
                    lead_collection_id=lead_collection_id if lead_collection_id else None,
                       # important (so it shows in your filter)
                )

            messages.success(request, "Task assigned successfully ")
            return redirect('teamlead_allocated_works')

        allocations = LeadAllocation.objects.filter(
            team_lead=employee
        ).exclude(
            work_type__iexact='single'
        ).select_related('work', 'task', 'lead_collection')

        context = {
            'teamlead_name': employee.name,
            'employee': employee,
            'user': user,
            'allocations': allocations,
            'tasks': ClientTask_Register.objects.all(),
            'employees': EmployeeRegister_Details.objects.filter(
    id__in=Allocation_Details.objects.filter(
        allocat_to_id=employee   # logged-in team lead
    ).values_list('allocatEmp_id', flat=True),
    designation__dashboard_id='Executive'
).distinct(),
            'leads': LeadCollection.objects.all(),
            'today': now().date()
        }

        return render(request, 'tl_work_allocation.html', context)

    except (LogRegister_Details.DoesNotExist, EmployeeRegister_Details.DoesNotExist):
        messages.error(request, 'Employee not found')
        return redirect('login')       
    

def teamlead_pending_works(request):
    user_id = request.session.get('user_id')

    if not user_id or request.session.get('position', '').lower() != 'team_lead':
        return redirect('login')

    try:
        user = LogRegister_Details.objects.get(id=user_id)
        employee = EmployeeRegister_Details.objects.get(login=user)

        Tasks = teamleadallocation.objects.filter(
            team_lead=employee
        ).select_related(
            'task',
            'assigned_to',
            'lead_collection'
        ).prefetch_related('dailywork_set').order_by('-allocated_date')

        for task in Tasks:
            daily_works = task.dailywork_set.all()

            total_verified = 0
            total_target = 0

            for dw in daily_works:
                total_verified += dw.verified_target or 0
                total_target += dw.target or 0

            if total_target > 0:
                percent = int((total_verified / total_target) * 100)
                percent = min(percent, 100)
            else:
                percent = 0

            task.progress_percent = percent
            task.total_verified_target = total_verified
            latest_work = daily_works.last()
            task.latest_verified_target = (
                latest_work.verified_target if latest_work else 0
            )

        context = {
            'teamlead_name': employee.name,
            'Tasks': Tasks,
            'employee': employee,
            'user': user
        }

        return render(request, 'tl_pending_work.html', context)

    except (LogRegister_Details.DoesNotExist, EmployeeRegister_Details.DoesNotExist):
        messages.error(request, "User not found")
        return redirect('login')


def executive_work_home(request):
    user_id = request.session.get('user_id')
    
    if not user_id or request.session.get('position', '').lower() != 'executive':
        return redirect('login')

    try:
        user = LogRegister_Details.objects.get(id=user_id)
        employee = EmployeeRegister_Details.objects.get(login=user)
        
        context = {
            'executive_name': employee.name,  
            'employee': employee,
            'user': user
        }
        
        return render(request, 'executive_work_home.html', context)
        
    except (LogRegister_Details.DoesNotExist, EmployeeRegister_Details.DoesNotExist):
        messages.error(request, 'Employee details not found')
        return redirect('login')
    
def executive_new_work(request):
    user_id = request.session.get('user_id')

    if not user_id or request.session.get('position', '').lower() != 'executive':
        return redirect('login')

    try:
        user = LogRegister_Details.objects.get(id=user_id)
        employee = EmployeeRegister_Details.objects.get(login=user)

        Tasks = teamleadallocation.objects.filter(
            assigned_to=employee,
            status='Pending'
        ).select_related(
            'work',
            'task',
            'assigned_to',
            'lead_collection'
        ).order_by('-allocated_date')

        context = {
            'executive_name': employee.name,
            'Tasks': Tasks,
            'employee': employee,
            'user': user
        }

        return render(request, 'executive_new_work.html', context)

    except (LogRegister_Details.DoesNotExist, EmployeeRegister_Details.DoesNotExist):
        messages.error(request, "User not found")
        return redirect('login')
    
def accept_task(request, id):
    task = get_object_or_404(teamleadallocation, id=id)

    task.status = 'Accepted'
    task.accepted_date = timezone.now()  
    task.save()

    messages.success(request, "Task Accepted Successfully")
    return redirect('executive_new_work')

def executive_task_ongoing(request):
    user_id = request.session.get('user_id')

    if not user_id or request.session.get('position', '').lower() != 'executive':
        return redirect('login')

    try:
        user = LogRegister_Details.objects.get(id=user_id)
        employee = EmployeeRegister_Details.objects.get(login=user)
        Tasks = teamleadallocation.objects.filter(
            assigned_to=employee,
            status='Accepted'
        ).select_related(
            'work',
            'task',
            'assigned_to',
            'lead_collection'
        ).order_by('-allocated_date')

        for allocation in Tasks:
            daily_works = DailyWork.objects.filter(allocation=allocation)

            total_verified = sum(dw.verified_target or 0 for dw in daily_works)
            total_target = sum(dw.target or 0 for dw in daily_works)

            if total_target > 0:
                allocation.progress_percent = int((total_verified / total_target) * 100)
                if allocation.progress_percent > 100:
                    allocation.progress_percent = 100
            else:
                allocation.progress_percent = 0

        context = {
            'executive_name': employee.name,
            'Tasks': Tasks,
            'employee': employee,
            'user': user
        }

        return render(request, 'execotive_task_ongoing.html', context)

    except (LogRegister_Details.DoesNotExist, EmployeeRegister_Details.DoesNotExist):
        messages.error(request, "User not found")
        return redirect('login')


def executive_adddaily_work(request, id):
    user_id = request.session.get('user_id')

    if not user_id or request.session.get('position', '').lower() != 'executive':
        return redirect('login')

    try:
        user = LogRegister_Details.objects.get(id=user_id)
        employee = EmployeeRegister_Details.objects.get(login=user)
        allocation = get_object_or_404(
            teamleadallocation.objects.select_related(
                'work', 'task', 'assigned_to', 'lead_collection'
            ),
            id=id,
            assigned_to=employee
        )
        daily_works = DailyWork.objects.filter(allocation=allocation)

        total_verified = 0
        total_target = 0

        for dw in daily_works:
            total_verified += dw.verified_target or 0
            total_target += dw.target or 0

        if total_target > 0:
            percent = int((total_verified / total_target) * 100)
            if percent > 100:
                percent = 100
        else:
            percent = 0

        allocation.progress_percent = percent  # attach dynamically for template

        context = {
            'executive_name': employee.name,
            'Task': allocation,   # send the allocation to template
            'employee': employee,
            'user': user
        }

        return render(request, 'executive_adddaily_work.html', context)

    except (LogRegister_Details.DoesNotExist, EmployeeRegister_Details.DoesNotExist):
        messages.error(request, "User not found")
        return redirect('login')
    
from datetime import date
from .models import DailyWork

def executive_adddailywork(request, id):
    user_id = request.session.get('user_id')

    # 🔐 Check login + role
    if not user_id or request.session.get('position', '').lower() != 'executive':
        return redirect('login')

    try:
        user = LogRegister_Details.objects.get(id=user_id)
        employee = EmployeeRegister_Details.objects.get(login=user)

        Task = get_object_or_404(
            teamleadallocation.objects.select_related(
                'work', 'task', 'assigned_to', 'lead_collection'
            ),
            id=id,
            assigned_to=employee
        )

        today = date.today()
        if request.method == "POST":
            title = request.POST.get('title')
            description = request.POST.get('description')
            target = request.POST.get('target')
            work_date = request.POST.get('work_date')
            file = request.FILES.get('file')

            if not title or not description:
                messages.error(request, "Title and Description are required")
                return redirect('executive_adddailywork', id=id)

            
            DailyWork.objects.create(
                employee=employee,
                task=Task.task,
                allocation=Task,
                title=title,
                description=description,
                target=target if target else 0,
                work_date=work_date,
                file=file
            )

            if target:
                Task.target += int(target)
                Task.save()

            messages.success(request, "Daily work added successfully ")

            return redirect('executive_task_ongoing')

        context = {
            'executive_name': employee.name,
            'Task': Task,
            'employee': employee,
            'user': user,
            'today': today
        }

        return render(request, 'executive_adddailywork.html', context)

    except (LogRegister_Details.DoesNotExist, EmployeeRegister_Details.DoesNotExist):
        messages.error(request, "User not found")
        return redirect('login')
    
def executive_view_dailywork(request, id):
    user_id = request.session.get('user_id')

    if not user_id or request.session.get('position', '').lower() != 'executive':
        return redirect('login')

    try:
        user = LogRegister_Details.objects.get(id=user_id)
        employee = EmployeeRegister_Details.objects.get(login=user)

        
        dailywork = DailyWork.objects.select_related('allocation').filter(
            allocation__id=id,
            employee=employee
        )

        from_date = request.GET.get('from_date')
        to_date = request.GET.get('to_date')

        if from_date:
            dailywork = dailywork.filter(work_date__gte=from_date)

        if to_date:
            dailywork = dailywork.filter(work_date__lte=to_date)

        context = {
            'executive_name': employee.name,
            'dailywork': dailywork,
            'employee': employee,
            'user': user,
        }

        return render(request, 'executive_view_dailywork.html', context)

    except (LogRegister_Details.DoesNotExist, EmployeeRegister_Details.DoesNotExist):
        messages.error(request, "User not found")
        return redirect('login')
    
def delete_allocation(request, id):
    allocation = get_object_or_404(teamleadallocation, id=id)

    allocation.delete()

    messages.success(request, "Allocation removed successfully")
    return redirect('teamlead_pending_works')

from django.http import JsonResponse

def get_dailyworks(request, task_id):
    works = DailyWork.objects.filter(allocation__id=task_id)
    data = []
    for work in works:
        
        data.append({
            "id": work.id,
            'date': work.work_date.strftime("%d %b %Y") if work.work_date else '',
            'title': work.title or '',
            'description': work.description or '',
            'target': work.target or '',
            'file': work.file.url if work.file else '',
            'verified_target': work.verified_target if work.verified_target is not None else '',
            'verification': work.verification or ''
            
        })

    return JsonResponse({'works': data})
    
def update_verification(request):
    if request.method == "POST":
        work_id = request.POST.get("work_id")
        verified = request.POST.get("verified_target")

        work = DailyWork.objects.get(id=work_id)
        work.verified_target = int(verified) if verified else 0 
        work.verification = "Verified"
        work.save()
    return redirect('teamlead_pending_works')
    
def tl_individualwork_home(request):
    user_id = request.session.get('user_id')
    
    if not user_id or request.session.get('position', '').lower() != 'team_lead':
        return redirect('login')

    try:
        user = LogRegister_Details.objects.get(id=user_id)
        employee = EmployeeRegister_Details.objects.get(login=user)
        
        context = {
            'teamlead_name': employee.name,  
            'employee': employee,
            'user': user
        }
        
        return render(request, 'tl_individualwork_home.html', context)
        
    except (LogRegister_Details.DoesNotExist, EmployeeRegister_Details.DoesNotExist):
        messages.error(request, 'Employee details not found')
        return redirect('login')
def tl_individualwork(request):
    user_id = request.session.get('user_id')

    if not user_id or request.session.get('position', '').lower() != 'team_lead':
        return redirect('login')

    try:
        user = LogRegister_Details.objects.get(id=user_id)
        employee = EmployeeRegister_Details.objects.get(login=user)

        allocations = LeadAllocation.objects.filter(
            team_lead=employee,
            work_type__iexact='single'
        ).exclude(status='Accepted').select_related(
            'work',
            'task',
            'lead_collection'
        ).order_by('-allocated_date')

        context = {
            'teamlead_name': employee.name,
            'employee': employee,
            'user': user,
            'allocations': allocations
        }

        return render(request, 'tl_individualwork.html', context)

    except (LogRegister_Details.DoesNotExist, EmployeeRegister_Details.DoesNotExist):
        messages.error(request, "User not found")
        return redirect('login')
 
    
def accept_task1(request, id):
    allocations = get_object_or_404(LeadAllocation, id=id)

    allocations.status = 'Accepted'
    allocations.accepted_date = timezone.now()  
    allocations.save()

    messages.success(request, "Task Accepted Successfully")
    return redirect('tl_individualwork')
from django.db.models import Sum


def tl_task_ongoing(request):
    user_id = request.session.get('user_id')

    if not user_id or request.session.get('position', '').lower() != 'team_lead':
        return redirect('login')

    try:
        user = LogRegister_Details.objects.get(id=user_id)
        employee = EmployeeRegister_Details.objects.get(login=user)

        allocations = LeadAllocation.objects.filter(
            team_lead=employee,
            status='Accepted'
        ).select_related(
            'work',
            'task',
            'lead_collection'
        ).annotate(
            total_verified_target=Sum('dailywork2__verified_target'),
            total_target=Sum('dailywork2__target')
        ).order_by('-allocated_date')

        for allocation in allocations:
            total_verified = allocation.total_verified_target or 0
            total_target = allocation.total_target or 0

            if total_target > 0:
                allocation.progress_percent = int((total_verified / total_target) * 100)
                allocation.progress_percent = min(allocation.progress_percent, 100)
            else:
                allocation.progress_percent = 0
        context = {
            'teamlead_name': employee.name,
            'allocations': allocations,
            'employee': employee,
            'user': user
        }

        return render(request, 'tl_ongoingwork.html', context)

    except (LogRegister_Details.DoesNotExist, EmployeeRegister_Details.DoesNotExist):
        messages.error(request, "User not found")
        return redirect('login')

from django.db.models import Sum

def tl_adddaily_work(request, id):
    user_id = request.session.get('user_id')

    if not user_id or request.session.get('position', '').lower() != 'team_lead':
        return redirect('login')

    try:
        user = LogRegister_Details.objects.get(id=user_id)
        employee = EmployeeRegister_Details.objects.get(login=user)

        allocation = get_object_or_404(
            LeadAllocation.objects.select_related(
                'work', 'task', 'lead_collection'
            ).annotate(
                total_verified_target=Sum('dailywork2__verified_target'),
                total_target=Sum('dailywork2__target')
            ),
            id=id,
            team_lead=employee
        )

        total_verified = allocation.total_verified_target or 0
        total_target = allocation.total_target or 0

        if total_target > 0:
            allocation.progress_percent = int((total_verified / total_target) * 100)
            allocation.progress_percent = min(allocation.progress_percent, 100)
        else:
            allocation.progress_percent = 0

        context = {
            'teamlead_name': employee.name,
            'Task': allocation,   # single allocation
            'employee': employee,
            'user': user
        }

        return render(request, 'tl_adddaily_work.html', context)

    except (LogRegister_Details.DoesNotExist, EmployeeRegister_Details.DoesNotExist):
        messages.error(request, "User not found")
        return redirect('login')

def tl_adddailywork(request, id):
    user_id = request.session.get('user_id')


    if not user_id or request.session.get('position', '').lower() != 'team_lead':
        return redirect('login')

    try:
        user = LogRegister_Details.objects.get(id=user_id)
        employee = EmployeeRegister_Details.objects.get(login=user)

        singleallocations = get_object_or_404(
            LeadAllocation.objects.select_related(
                'work', 'task','lead_collection'
            ),
            id=id,
            team_lead=employee
        )

        today = date.today()
        if request.method == "POST":
            title = request.POST.get('title')
            description = request.POST.get('description')
            target = request.POST.get('target')
            work_date = request.POST.get('work_date')
            file = request.FILES.get('file')

            if not title or not description:
                messages.error(request, "Title and Description are required")
                return redirect('tl_adddailywork', id=id)

            
            DailyWork2.objects.create(
                employee=employee,
                task=singleallocations.task,
                singleallocation=singleallocations,
                title=title,
                description=description,
                target=target if target else 0,
                work_date=work_date,
                file=file
            )

            if target:
                singleallocations.target += int(target)
                singleallocations.save()

            messages.success(request, "Daily work added successfully ")

            return redirect('tl_task_ongoing')

        context = {
            'teamlead_name': employee.name,
            'Task': singleallocations,
            'employee': employee,
            'user': user,
            'today': today
        }

        return render(request, 'tl_adddailywork.html', context)

    except (LogRegister_Details.DoesNotExist, EmployeeRegister_Details.DoesNotExist):
        messages.error(request, "User not found")
        return redirect('login')
    
def tl_view_dailywork(request, id):
    user_id = request.session.get('user_id')

    if not user_id or request.session.get('position', '').lower() != 'team_lead':
        return redirect('login')

    try:
        user = LogRegister_Details.objects.get(id=user_id)
        employee = EmployeeRegister_Details.objects.get(login=user)

        
        dailywork = DailyWork2.objects.select_related('singleallocation').filter(
            singleallocation__id=id,
            employee=employee
        )

        from_date = request.GET.get('from_date')
        to_date = request.GET.get('to_date')

        if from_date:
            dailywork = dailywork.filter(work_date__gte=from_date)

        if to_date:
            dailywork = dailywork.filter(work_date__lte=to_date)

        context = {
            'teamlead_name': employee.name,
            'dailywork': dailywork,
            'employee': employee,
            'user': user,
        }

        return render(request, 'tl_view_dailywork.html', context)

    except (LogRegister_Details.DoesNotExist, EmployeeRegister_Details.DoesNotExist):
        messages.error(request, "User not found")
        return redirect('login')

def executive_weekly_progress(request):
    user_id = request.session.get('user_id')

    if not user_id or request.session.get('position', '').lower() != 'executive':
        return redirect('login')

    try:
        user = LogRegister_Details.objects.get(id=user_id)
        employee = EmployeeRegister_Details.objects.get(login=user)
        allocation = teamleadallocation.objects.filter(assigned_to=employee).first()
       

        reports = progressreport.objects.filter(employee=employee,report_type="weekly")
        

        from_date = request.GET.get('from_date')
        to_date = request.GET.get('to_date')

        if from_date:
            reports = reports.filter(from_date__gte=from_date)

        if to_date:
            reports = reports.filter(to_date__lte=to_date)

        context = {
            'executive_name': employee.name,
            'reports': reports,
            
            'employee': employee,
            'user': user,
            'allocation': allocation,
            
        }
        return render(request, 'executive_weeklyprogress.html', context)

    except:
        messages.error(request, "User not found")
        return redirect('login')
    
def executive_add_weekly_report(request):
    user_id = request.session.get('user_id')

    if request.method == "POST":
        user = LogRegister_Details.objects.get(id=user_id)
        employee = EmployeeRegister_Details.objects.get(login=user)

        allocation_id = request.POST.get('allocation_id')
        if not allocation_id:
            messages.error(request, "Allocation ID missing")
            return redirect('executive_weekly_progress')

        from_date = request.POST.get('from_date')
        to_date = request.POST.get('to_date')
        progress = request.POST.get('progress')
        description = request.POST.get('description')
        file = request.FILES.get('file')
        report_type = request.POST.get('report_type')
        allocation = teamleadallocation.objects.get(id=allocation_id)

        progressreport.objects.create(
            employee=employee,
            allocation=allocation,
            from_date=from_date,
            to_date=to_date,
            progress=progress,
            description=description,
            file=file
        )

        messages.success(request, "Weekly report saved")
        return redirect('executive_weekly_progress')
    
def executive_monthly_progress(request):
    user_id = request.session.get('user_id')

    if not user_id or request.session.get('position', '').lower() != 'executive':
        return redirect('login')

    try:
        user = LogRegister_Details.objects.get(id=user_id)
        employee = EmployeeRegister_Details.objects.get(login=user)
        allocation = teamleadallocation.objects.filter(assigned_to=employee).first()
        reports = progressreport.objects.filter(employee=employee,report_type="monthly")
        

        from_date = request.GET.get('from_date')
        to_date = request.GET.get('to_date')

        if from_date:
            reports = reports.filter(from_date__gte=from_date)

        if to_date:
            reports = reports.filter(to_date__lte=to_date)

        context = {
            'executive_name': employee.name,
            'reports': reports,
            
            'employee': employee,
            'user': user,
            'allocation': allocation,
            
        }

        return render(request, 'executive_monthly_progress.html', context)

    except:
        messages.error(request, "User not found")
        return redirect('login')
    
def executive_add_monthly_report(request):
    user_id = request.session.get('user_id')

    if request.method == "POST":
        user = LogRegister_Details.objects.get(id=user_id)
        employee = EmployeeRegister_Details.objects.get(login=user)

        allocation_id = request.POST.get('allocation_id')
        if not allocation_id:
            messages.error(request, "Allocation ID missing")
            return redirect('executive_monthly_progress')

        from_date = request.POST.get('from_date')
        to_date = request.POST.get('to_date')
        progress = request.POST.get('progress')
        description = request.POST.get('description')
        file = request.FILES.get('file')
        report_type = request.POST.get('report_type')

        allocation = teamleadallocation.objects.get(id=allocation_id)

        progressreport.objects.create(
            employee=employee,
            allocation=allocation,
            from_date=from_date,
            to_date=to_date,
            progress=progress,
            description=description,
            file=file,
            report_type=report_type
        )

        messages.success(request, "monthly report saved")
        return redirect('executive_monthly_progress')

def head_weekly_progress(request):
    user_id = request.session.get('user_id')

    if not user_id or request.session.get('position', '').lower() != 'digital_marketing_head':
        return redirect('login')

    try:
        user = LogRegister_Details.objects.get(id=user_id)

        head_employee = EmployeeRegister_Details.objects.filter(login=user).first()
        if not head_employee:
            messages.error(request, "Employee not found")
            return redirect('login')

        employees = EmployeeRegister_Details.objects.filter(
            company=head_employee.company,
            designation__dashboard_id__in=['Executive', 'Team Lead']
        )

        reports = progressreport.objects.filter(
            report_type='weekly',
            employee__company=head_employee.company
        ).select_related('employee')

        employee_id = request.GET.get('employee')
        from_date = request.GET.get('from_date')
        to_date = request.GET.get('to_date')

        if employee_id:
            reports = reports.filter(employee__id=employee_id)

        if from_date:
            reports = reports.filter(from_date__gte=from_date)

        if to_date:
            reports = reports.filter(to_date__lte=to_date)

        context = {
            'head_name': head_employee.name,
            'reports': reports,
            'employees': employees, 
            'user': user,
        }

        return render(request, 'head_weekly_progress.html', context)

    except Exception as e:
        print("ERROR:", e)
        messages.error(request, "Something went wrong")
        return redirect('login')
    
def head_monthly_progress(request):
    user_id = request.session.get('user_id')

    if not user_id or request.session.get('position', '').lower() != 'digital_marketing_head':
        return redirect('login')

    try:
        user = LogRegister_Details.objects.get(id=user_id)

        head_employee = EmployeeRegister_Details.objects.filter(login=user).first()
        if not head_employee:
            messages.error(request, "Employee not found")
            return redirect('login')
        employees = EmployeeRegister_Details.objects.filter(
            company=head_employee.company,
            designation__dashboard_id__in=['Executive', 'Team Lead']
        )
        reports = progressreport.objects.filter(
            report_type='monthly',
            employee__company=head_employee.company
        ).select_related('employee')

        employee_id = request.GET.get('employee')
        from_date = request.GET.get('from_date')
        to_date = request.GET.get('to_date')

        if employee_id:
            reports = reports.filter(employee__id=employee_id)

        if from_date:
            reports = reports.filter(from_date__gte=from_date)

        if to_date:
            reports = reports.filter(to_date__lte=to_date)

        context = {
            'head_name': head_employee.name,
            'reports': reports,
            'employees': employees, 
            'user': user,
        }
        return render(request, 'head_monthly_progress.html', context)

    except Exception as e:
        print("ERROR:", e)
        messages.error(request, "Something went wrong")
        return redirect('login')


def verify_report(request):
    user_id = request.session.get('user_id')
    if not user_id or request.session.get('position', '').lower() != 'digital_marketing_head':
        return JsonResponse({'success': False, 'error': 'Unauthorized'})

    if request.method == "POST":
        report_id = request.POST.get('id')

        try:
            report = progressreport.objects.get(id=report_id)
            report.status = 'verified'
            report.save()

            return JsonResponse({'success': True})
        except progressreport.DoesNotExist:
            return JsonResponse({'success': False, 'error': 'Not found'})

    return JsonResponse({'success': False, 'error': 'Invalid request'})

def progress_updation(request):
    if request.method == "POST":
        allocation_id = request.POST.get("work_id")
        verified = int(request.POST.get("verified_target") or 0)

        daily_works = DailyWork2.objects.filter(singleallocation_id=allocation_id)

        for work in daily_works:
            work.verified_target = verified
            work.verification = "Verified"
            work.save()

    return redirect('tl_task_ongoing')
    
def data_manager_dataBank(request):
    client = request.GET.get('client')
    category = request.GET.get('category')
    hr = request.GET.get('hr')
    status = request.GET.get('status')
    start = request.GET.get('start')
    end = request.GET.get('end')

    company_id = request.session.get('company_id')
    leads = LeadRow.objects.filter(is_transferred=True,
    lead_collection__work_Id__clientId__compId_id=company_id)

    # Client filter
    if client:
        leads = leads.filter(
            lead_collection__work_Id__clientId_id=client
        )

    # Category filter
    if category:
        work_ids = ClientTask_Register.objects.filter(
            id=category
        ).values_list('work_Id_id', flat=True)

        leads = leads.filter(
            lead_collection__work_Id_id__in=work_ids
        )

    # HR filter
    if hr:
        leads = leads.filter(collected_by_id=hr)

    if status:
        if status == "Not Allocated":
            leads = leads.filter(status="Not Attended")  
        else:
            leads = leads.filter(status=status)

    # Date filters
    if start:
        leads = leads.filter(created_at__date__gte=start)

    if end:
        leads = leads.filter(created_at__date__lte=end)

    leads = leads.distinct()
    
    page_size = request.GET.get('page_size', 20)  # default 20
    page_size = int(page_size)

    paginator = Paginator(leads, page_size)
    page_number = request.GET.get('page')
    page_obj = paginator.get_page(page_number)

    company_id = request.session.get('company_id')
    employee_id = request.session.get('employee_id')

    clients = ClientRegister.objects.filter(compId_id=company_id).distinct()

    logged_employee = EmployeeRegister_Details.objects.get(id=employee_id)

    telecallerList = EmployeeRegister_Details.objects.filter(
        company=logged_employee.company,
        department=logged_employee.department,
        login__position='Telecaller'
    ).exclude(id=logged_employee.id)

    return render(request, 'data_manager_dataBank.html', {
        "leads": page_obj,
        "clients": clients,
        "telecallers": telecallerList,
        "page_obj": page_obj,
        "page_size": page_size
    })

def data_manager_allocation(request):
    company_id = request.session.get('company_id')
    employee_id = request.session.get('employee_id')
    logged_employee = EmployeeRegister_Details.objects.get(id=employee_id)

    clients = ClientRegister.objects.filter(compId_id=company_id).distinct() 

    executives = EmployeeRegister_Details.objects.filter(
        company=logged_employee.company,
        department=logged_employee.department,
        login__position__in=['Executive', 'Team_Lead', 'Digital_Marketing_Head']
    ).exclude(id=logged_employee.id)

    leads = LeadRow.objects.filter(
    lead_collection__work_Id__clientId__compId_id=company_id,is_transferred=True
    ).exclude(status="Allocated").select_related('lead_collection','collected_by').order_by('-created_at')

     # GET params
    client = request.GET.get("client")
    category = request.GET.get("category")
    hr = request.GET.get("hr")
    start_date = request.GET.get("start_date")
    end_date = request.GET.get("end_date")
    status = request.GET.get("status")
    executive = request.GET.get("executive")

    if client:
        leads = leads.filter(lead_collection__work_Id__clientId_id=client)

    if category:
        leads = leads.filter(lead_collection__work_Id__id=category)

    if hr:
        leads = leads.filter(transferred_to_id=hr)

    if start_date:
        leads = leads.filter(created_at__date__gte=start_date)

    if end_date:
        leads = leads.filter(created_at__date__lte=end_date)
    if status:
        leads = leads.filter(status=status)
    if executive:
        leads = leads.filter(collected_by_id = executive)

    page_size = request.GET.get('page_size', 20)  # default 20
    page_size = int(page_size)

    paginator = Paginator(leads, page_size)
    page_number = request.GET.get('page')
    page_obj = paginator.get_page(page_number)

    telecallerList = EmployeeRegister_Details.objects.filter(
        company=logged_employee.company,
        department=logged_employee.department,
        login__position='Telecaller'
    ).exclude(id=logged_employee.id)
    
    return render(request,'data_manager_allocation.html',{'clients':clients,'executives':executives,'leads':page_obj,'telecallers':telecallerList,
                                                          'page_size': page_size})


def data_manager_followup(request):
    company_id = request.session.get('company_id')
    leads = LeadRow.objects.filter(lead_collection__work_Id__clientId__compId_id=company_id,status = "Allocated")     
    employee_id = request.session.get('employee_id')
    logged_employee = EmployeeRegister_Details.objects.get(id=employee_id)
    telecallerList = EmployeeRegister_Details.objects.filter(
        company=logged_employee.company,
        department=logged_employee.department,
        login__position='Telecaller'
    ).exclude(id=logged_employee.id)
    statuses = FollowUpStatus.objects.all()

    status = request.GET.get("status")
    hr = request.GET.get("hr")
    start = request.GET.get("start")
    end = request.GET.get("end")


    if status:
        leads = leads.filter(status=status)

    if hr:
        leads = leads.filter(transferred_to_id=hr)

    if start and end:
        leads = leads.filter(created_at__date__range=[start, end])
        
    page_size = request.GET.get('page_size', 20)  # default 20
    page_size = int(page_size)

    paginator = Paginator(leads, page_size)
    page_number = request.GET.get('page')
    page_obj = paginator.get_page(page_number)

    return render(request, 'data_manager_followUp.html', {
        'leads': page_obj,
        'telecallers': telecallerList,
        'followup_statuses':statuses,
        'page_size':page_size
    })


def get_categories_by_client(request):
    client_id = request.GET.get('client_id')
    tasks = LeadCollection.objects.filter(work_Id__clientId_id = client_id).values('id', 'collection_head')
    return JsonResponse(list(tasks), safe=False)


def allocate_leads(request):
    if request.method == "POST":
        data = json.loads(request.body)

        lead_ids = data.get("lead_ids", [])
        hr_id = data.get("hr_id")

        try:
            hr = EmployeeRegister_Details.objects.get(id=hr_id)

            leads = LeadRow.objects.filter(id__in=lead_ids)

            leads.update(
                transferred_to=hr,   # field name adjust if different
                status="Allocated",
                is_transferred=True
            )

            return JsonResponse({"success": True})

        except Exception as e:
            return JsonResponse({"success": False, "message": str(e)})

    return JsonResponse({"success": False, "message": "Invalid request"})



def lead_details(request, id):
    lead = get_object_or_404(
        LeadRow.objects.select_related(
            'lead_collection__work_Id__clientId',
            'transferred_to',
            'collected_by'
        ),
        id=id
    )

    data = {
        # 🔹 Lead Details
        "full_name": lead.full_name,
        "email": lead.email,
        "contact": lead.contact,
        "status": lead.status,
        "source": lead.source,
        "created_at": lead.created_at.strftime("%d %b %Y"),

        # 🔹 Client Details (your requirement)
        "client_name": (
            lead.lead_collection.work_Id.clientId.client_name
            if lead.lead_collection and lead.lead_collection.work_Id and lead.lead_collection.work_Id.clientId
            else ""
        ),
        "collection_head":(
            lead.lead_collection.collection_head if lead.lead_collection else ""
        ),

        # 🔹 Assigned HR (transferred_to)
        "assigned_to": (
            lead.transferred_to.name
            if lead.transferred_to else "Not Assigned"
        ),

        # 🔹 Collected By
        "collected_by": (
            lead.collected_by.name
            if lead.collected_by else ""
        ),

        # 🔹 More Details
        "date": (
            lead.created_at.strftime("%d %b %Y")
            if lead.created_at else ""
        ),
        "time":(
             lead.created_at.strftime("%I:%M %p") if lead.created_at else ""
        )


    }

    return JsonResponse(data)


def unassign_lead(request, id):
    try:
        lead = LeadRow.objects.get(id=id)
        lead.transferred_to = None
        lead.is_transferred = False
        lead.transfer_date = None
        lead.status = "Unverified"
        lead.save()

        return JsonResponse({"success": True})
    except LeadRow.DoesNotExist:
        return JsonResponse({"success": False})
    

def bulk_unassign_leads(request):
    data = json.loads(request.body)
    ids = data.get("ids", [])

    leads = LeadRow.objects.filter(id__in=ids)

    for lead in leads:
        lead.transferred_to = None
        lead.status = "Not Attended"
        lead.save()

    return JsonResponse({"success": True})


@require_POST
def add_followup_status(request):
    data = json.loads(request.body)
    name = data.get("name")

    status = FollowUpStatus.objects.create(name=name)

    return JsonResponse({
        "success": True,
        "id": status.id,
        "name": status.name
    })


@require_POST
def delete_followup_status(request, id):
    FollowUpStatus.objects.filter(id=id).delete()
    return JsonResponse({"success": True})


@csrf_exempt
def accept_leads(request):
    if request.method == "POST":
        data = json.loads(request.body)
        ids = data.get("ids", [])

        LeadRow.objects.filter(id__in=ids).update(status="Opened")

        return JsonResponse({"success": True})



def telecaller_allLeads(request):
    user_id = request.session.get('user_id')
    company_id = request.session.get('company_id')

    if not user_id:
        return redirect('login')

    user = LogRegister_Details.objects.get(id=user_id)
    employee = EmployeeRegister_Details.objects.get(login=user)

    leads = LeadRow.objects.filter(lead_collection__work_Id__clientId__compId_id=company_id,transferred_to=employee).order_by('id')

    #  Filters
    status = request.GET.get('status')
    start = request.GET.get('start')
    end = request.GET.get('end')

    if status:
        leads = leads.filter(status=status)

    if start:
        leads = leads.filter(created_at__date__gte=start)

    if end:
        leads = leads.filter(created_at__date__lte=end)

    #Pagination
    page_size = request.GET.get('page_size', 20)
    try:
        page_size = int(page_size)
    except:
        page_size = 20

    paginator = Paginator(leads, page_size)
    page_number = request.GET.get('page')

    page_obj = paginator.get_page(page_number)

    context = {
        'leads': page_obj,   
        'page_obj': page_obj,
        'page_size': page_size,
        'status': status,
        'start': start,
        'end': end,
    }

    return render(request, 'telecaller_allLeads.html', context)
    
def admin_action_taken_list(request):
    user_id = request.session.get('user_id')
    user = LogRegister_Details.objects.get(id=user_id)
    business = BusinessRegister_Details.objects.get(login=user) 
    employees = EmployeeRegister_Details.objects.filter(company=business)
    selected_emp_id = request.GET.get('employee_id')
    from_date = request.GET.get('from_date')
    to_date = request.GET.get('to_date')

    if selected_emp_id:
        actions = ActionTaken.objects.filter(act_emp_id__company=business, act_emp_id_id=selected_emp_id)
        
        if from_date and to_date:
            actions = actions.filter(action_date__range=[from_date, to_date])
        elif from_date:
            actions = actions.filter(action_date__gte=from_date)
        elif to_date:
            actions = actions.filter(action_date__lte=to_date)
            
        actions = actions.order_by('-action_date')
    else:
        
        actions = ActionTaken.objects.none()
    return render(request, 'admin_action_taken.html', {
        'employees': employees,
        'actions': actions,
        'selected_emp_id': selected_emp_id,
        'company_name': business.company_name,
        'owner_full_name': f"{business.owner_fname} {business.owner_lname}"
    })
    
    
def admin_feedback_list(request):
    user_id = request.session.get('user_id')
    user = LogRegister_Details.objects.get(id=user_id)
    business = BusinessRegister_Details.objects.get(login=user)
    employees = EmployeeRegister_Details.objects.filter(company=business)
    selected_emp_id = request.GET.get('employee_id')
    from_date = request.GET.get('from_date')
    to_date = request.GET.get('to_date')
    
    
    if selected_emp_id:
        feedbacks = Feedback.objects.filter(feedback_emp_id_id=selected_emp_id)
        if from_date and to_date:
            feedbacks = feedbacks.filter(feedback_date__range=[from_date, to_date])
        elif from_date:
            feedbacks = feedbacks.filter(feedback_date__gte=from_date)
        elif to_date:
            feedbacks = feedbacks.filter(feedback_date__lte=to_date)
            
        feedbacks = feedbacks.order_by('-feedback_date')
    else:
        feedbacks = Feedback.objects.none()

    return render(request, 'admin_feedback.html', {
        'employees': employees,
        'feedbacks': feedbacks,
        'selected_emp_id': selected_emp_id,
        'company_name': business.company_name,
        'owner_full_name': f"{business.owner_fname} {business.owner_lname}"
    })
    
    
    
def admin_complaints_list(request):
    user_id = request.session.get('user_id')
    user = LogRegister_Details.objects.get(id=user_id)
    business = BusinessRegister_Details.objects.get(login=user)   
    employees = EmployeeRegister_Details.objects.filter(company=business)  
    selected_emp_id = request.GET.get('employee_id')
    from_date = request.GET.get('from_date')
    to_date = request.GET.get('to_date')
    if selected_emp_id:
        complaints = Complaints.objects.filter(complaint_emp_id_id=selected_emp_id)
        
        if from_date and to_date:
            complaints = complaints.filter(complaint_date__range=[from_date, to_date])
        elif from_date:
            complaints = complaints.filter(complaint_date__gte=from_date)
        elif to_date:
            complaints = complaints.filter(complaint_date__lte=to_date)
            
        complaints = complaints.order_by('-complaint_date')
    else:
        complaints = Complaints.objects.none()

    return render(request, 'admin_complaints.html', {
        'employees': employees,
        'complaints': complaints,
        'selected_emp_id': selected_emp_id,
        'company_name': business.company_name,
        'owner_full_name': f"{business.owner_fname} {business.owner_lname}"
    })

def telecaller_followups(request):
    user_id = request.session.get('user_id')

    if not user_id:
        messages.error(request, "Please login first")
        return redirect("login")

    
    user = LogRegister_Details.objects.get(id=user_id)
    employee = EmployeeRegister_Details.objects.get(login=user)
    follow_reason=FollowUpStatus.objects.all()

    from_date = request.GET.get('from')
    to_date= request.GET.get('to')

    next_date=request.GET.get('next')
    status=request.GET.get('filter_status')
    try:
        statvalue=FollowUpStatus.objects.get(id=status)
    except:
        statvalue=None

    page_items=request.GET.get('pagination')

    leads=LeadRow.objects.filter(transferred_to=employee,status="Opened",is_waste=False).order_by('id')

    if from_date and to_date:
        from_date_obj = datetime.strptime(from_date, "%Y-%m-%d").date()
        to_date_obj = datetime.strptime(to_date, "%Y-%m-%d").date()
        leads = leads.filter(transfer_date__date__range=[from_date_obj, to_date_obj]).order_by('id')

    elif from_date:
        from_date_obj = datetime.strptime(from_date, "%Y-%m-%d").date()
        leads = leads.filter(transfer_date__date__gte=from_date_obj).order_by('id') 

    elif to_date:
        to_date_obj = datetime.strptime(to_date, "%Y-%m-%d").date()
        leads = leads.filter(transfer_date__date__lte=to_date_obj).order_by('id')

    if next_date:
        leads = leads.filter(followupupdates__next_update_date=next_date).distinct().order_by('-id')

    if status and statvalue:
        latest_followup = followUpUpdates.objects.filter(
            lead=OuterRef('pk')
        ).order_by('-updated')

        leads = leads.annotate(
            latest_followup_response_id=Subquery(latest_followup.values('response_id')[:1]),
            latest_followup_response_name=Subquery(latest_followup.values('response__name')[:1])
        ).filter(latest_followup_response_id=statvalue.id).distinct().order_by('-id')

    followups_count = leads.count()

    if page_items in ['10', '50', '100']:
        paginator = Paginator(leads, int(page_items))
        leads = paginator.get_page(request.GET.get('page'))

    elif page_items =='all':
        paginator = Paginator(leads, leads.count())
        leads = paginator.get_page(request.GET.get('page'))
    else:
        page_items = '10'
        paginator = Paginator(leads, 10)
        leads = paginator.get_page(request.GET.get('page'))


    return render(request, 'telecaller_followups.html',{'telecaller_name':employee.name.upper(),
                                                            'leads':leads,
                                                            'followups_count':followups_count,
                                                            "follow_reason":follow_reason,
                                                            "selected_from":from_date,
                                                            "selected_to":to_date,
                                                            "selected_next":next_date,
                                                            "selected_status":status,
                                                            "selected_pagination":page_items})


def telecaller_followup_updates(request,id):
    user_id = request.session.get('user_id')

    if not user_id:
        messages.error(request, "Please login first")
        return redirect("login")

    
    user = LogRegister_Details.objects.get(id=user_id)
    employee = EmployeeRegister_Details.objects.get(login=user)
    
    lead=LeadRow.objects.get(id=id)
    follow_reason=FollowUpStatus.objects.all()

    try:
        last_updates=followUpUpdates.objects.filter(lead=lead).order_by('-id')[0]
    except:
        last_updates=None

    try:
        all_followups=followUpUpdates.objects.filter(lead=lead).order_by('-id')
    except:
        all_followups=None
    try:
        collections=LeadRowValue.objects.filter(lead_row=lead)
    except:
        collections=None

    if request.method=="POST":
        response_id=request.POST.get('response')
        next_update_date=request.POST.get('next_update_date')
        reason=request.POST.get('reason')
        call_record=request.FILES.get('call_record')

        try:
            response=FollowUpStatus.objects.get(id=response_id)
        except:
            response=None
        try:
            followUpUpdates.objects.create(
                lead=lead,
                response=response,
                next_update_date=next_update_date,
                reason=reason,
                call_record=call_record,
                user=employee,
                notes='Lead Status Changed',
                updated=date.today()
            )
            messages.success(request,'followup updated successfully')
            return redirect('telecaller_followup_updates',id)
        except:
            messages.error(request,'something went wrong')
            return redirect('telecaller_followup_updates',id)


    return render(request,'telecaller_followup_updates.html',
                  {'id':id,
                    "lead":lead,
                    "follow_reason":follow_reason,
                    "last_updates":last_updates,
                    'telecaller_name':employee.name.upper(),
                    "all_followups":all_followups,
                    'collections':collections
                    })   


def change_lead_waste(request,id):
    if request.method=="POST":
        waste_reason=request.POST.get('waste_reason')
        lead=LeadRow.objects.get(id=id)
        lead.is_waste=True
        lead.waste_reason=waste_reason
        lead.waste_marked_date=date.today()
        lead.save()
        messages.success(request,'lead marked as waste')
        return redirect('telecaller_followups')
    
    return redirect('telecaller_followups')


def change_lead_close(request,id):
    lead=LeadRow.objects.get(id=id)
    lead.status="Closed"
    lead.save()
    user_id = request.session.get('user_id')

    if not user_id:
        messages.error(request, "Please login first")
        return redirect("login")

    
    user = LogRegister_Details.objects.get(id=user_id)
    employee = EmployeeRegister_Details.objects.get(login=user)
    followUpUpdates.objects.create(
                lead=lead,
                
                user=employee,
                notes='Lead is closed',
                is_close=True,
                close_date=date.today()
            )
    messages.success(request,'lead closed')
    return redirect('telecaller_followups')




def change_lead_recall(request,id):
    lead=LeadRow.objects.get(id=id)
    lead.status="Opened"
    lead.is_waste=False
    lead.waste_reason=None
    lead.is_recall=True
    lead.save()
    lastfollow=followUpUpdates.objects.filter(lead=lead).order_by('-id')[0]
    lastfollow.delete()

    messages.success(request,'lead recalled')
    return redirect('telecaller_followups')


def change_lead_joined(request,id):
    lead=LeadRow.objects.get(id=id)
    lead.status='Joined'
    lead.save()
    messages.success(request,'lead joined')
    return redirect('telecaller_followups')




def telecaller_wasteLeads(request):
    user_id = request.session.get('user_id')

    if not user_id:
        messages.error(request, "Please login first")
        return redirect("login")

    
    user = LogRegister_Details.objects.get(id=user_id)
    employee = EmployeeRegister_Details.objects.get(login=user)

    waste_leads=LeadRow.objects.filter(transferred_to=employee,is_waste=True).order_by('id')
    waste_count=waste_leads.count()

    from_date=request.GET.get('from')
    to_date=request.GET.get('to')
    status=request.GET.get('status')
    pagination=request.GET.get('pagination')
    if status:
        waste_leads=waste_leads.filter(waste_status=0 if status=='pending' else 1)
    
    if from_date and to_date:
        from_date_obj=datetime.strptime(from_date,"%Y-%m-%d").date()
        to_date_obj=datetime.strptime(to_date,"%Y-%m-%d").date()

        waste_leads=waste_leads.filter(waste_marked_date__date__range=[from_date_obj,to_date_obj]).order_by('id')
    elif from_date:
        from_date_obj=datetime.strptime(from_date,"%Y-%m-%d").date()
        waste_leads=waste_leads.filter(waste_marked_date__date__gte=from_date_obj).order_by('id')
    elif to_date:
        to_date_obj=datetime.strptime(to_date,"%Y-%m-%d").date()
        waste_leads=waste_leads.filter(waste_marked_date__date__lte=to_date_obj).order_by('id')

    if pagination in ['10','50','100']:
        paginator=Paginator(waste_leads,int(pagination))
        waste_leads=paginator.get_page(request.GET.get('page'))
    elif pagination == 'All':
        paginator=Paginator(waste_leads,waste_leads.count() or 1)
        waste_leads=paginator.get_page(request.GET.get('page'))
    else:
        pagination = '10'
        paginator=Paginator(waste_leads,10)
        waste_leads=paginator.get_page(request.GET.get('page'))
    

    return render(request,'telecaller_wasteLeads.html',
                  {"waste_leads":waste_leads,
                  "waste_count":waste_count,
                  'telecaller_name':employee.name.upper(),
                  'selected_from': from_date,
                  'selected_to': to_date,
                  'selected_status': status,
                  'selected_pagination': pagination})




def telecaller_closedLeads(request):
    user_id = request.session.get('user_id')

    if not user_id:
        messages.error(request, "Please login first")
        return redirect("login")

    
    user = LogRegister_Details.objects.get(id=user_id)
    employee = EmployeeRegister_Details.objects.get(login=user)
    follow_reason=FollowUpStatus.objects.all()


    closed_leads=list(LeadRow.objects.filter(transferred_to=employee,status="Closed",is_waste=False).order_by('id'))
    for lead in closed_leads:
        updates = list(followUpUpdates.objects.filter(lead=lead).select_related('response').order_by('-updated'))
        lead.second_last_followup_response = updates[1].response.name if len(updates) > 1 and updates[1].response else None
    from_date=request.GET.get('from')
    to_date=request.GET.get('to')
    pagination=request.GET.get('pagination')
    status=request.GET.get('status')

    if from_date and to_date:
        from_date_obj=datetime.strptime(from_date,"%Y-%m-%d").date()
        to_date_obj=datetime.strptime(to_date,"%Y-%m-%d").date()
        closed_leads=[lead for lead in closed_leads if  from_date_obj <= lead.transfer_date.date() <= to_date_obj]
    elif from_date:
        from_date_obj=datetime.strptime(from_date,"%Y-%m-%d").date()
        closed_leads=[lead for lead in closed_leads if  lead.transfer_date.date() >= from_date_obj]
    elif to_date:
        to_date_obj=datetime.strptime(to_date,"%Y-%m-%d").date()
        closed_leads=[lead for lead in closed_leads if lead.transfer_date.date() <= to_date_obj]

    if status:
        closed_leads=[lead for lead in closed_leads if lead.second_last_followup_response == status]

    closed_count=len(closed_leads)

    if pagination in ['10','50','100']:
        paginator=Paginator(closed_leads,int(pagination))
        closed_leads=paginator.get_page(request.GET.get('page'))
    elif pagination == 'All':
        paginator=Paginator(closed_leads,len(closed_leads) or 1)
        closed_leads=paginator.get_page(request.GET.get('page'))
    else:
        pagination = '10'
        paginator=Paginator(closed_leads,10)
        closed_leads=paginator.get_page(request.GET.get('page'))



    if request.method=="POST":
        lead_id=request.POST.getlist('lead_id[]')
        waste_reason=request.POST.get('waste_reason')
        print(waste_reason)
        for id in lead_id:
            lead=LeadRow.objects.get(id=id)
            lead.is_waste=True
            lead.waste_reason=waste_reason
            lead.waste_marked_date=date.today()
            lead.save()
            return redirect('telecaller_wasteLeads')

    return render(request,'telecaller_closedLeads.html',
                  {"closed_leads":closed_leads,
                   "closed_count":closed_count,
                   'telecaller_name':employee.name.upper(),
                   'follow_reason':follow_reason,
                   'selected_from': from_date,
                   'selected_to': to_date,
                   'selected_status': status,
                   'selected_pagination': pagination}) 




def telecaller_followup_details(request,id):
    user_id=request.session.get('user_id')

    if not user_id:
        messages.error(request,"Please login first")
        return redirect('login')
    user=LogRegister_Details.objects.get(id=user_id)
    employee=EmployeeRegister_Details.objects.get(login=user)

    lead=LeadRow.objects.get(id=id)
    followups=followUpUpdates.objects.filter(lead=lead).order_by('-id')
    last_followup = followups.first()
    second_last_followup = followups[1] if followups.count() > 1 else None
    try:
        collections=LeadRowValue.objects.filter(lead_row=lead)
    except:
        collections=None

    return render(request,'telecaller_followup_details.html',
                  {'id':id,
                   'lead':lead,
                   'followups':followups,
                   'last_followup': last_followup,
                   'second_last_followup': second_last_followup,
                   'telecaller_name':employee.name.upper(),
                   'collections':collections,
                   })




def telecaller_reportLeads(request):
    user_id = request.session.get('user_id')

    if not user_id:
        messages.error(request, "Please login first")
        return redirect("login")

    
    user = LogRegister_Details.objects.get(id=user_id)
    employee = EmployeeRegister_Details.objects.get(login=user)

    closed_leads = LeadRow.objects.filter(transferred_to=employee,status__in=['Closed','Joined','Opened','Waste'])

    if request.method == "POST":
        try:
            data={}
            data = json.loads(request.body)
            selected_date = data.get("date")
            leads = closed_leads.filter(transfer_date__date=selected_date).order_by("id")
            rdata = [
                {
                    "id": i.id if i.id else '',  
                    "full_name": i.full_name if i.full_name else 'not available',
                    "email": i.email if i.email else 'not available',
                    "contact": i.contact if i.contact else 'not available',
                    "transfer_date": i.transfer_date.strftime("%Y-%m-%d") if i.transfer_date else "not available",
                    "status": i.followupupdates_set.order_by('-updated').first().response.name if i.followupupdates_set.exists() and i.followupupdates_set.order_by('-updated').first().response else ('Closed' if i.status == 'Closed' else "not available")
                } 
                for i in leads
            ]
            count=len(rdata)
            return JsonResponse({"status": "success", "data": rdata,"count":count})
        except :
            return JsonResponse({"status": "error", "message": "Somthing went wrong"})

    distinct_dates = (closed_leads.filter(transfer_date__isnull=False).annotate(date=TruncDate('transfer_date')).values('date').annotate(count=Count('id')).order_by('date'))

    new_dates = []
    for i in distinct_dates:
        if not i["date"]:
            continue

        new_dates.append({
            'date': i["date"].isoformat(),
            'label': i["date"].strftime("%B %d, %Y"),
            'count': i['count']
        })
    print(new_dates)
    
    return render(request,'telecaller_reportLeads.html',
                  {"closed_leads":closed_leads,
                   "distinct_dates": new_dates,
                   'telecaller_name':employee.name.upper()})



def telecaller_report_details(request,id):
    user_id=request.session.get('user_id')

    if not user_id:
        messages.error(request,"Please login first")
        return redirect('login')
    user=LogRegister_Details.objects.get(id=user_id)
    employee=EmployeeRegister_Details.objects.get(login=user)

    lead=LeadRow.objects.get(id=id)
    followups=followUpUpdates.objects.filter(lead=lead).order_by('-id')
    last_followup = followups.first()
    try:
        collections=LeadRowValue.objects.filter(lead_row=lead)
    except:
        collections=None

    return render(request,'telecaller_report_details.html',
                  {'id':id,
                   'lead':lead,
                   'followups':followups,
                   'telecaller_name':employee.name.upper(),
                   'collections':collections,
                   'last_followup': last_followup,
                   })

def platform_management(request):
    user_id = request.session.get('user_id')
    company_id = request.session.get('company_id')
    
    if not user_id:
        return redirect('login')

    user = LogRegister_Details.objects.get(id=user_id)
    employee = EmployeeRegister_Details.objects.get(login=user)
    company = BusinessRegister_Details.objects.get(id=company_id)

   
    if request.method == "POST":
        platform_name = request.POST.get('platform_name')
        if platform_name:
            Platform.objects.get_or_create(company=company, name=platform_name)
            messages.success(request, "Platform added successfully!")
            return redirect('platform_management')

    platforms = Platform.objects.filter(company=company).annotate(
        total_leads=Count('company__clientregister__workregister__leadcollection__rows', 
        filter=models.Q(company__clientregister__workregister__leadcollection__rows__source__iexact=models.F('name')))
    )

    context = {
        'platforms': platforms,
        'datamanager_name': employee.name,
        'employee': employee
    }
    return render(request, 'platform_management.html', context)






import openpyxl
from django.http import HttpResponse
from django.template.loader import render_to_string
from xhtml2pdf import pisa

def platform_leads_view(request):
    user_id = request.session.get('user_id')
    company_id = request.session.get('company_id')
    source_name = request.GET.get('source', '').strip()

    if not user_id or not company_id or not source_name:
        return redirect('data_manager_dashboard')

    leads = LeadRow.objects.filter(
        lead_collection__work_Id__wcompId_id=company_id,
        source__iexact=source_name
    ).select_related('lead_collection__work_Id__clientId', 'collected_by').order_by('-created_at')
    

    search_query = request.GET.get('search', '').strip()
    start_date = request.GET.get('start_date', '').strip()
    end_date = request.GET.get('end_date', '').strip()
    export_format = request.GET.get('export', '')


    if search_query:
        leads = leads.filter(
            Q(full_name__icontains=search_query) | 
            Q(email__icontains=search_query) | 
            Q(contact__icontains=search_query)
        )
    
    if start_date:
        leads = leads.filter(created_at__date__gte=start_date)
    if end_date:
        leads = leads.filter(created_at__date__lte=end_date)

    if export_format == 'excel':
        response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = f'attachment; filename="{source_name}_leads.xlsx"'
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(['No', 'Collected Date', 'Full name', 'Email', 'Phone', 'Client Name', 'Collected For', 'Collected By', 'Status'])
        for i, lead in enumerate(leads, 1):
            client_name = ""
            category_name = ""
            collected_by_name = ""
            
            if lead.lead_collection and lead.lead_collection.work_Id and lead.lead_collection.work_Id.clientId:
                client_name = lead.lead_collection.work_Id.clientId.client_name
            
            if lead.lead_collection:
                category_name = lead.lead_collection.collection_head
            
            if lead.collected_by:
                collected_by_name = lead.collected_by.name or lead.collected_by.employee_name

            ws.append([
                i, 
                lead.created_at.strftime('%B %d, %Y'), 
                lead.full_name, 
                lead.email, 
                lead.contact, 
                client_name, 
                category_name, 
                collected_by_name, 
                lead.status
            ])
        wb.save(response)
        return response

    if export_format == 'pdf':
        context = {'leads': leads, 'source_name': source_name, 'leads_count': leads.count()}
        html = render_to_string('platform_leads_pdf.html', context)
        response = HttpResponse(content_type='application/pdf')
        response['Content-Disposition'] = f'attachment; filename="{source_name}_leads.pdf"'
        pisa_status = pisa.CreatePDF(html, dest=response)
        return response if not pisa_status.err else HttpResponse("Error generating PDF")

    return render(request, 'platform_leads.html', {
        'leads': leads,
        'source_name': source_name,
        'leads_count': leads.count(),
        'start_date': start_date, 
        'end_date': end_date,     
        'search_query': search_query,
        'datamanager_name': request.session.get('employee_name')
    })



from django.db.models.functions import TruncDate

def platform_reports(request):
    user_id = request.session.get('user_id')
    company_id = request.session.get('company_id')
    
    if not user_id:
        return redirect('login')

    
    from_date = request.GET.get('from_date')
    to_date = request.GET.get('to_date')
    total_company_leads = LeadRow.objects.filter(
        lead_collection__work_Id__wcompId_id=company_id
    ).count()


    platforms = Platform.objects.filter(company_id=company_id)

    report_data = []
    for platform in platforms:
        platform_leads = LeadRow.objects.filter(
            lead_collection__work_Id__wcompId_id=company_id,
            source__iexact=platform.name
        )

        
        if from_date and to_date:
            platform_leads = platform_leads.filter(created_at__date__range=[from_date, to_date])
        elif from_date:
            platform_leads = platform_leads.filter(created_at__date__gte=from_date)
        elif to_date:
            platform_leads = platform_leads.filter(created_at__date__lte=to_date)

       
        daily_breakdown = platform_leads.annotate(date=TruncDate('created_at')) \
            .values('date') \
            .annotate(count=Count('id')) \
            .order_by('-date')

        report_data.append({
            'name': platform.name,
            'collected_count': platform_leads.count(),
            'daily_breakdown': daily_breakdown
        })

    context = {
        'report_data': report_data,
        'total_company_leads': total_company_leads,
        'platforms_count': platforms.count(),
        'from_date': from_date,  
        'to_date': to_date,      
        'datamanager_name': request.session.get('employee_name')
    }
    return render(request, 'platform_reports.html', context)




def notification_management(request):
    user_id = request.session.get('user_id')
    company_id = request.session.get('company_id')
    today = timezone.now().date()

    if not user_id or not company_id:
        return redirect('login')

  
    today_new_leads = LeadRow.objects.filter(
        lead_collection__work_Id__wcompId_id=company_id,
        created_at__date=today
    ).count()

    today_allocations = LeadRow.objects.filter(
        lead_collection__work_Id__wcompId_id=company_id,
        is_transferred=True,
        created_at__date=today 
    ).count()


    follow_up_leads = LeadRow.objects.filter(
        lead_collection__work_Id__wcompId_id=company_id,
        status__in=["Recalled", "Follow Up"] # Adjust based on your status names
    ).count()

    context = {
        'today_new_leads': today_new_leads,
        'today_allocations': today_allocations,
        'follow_up_leads': follow_up_leads,
        'datamanager_name': request.session.get('employee_name'),
    }
    return render(request, 'notification_management.html', context)



def today_new_leads_view(request):
    user_id = request.session.get('user_id')
    company_id = request.session.get('company_id')
    today = timezone.now().date()

    if not user_id or not company_id:
        return redirect('login')

    leads = LeadRow.objects.filter(
        lead_collection__work_Id__wcompId_id=company_id,
        created_at__date=today
    ).select_related('lead_collection__work_Id__clientId', 'collected_by').order_by('-created_at')

    client_f = request.GET.get('client')
    category_f = request.GET.get('category')
    exec_f = request.GET.get('executive')
    source_f = request.GET.get('source')
    search_q = request.GET.get('q', '').strip()

    if client_f:
        leads = leads.filter(lead_collection__work_Id__clientId_id=client_f)
    if category_f:
        leads = leads.filter(lead_collection__collection_head__icontains=category_f)
    if exec_f:
        leads = leads.filter(collected_by_id=exec_f)
    if source_f:
        leads = leads.filter(source__iexact=source_f)
    if search_q:
        leads = leads.filter(
            Q(full_name__icontains=search_q) | Q(email__icontains=search_q) | Q(contact__icontains=search_q)
        )

    export_format = request.GET.get('export', '')

    if export_format == 'excel':
        response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = 'attachment; filename="today_new_leads.xlsx"'
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(['No', 'Status', 'Added Date', 'Lead Name', 'Email', 'Phone Number', 'Data Sourse', 'Client name', 'Lead Category'])
        for i, lead in enumerate(leads, 1):
            client_name = ""
            category_name = ""
            if lead.lead_collection and lead.lead_collection.work_Id and lead.lead_collection.work_Id.clientId:
                client_name = lead.lead_collection.work_Id.clientId.client_name
            if lead.lead_collection:
                category_name = lead.lead_collection.collection_head

            ws.append([i, lead.status, lead.created_at.strftime('%B %d, %Y'), lead.full_name, lead.email, lead.contact, 
                       lead.source, client_name, category_name])
        wb.save(response)
        return response

    if export_format == 'pdf':
        context = {'leads': leads, 'source_name': 'Today New Leads', 'leads_count': leads.count()}
        html = render_to_string('today_new_leads_pdf.html', context)
        response = HttpResponse(content_type='application/pdf')
        response['Content-Disposition'] = 'attachment; filename="today_new_leads.pdf"'
        pisa_status = pisa.CreatePDF(html, dest=response)
        return response if not pisa_status.err else HttpResponse("Error generating PDF")

   
    clients = ClientRegister.objects.filter(compId_id=company_id)
    
    if client_f:
        categories = LeadCollection.objects.filter(work_Id__clientId_id=client_f).values_list('collection_head', flat=True).distinct()
    else:
        categories = LeadCollection.objects.filter(work_Id__wcompId_id=company_id).values_list('collection_head', flat=True).distinct()
    executives = EmployeeRegister_Details.objects.filter(company_id=company_id)
    
    if client_f:
        client_sources = LeadRow.objects.filter(lead_collection__work_Id__clientId_id=client_f).values_list('source', flat=True).distinct()
        platforms = Platform.objects.filter(company_id=company_id, name__in=client_sources)
    else:
        platforms = Platform.objects.filter(company_id=company_id)

    context = {
        'leads': leads,
        'leads_count': leads.count(),
        'clients': clients,
        'categories': categories,
        'executives': executives,
        'platforms': platforms,
        'today_date': today,
        'datamanager_name': request.session.get('employee_name'),
    }
    return render(request, 'today_new_leads.html', context)




def today_allocations_view(request):
    user_id = request.session.get('user_id')
    company_id = request.session.get('company_id')
    today = timezone.now().date()

    if not user_id or not company_id:
        return redirect('login')

  
    leads = LeadRow.objects.filter(
        lead_collection__work_Id__wcompId_id=company_id,
        is_transferred=True,
        transfer_date__date=today 
    ).select_related('lead_collection__work_Id__clientId', 'collected_by', 'transferred_to').order_by('-transfer_date')

   
    status_f = request.GET.get('status')
    employee_f = request.GET.get('employee')
    search_q = request.GET.get('q', '').strip()

    if status_f:
        leads = leads.filter(status=status_f)
    if employee_f:
        leads = leads.filter(transferred_to_id=employee_f)
    if search_q:
        leads = leads.filter(
            Q(full_name__icontains=search_q) | Q(email__icontains=search_q) | Q(contact__icontains=search_q)
        )

    export_format = request.GET.get('export', '')

    if export_format == 'excel':
        response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = 'attachment; filename="today_allocations.xlsx"'
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(['No', 'Status', 'Added Date', 'Lead Name', 'Email', 'Phone Number', 'Data Sourse', 'Client name', 'Lead Category'])
        for i, lead in enumerate(leads, 1):
            client_name = ""
            category_name = ""
            if lead.lead_collection and lead.lead_collection.work_Id and lead.lead_collection.work_Id.clientId:
                client_name = lead.lead_collection.work_Id.clientId.client_name
            if lead.lead_collection:
                category_name = lead.lead_collection.collection_head

            ws.append([i, lead.status, lead.created_at.strftime('%B %d, %Y'), lead.full_name, lead.email, lead.contact, 
                       lead.source, client_name, category_name])
        wb.save(response)
        return response

    if export_format == 'pdf':
        context = {'leads': leads, 'source_name': 'Today Allocations', 'leads_count': leads.count()}
        html = render_to_string('today_allocations_pdf.html', context)
        response = HttpResponse(content_type='application/pdf')
        response['Content-Disposition'] = 'attachment; filename="today_allocations.pdf"'
        pisa_status = pisa.CreatePDF(html, dest=response)
        return response if not pisa_status.err else HttpResponse("Error generating PDF")

    employees = EmployeeRegister_Details.objects.filter(company_id=company_id, login__position='Telecaller')
    
    context = {
        'leads': leads,
        'leads_count': leads.count(),
        'employees': employees,
        'search_query': search_q,
        'selected_status': status_f,
        'selected_employee': employee_f,
        'datamanager_name': request.session.get('employee_name'),
    }
    return render(request, 'today_allocations.html', context)

def delete_all_followups(request,id):
    folloups=followUpUpdates.objects.filter(lead=id)
    folloups.delete()
    return redirect('telecaller_followups')



    